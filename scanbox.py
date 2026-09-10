import ctypes
from ctypes import wintypes
import base64
import io
from collections import Counter
import asyncio
import json
import logging
import os
from pathlib import Path
import re
import shutil
import shlex
import signal
import socket
import subprocess
import sys
import threading
import time
import urllib.request
import urllib.error
from urllib.parse import unquote, urlsplit
import uuid
import platform
import zipfile
import tarfile
import tempfile
from concurrent.futures import ThreadPoolExecutor

from get_active_tab_url import (
    ActiveTabUrlError,
    download_active_browser_document_windows,
    get_active_tab_url,
    is_session_bound_download_url,
    make_browser_download_request,
    open_with_browser_session,
    save_active_browser_document_windows,
    warm_up_active_tab_url_reader,
)

if sys.platform == "win32":
    import winreg
else:
    winreg = None

import docx
if sys.platform == "win32":
    from pdf_to_word import pdf_to_docx
else:
    pdf_to_docx = None
import wx
import wx.adv
from fpdf import FPDF
from PIL import Image, ImageGrab, ImageOps, UnidentifiedImageError

APP_NAME = "ScanBox"
APP_VERSION = "2026.9.3"
UPDATE_MANIFEST_URL = os.environ.get(
    "SCANBOX_UPDATE_MANIFEST_URL",
    "https://api.github.com/repos/sjtaylor82/scanbox/releases/latest",
).strip()
MTMD_RUNTIME_RELEASE_URL = os.environ.get(
    "SCANBOX_MTMD_RUNTIME_RELEASE_URL",
    "https://api.github.com/repos/ggml-org/llama.cpp/releases/tags/b10216",
).strip()


def _release_asset_url(manifest, platform_name=None):
    """Return the current platform's portable ZIP from a GitHub release."""
    platform_name = platform_name or sys.platform
    expected_suffix = (
        "-Windows-x64.zip" if platform_name == "win32"
        else "-macOS.zip" if platform_name == "darwin"
        else ""
    )
    if expected_suffix:
        for asset in manifest.get("assets", []):
            name = str(asset.get("name", ""))
            url = str(asset.get("browser_download_url", "")).strip()
            if name.endswith(expected_suffix) and url.startswith("https://"):
                return url
    return str(manifest.get("html_url") or manifest.get("url", "")).strip()


def _prepare_update_payload(archive_path, platform_name=None):
    """Validate and extract an official portable update into system temp."""
    platform_name = platform_name or sys.platform
    expected = (
        "ScanBox/ScanBox.exe" if platform_name == "win32"
        else "ScanBox.app/Contents/MacOS/ScanBox" if platform_name == "darwin"
        else ""
    )
    if not expected:
        raise RuntimeError("Automatic installation is unavailable on this platform.")
    with zipfile.ZipFile(archive_path) as bundle:
        members = bundle.namelist()
        for name in members:
            path = Path(name.replace("\\", "/"))
            if path.is_absolute() or ".." in path.parts:
                raise ValueError("The update archive contains an unsafe path.")
        if expected not in members:
            raise ValueError("The update archive does not contain the expected application.")
        staging = os.path.join(
            tempfile.gettempdir(), f"scanbox-update-{uuid.uuid4().hex}"
        )
        os.makedirs(staging)
        bundle.extractall(staging)
    return os.path.join(staging, "ScanBox" if platform_name == "win32" else "ScanBox.app")


def _launch_portable_updater(payload_path):
    """Replace this portable build after its running process has exited."""
    if not getattr(sys, "frozen", False):
        raise RuntimeError("Automatic installation is available only in packaged builds.")
    process_id = os.getpid()
    token = uuid.uuid4().hex
    if sys.platform == "win32":
        install_dir = os.path.abspath(BASE)
        script_path = os.path.join(tempfile.gettempdir(), f"scanbox-update-{token}.ps1")

        ready_path = script_path + ".ready"
        log_path = os.path.join(DATA_DIR, "update.log")
        source = os.path.join(RESOURCE_BASE, "portable_updater.ps1")
        # A BOM lets Windows PowerShell correctly read non-ASCII paths.
        with open(source, encoding="utf-8-sig") as source_file:
            script = source_file.read()
        with open(script_path, "w", encoding="utf-8-sig") as script_file:
            script_file.write(script)
        logger.info("Starting portable updater; update log: %s", log_path)
        with open(log_path, "ab") as update_log:
            helper = subprocess.Popen(
                [
                    "powershell.exe", "-NoProfile", "-NonInteractive",
                    "-ExecutionPolicy", "Bypass", "-File", script_path,
                    "-ScanBoxProcessId", str(process_id),
                    "-PayloadPath", payload_path, "-AppDirectory", install_dir,
                    "-ReadyPath", ready_path, "-LogPath", log_path,
                ],
                creationflags=subprocess.CREATE_NO_WINDOW,
                close_fds=True, stdout=update_log, stderr=subprocess.STDOUT,
            )
        deadline = time.monotonic() + 300
        while not os.path.isfile(ready_path):
            if helper.poll() is not None:
                raise RuntimeError(f"Updater stopped during preparation. See {log_path}")
            if time.monotonic() >= deadline:
                helper.terminate()
                helper.wait(timeout=10)
                raise RuntimeError(f"Updater preparation timed out. See {log_path}")
            time.sleep(0.2)
        status = Path(ready_path).read_text(encoding="ascii").strip()
        os.remove(ready_path)
        if status != "ready":
            raise RuntimeError(f"Updater preparation failed. See {log_path}")
        return
    if sys.platform == "darwin":
        current_app = str(Path(sys.executable).resolve().parents[2])
        backup_app = current_app + f".previous-{token}"
        script_path = os.path.join(tempfile.gettempdir(), f"scanbox-update-{token}.sh")
        script = (
            "#!/bin/sh\nset -e\n"
            f"while kill -0 {process_id} 2>/dev/null; do sleep 1; done\n"
            f"mv {shlex.quote(current_app)} {shlex.quote(backup_app)}\n"
            f"mv {shlex.quote(payload_path)} {shlex.quote(current_app)}\n"
            f"open {shlex.quote(current_app)}\n"
        )
        with open(script_path, "w", encoding="utf-8", newline="\n") as script_file:
            script_file.write(script)
        os.chmod(script_path, 0o700)
        subprocess.Popen(
            ["/bin/sh", script_path],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
        return
    raise RuntimeError("Automatic installation is unavailable on this platform.")


if sys.platform == "win32":
    class _NamedPageAccessible(wx.Accessible):
        """Expose a wx.Notebook page title to MSAA clients such as JAWS."""

        def __init__(self, window, title):
            super().__init__(window)
            self._title = title

        def GetName(self, child_id):
            if child_id == wx.ACC_SELF:
                return wx.ACC_OK, self._title
            return wx.ACC_NOT_IMPLEMENTED, ""

        def GetRole(self, child_id):
            if child_id == wx.ACC_SELF:
                return wx.ACC_OK, wx.ROLE_SYSTEM_PROPERTYPAGE
            return wx.ACC_NOT_IMPLEMENTED, wx.ROLE_SYSTEM_PANE


def _set_named_page_accessible(window, title):
    """Add the Windows-specific notebook-page accessibility provider."""
    if sys.platform == "win32":
        window.SetAccessible(_NamedPageAccessible(window, title))

try:
    import numpy as np
except ImportError:
    # Needed by page/edge auto-crop and by the optional Florence-2 vision
    # engine below. Degrade gracefully if not bundled: both features just
    # become unavailable rather than crashing the app.
    np = None

try:
    import cv2
except ImportError:
    # Page/edge detection (for camera-based scanners: document cameras,
    # handheld/overhead scanners such as Pearl/IRIScan-style devices).
    # Degrade gracefully if not bundled: scanning still works, it just
    # skips the auto-crop step.
    cv2 = None

try:
    from pygrabber.dshow_graph import FilterGraph as _DirectShowFilterGraph
except Exception:
    # pygrabber initializes legacy DirectShow COM type libraries during
    # import and can fail even when installed (for example, when its generated
    # wrapper directory is not writable). Camera capture still works through
    # OpenCV; only friendly DirectShow device names become unavailable.
    _DirectShowFilterGraph = None

try:
    import onnxruntime as ort
except ImportError:
    # Optional at source level; packaged builds include the Florence runtime.
    ort = None

try:
    from tokenizers import Tokenizer as _HFTokenizer
except ImportError:
    _HFTokenizer = None

try:
    import fitz
except ImportError:
    fitz = None

if getattr(sys, "frozen", False):
    BASE = os.path.dirname(sys.executable)
    RESOURCE_BASE = getattr(sys, "_MEIPASS", BASE)
else:
    BASE = os.path.dirname(os.path.abspath(__file__))
    RESOURCE_BASE = BASE

MANUAL_PATH = os.path.join(RESOURCE_BASE, "manual.html")
LICENSE_PATH = os.path.join(RESOURCE_BASE, "LICENSE")
SHUTTER_SOUND_PATH = os.path.join(RESOURCE_BASE, "SHUTTER.WAV")
DONATE_URL = (
    "https://www.paypal.com/donate?"
    "business=samtaylor9%40me.com&currency_code=AUD&item_name=ScanBox"
)
MACOS_CAPTURE_HELPER = os.environ.get(
    "SCANBOX_MACOS_CAPTURE_HELPER",
    os.path.join(RESOURCE_BASE, "macos", "scanbox-macos-helper"),
)


def _is_writable(directory):
    try:
        os.makedirs(directory, exist_ok=True)
        probe = os.path.join(directory, f".write_test_{uuid.uuid4().hex}")
        with open(probe, "w"):
            pass
        os.remove(probe)
        return True
    except Exception:
        return False


# Writable per-user state lives next to the app for portable/source use. A
# packaged read-only installation uses the conventional per-user application
# data directory for its operating system.
if sys.platform == "darwin" and getattr(sys, "frozen", False):
    # Never probe or write inside a signed .app bundle: changing even a
    # temporary file beneath Contents invalidates its code signature.
    DATA_DIR = os.path.expanduser("~/Library/Application Support/ScanBox")
elif _is_writable(BASE):
    DATA_DIR = BASE
elif sys.platform == "darwin":
    DATA_DIR = os.path.expanduser("~/Library/Application Support/ScanBox")
else:
    DATA_DIR = os.path.join(
        os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "ScanBox"
    )

TEMP_DIR = os.path.join(DATA_DIR, "temp")
IMAGES_DIR = os.path.join(DATA_DIR, "images")
OUTPUT_DIR = os.path.join(DATA_DIR, "output")
CONFIG_DIR = os.path.join(DATA_DIR, "config")
APP_SETTINGS_CONFIG = os.path.join(CONFIG_DIR, "settings.json")
LOG_FILE = os.environ.get(
    "SCANBOX_LOG_FILE",
    os.path.join(OUTPUT_DIR, "scanbox.log"),
)
PHOTO_LIBRARY_MANIFEST = os.path.join(CONFIG_DIR, "photo_descriptions.json")
VISION_DIR = os.path.join(DATA_DIR, "engines", "vision")
BUNDLED_VISION_DIR = os.path.join(RESOURCE_BASE, "engines", "vision")
FLORENCE_SUBDIR = "florence2-base"
FLORENCE_MANIFEST_FILENAME = "vision_pack_florence2_base.json"
MTMD_RUNTIME_DIR = os.path.join(VISION_DIR, "runtime")
VISION_MODELS = {
    "florence_base": {
        "name": "Florence-2 Base",
        "choice_label": "Smaller model — Florence-2 Base",
        "subdir": FLORENCE_SUBDIR,
        "manifest": FLORENCE_MANIFEST_FILENAME,
        "mac_manifest": "vision_pack_florence2_base_macos.json",
        "runner": "florence",
    },
    "qwen3_vl_2b": {
        "name": "Qwen3-VL 2B",
        "choice_label": "Larger, more accurate model — Qwen3-VL 2B",
        "subdir": "qwen3-vl-2b",
        "manifest": "vision_pack_qwen3_vl_2b.json",
        "runner": "mtmd",
    },
}
if sys.platform == "darwin":
    # Every community-published Florence-2 ONNX decoder export tried for
    # macOS has failed differently: INT8 needs the ConvInteger op, which
    # has no kernel in ONNX Runtime's Apple Silicon CPU provider; the FP16
    # decoder_with_past and decoder_model_merged exports each load (or in
    # the merged case, fail to load - a malformed subgraph) but are
    # statically shaped in ways incompatible with plain one-token
    # generation. Without a working cached decoder, quality and speed both
    # trail Windows significantly even before those export bugs. Qwen3-VL
    # 2B already works reliably on macOS, so it is the only local AI model
    # offered there rather than continuing to ship a degraded Florence-2
    # Base experience.
    del VISION_MODELS["florence_base"]
DEFAULT_VISION_MODEL_ID = "qwen3_vl_2b" if sys.platform == "darwin" else "florence_base"
DEFAULT_APP_SETTINGS = {
    "delete_output_files_on_exit": False,
    "check_for_updates_on_startup": True,
    # Windows has notification-area background behavior. macOS uses the Dock
    # and its native global toggle shortcut, so this Windows preference stays
    # off there.
    # Ordinary minimising must remain visible to Alt+Tab. Only the explicit
    # global toggle shortcut hides the ScanBox window completely.
    "minimize_to_notification_area": False,
    "open_word_after_pdf_conversion": False,
    "render_converted_in_window": True,
    "diagnostic_logging": False,
    "append_text_to_buffer": False,
    # Controls OCR for imported PDF documents; see self.pdf_ocr_radio.
    "pdf_ocr_mode": "none",
    "show_page_headings": False,
    "last_import_dir": "",
    "camera_delay_seconds": 5,
    "camera_interval_seconds": 5,
    # Zero means repeat until the user chooses Stop Camera Capture.
    "camera_capture_count": 1,
    "scanner_id": "",
    "scanner_name": "Ask me each time",
    "vision_model": DEFAULT_VISION_MODEL_ID,
    "ai_provider": "builtin",
    "external_ai_kind": "",
    "external_ai_url": "",
    "external_ai_model": "",
    "external_ai_label": "",
    "mac_permissions_prompted": False,
}

CAMERA_SETTLE_SECONDS = 1.5
CAMERA_FRAME_TIMEOUT_SECONDS = 6.0
CAMERA_READ_INTERVAL_SECONDS = 0.05


def _camera_frame_is_usable(frame):
    """Reject missing and uniform placeholder frames from camera bridges."""
    if frame is None or getattr(frame, "size", 0) == 0:
        return False
    if np is None:
        return True
    try:
        # DirectShow bridge cameras can initially return an all-black,
        # all-white, or otherwise uniform placeholder.
        return float(np.max(frame)) - float(np.min(frame)) >= 2.0
    except (TypeError, ValueError):
        return False


def _read_settled_camera_frame(
    camera,
    settle_seconds=CAMERA_SETTLE_SECONDS,
    timeout_seconds=CAMERA_FRAME_TIMEOUT_SECONDS,
    read_interval=CAMERA_READ_INTERVAL_SECONDS,
):
    """Read over elapsed time so slower DirectShow bridges can deliver a frame."""
    started = time.monotonic()
    deadline = started + max(float(timeout_seconds), 0.0)
    latest = None
    attempts = 0
    while True:
        attempts += 1
        ok, candidate = camera.read()
        if ok and _camera_frame_is_usable(candidate):
            latest = candidate
        now = time.monotonic()
        if latest is not None and now - started >= max(float(settle_seconds), 0.0):
            return latest, attempts
        if now >= deadline:
            return latest, attempts
        time.sleep(max(float(read_interval), 0.0))

# Leave processor capacity for the desktop and screen reader while giving the
# larger multimodal model enough parallelism to avoid looking as though the
# application has frozen. The previous single-thread limit made Qwen spend
# roughly 80 seconds encoding one ordinary photograph.
VISION_THREADS = min(6, max(2, (os.cpu_count() or 4) // 2))
MTMD_IMAGE_TOKENS = 1024
MTMD_TIMEOUT_SECONDS = 180
QWEN_PRELOAD_MINIMUM_RAM = 8 * 1024**3

# Set once a GPU run has failed with a Metal "out of memory" error
# (kIOGPUCommandBufferCallbackErrorOutOfMemory). That error means the Mac's
# GPU genuinely doesn't have enough memory available for a fully
# GPU-offloaded run right now, not that anything is misconfigured, so once
# seen this process falls back to CPU-only inference (-ngl 0) for the rest
# of the session rather than repeating the same failed attempt on every
# photo.
_mtmd_gpu_out_of_memory = False
_mtmd_server_backend = None
_mtmd_active_backend = None
_mtmd_last_backend = None

_MTMD_GPU_OOM_MARKERS = (
    "kIOGPUCommandBufferCallbackErrorOutOfMemory",
    "Insufficient Memory",
)


def _mtmd_diagnostics_show_gpu_oom(diagnostics):
    return diagnostics and any(
        marker in diagnostics for marker in _MTMD_GPU_OOM_MARKERS
    )


_mtmd_server_lock = threading.Lock()
_mtmd_server_process = None
_mtmd_server_port = None
_mtmd_server_model_id = None
_mtmd_server_loading = None
_mtmd_server_job = None

os.makedirs(TEMP_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(CONFIG_DIR, exist_ok=True)
os.makedirs(VISION_DIR, exist_ok=True)

logger = logging.getLogger("scanbox")
logger.setLevel(logging.DEBUG)
logger.propagate = False


def configure_logging(enabled):
    """Turn the diagnostic file log on or off without duplicating handlers."""
    for handler in list(logger.handlers):
        handler.close()
        logger.removeHandler(handler)
    if enabled:
        handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
        handler.setFormatter(
            logging.Formatter("%(asctime)s %(levelname)s %(threadName)s %(message)s")
        )
        logger.addHandler(handler)
        logger.info("Diagnostic logging enabled; data directory: %s", DATA_DIR)

try:
    import win32com.client
    import pythoncom
except ImportError:
    win32com = None
    pythoncom = None

try:
    # PyObjC, used only for a single, non-interactive call at startup:
    # forcing ScanBox to become the frontmost/active application. This is
    # a different situation from the macOS Import picker, where PyObjC
    # turned out not to help - that needed genuine keyboard-event
    # interaction with a panel, which wx's ownership of the process's one
    # NSApplication/event loop blocked no matter which API built the
    # panel. Activating the app is a single fire-and-forget request with
    # no follow-up keyboard interaction, so there's nothing for wx's event
    # loop to interfere with here.
    from AppKit import NSApplicationActivateIgnoringOtherApps, NSRunningApplication
except ImportError:
    NSRunningApplication = None
    NSApplicationActivateIgnoringOtherApps = None

# Windows GUI builds suppress child-process consoles.
if sys.platform == "win32":
    NO_WINDOW = {"creationflags": subprocess.CREATE_NO_WINDOW}
else:
    NO_WINDOW = {}


_win_speaker = None
_mac_announce_process = None
_mac_announce_lock = threading.Lock()
_shutter_sound = None


def _make_speaker():
    """Build a speak(text) callable. Prefer accessible_output2, which routes to
    whichever screen reader is active (NVDA, JAWS, System Access, SAPI). Fall
    back to the bundled NVDA controller DLL, then to a silent no-op."""
    if sys.platform == "darwin":
        def speak_with_voiceover(text):
            # This is the same public VoiceOver ``output`` Apple event used
            # by accessible_output2. Unlike an accessibility notification
            # from a background helper, VoiceOver treats it as speech.
            script = (
                'on run argv\n'
                'tell application "VoiceOver" to output item 1 of argv\n'
                'end run'
            )
            try:
                subprocess.Popen(
                    ["osascript", "-e", script, text],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
            except OSError:
                logger.exception("Could not send a VoiceOver announcement")

        return speak_with_voiceover
    try:
        from accessible_output2.outputs.auto import Auto

        speaker = Auto()
        return lambda text: speaker.output(text)
    except Exception:
        pass
    try:
        dll_path = os.path.join(RESOURCE_BASE, "nvdaControllerClient64.dll")
        nvda = ctypes.windll.LoadLibrary(dll_path)

        def speak(text):
            if nvda.nvdaController_testIfRunning() == 0:
                nvda.nvdaController_speakText(text)

        return speak
    except Exception:
        return lambda text: None


def announce(text):
    """Announce text through the active screen reader.
    Best-effort, never raises, and silent when no screen reader is running."""
    global _win_speaker
    text = text or ""
    if not text:
        return
    try:
        if _win_speaker is None:
            _win_speaker = _make_speaker()
        _win_speaker(text)
    except Exception:
        pass


def play_shutter_sound():
    """Play the capture cue without interrupting screen-reader speech."""
    global _shutter_sound
    if not os.path.isfile(SHUTTER_SOUND_PATH):
        return
    try:
        if _shutter_sound is None:
            _shutter_sound = wx.adv.Sound(SHUTTER_SOUND_PATH)
        if _shutter_sound.IsOk():
            _shutter_sound.Play(wx.adv.SOUND_ASYNC)
    except Exception:
        logger.exception("Could not play shutter sound")


def grab_foreground_window():
    """Capture the visible pixels occupied by the active top-level window.

    Do not use Pillow's ``window=hwnd`` capture here. Chromium and other
    hardware-accelerated applications can return a black or patterned
    placeholder through that window-capture route even while their content is
    plainly visible. Capturing the foreground window's screen rectangle reads
    the composed desktop pixels instead, including GPU-rendered web content.
    """
    if sys.platform != "win32":
        raise RuntimeError(
            "Active-window capture is not available in this macOS preview. "
            "Import an image instead."
        )
    hwnd = ctypes.windll.user32.GetForegroundWindow()
    if not hwnd:
        raise RuntimeError("Windows could not identify the active window.")
    rect = wintypes.RECT()
    # GetWindowRect is DPI-virtualised and can describe only the upper-left
    # portion of a window on a scaled display when its coordinates are passed
    # to a physical-pixel capture API. DWM's extended frame bounds are physical
    # screen coordinates and keep the whole foreground window in the bitmap.
    DWMWA_EXTENDED_FRAME_BOUNDS = 9
    dwm_result = ctypes.windll.dwmapi.DwmGetWindowAttribute(
        hwnd,
        DWMWA_EXTENDED_FRAME_BOUNDS,
        ctypes.byref(rect),
        ctypes.sizeof(rect),
    )
    if dwm_result != 0 and not ctypes.windll.user32.GetWindowRect(
        hwnd, ctypes.byref(rect)
    ):
        raise RuntimeError("Windows could not determine the active window bounds.")
    if rect.right <= rect.left or rect.bottom <= rect.top:
        raise RuntimeError("The active window has no visible capture area.")
    image = ImageGrab.grab(
        bbox=(rect.left, rect.top, rect.right, rect.bottom),
        all_screens=True,
    )
    logger.info(
        "Captured foreground window bounds left=%d top=%d right=%d bottom=%d "
        "image=%dx%d",
        rect.left,
        rect.top,
        rect.right,
        rect.bottom,
        image.width,
        image.height,
    )
    return image


def microsoft_word_available():
    # PDF-to-Word is intentionally Windows-only for the initial macOS port.
    if sys.platform != "win32":
        return False
    if shutil.which("WINWORD.EXE"):
        return True
    subkey = (
        r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE"
    )
    for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
        try:
            with winreg.OpenKey(hive, subkey) as key:
                executable, _kind = winreg.QueryValueEx(key, None)
            if executable and os.path.exists(executable):
                return True
        except OSError:
            continue
    return False


def open_with_default_application(path):
    """Open a file or folder through the current operating system."""
    if sys.platform == "win32":
        os.startfile(path)
        return
    if sys.platform == "darwin":
        subprocess.Popen(["open", path])
        return
    if not wx.LaunchDefaultApplication(path):
        raise OSError(f"Could not open: {path}")


def reveal_in_file_manager(path):
    """Show a file in Explorer or Finder."""
    if sys.platform == "win32":
        subprocess.Popen(["explorer.exe", "/select,", path])
        return
    if sys.platform == "darwin":
        subprocess.Popen(["open", "-R", path])
        return
    open_with_default_application(os.path.dirname(path))


def open_macos_privacy_pane(anchor):
    """Open a macOS Privacy & Security pane."""
    url = f"x-apple.systempreferences:com.apple.preference.security?{anchor}"
    try:
        subprocess.Popen(
            ["open", "-b", "com.apple.systempreferences", url],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except Exception:
        subprocess.Popen(
            ["open", url], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
        )
    subprocess.Popen(
        ["open", "-a", "System Settings"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def reset_macos_scanbox_permissions():
    """Reset this app's user-level TCC entries so macOS will prompt again."""
    bundle_id = "au.com.scanbox.ScanBox"
    result = subprocess.run(
        ["tccutil", "reset", "All", bundle_id],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=30,
    )
    if result.returncode != 0:
        detail = (result.stderr or result.stdout or "").strip()
        raise RuntimeError(detail or "tccutil reset failed.")


def read_app_settings():
    """Load settings.json, coercing each value to match its default's type
    (bool defaults stay bool, string defaults stay string, etc). Earlier this
    unconditionally cast everything to bool, which silently corrupted string
    settings like last_import_dir into True/False."""
    settings = dict(DEFAULT_APP_SETTINGS)
    loaded = {}
    if os.path.exists(APP_SETTINGS_CONFIG):
        try:
            with open(APP_SETTINGS_CONFIG, "r", encoding="utf-8") as f:
                loaded = json.load(f)
            for key, value in loaded.items():
                if key not in settings:
                    continue
                default = DEFAULT_APP_SETTINGS[key]
                if isinstance(default, bool):
                    settings[key] = bool(value)
                elif isinstance(default, str):
                    settings[key] = value if isinstance(value, str) else default
                else:
                    settings[key] = value
        except Exception:
            pass
    # Migrate the former two-state checkbox. Checked meant OCR pages lacking
    # selectable text; unchecked meant no OCR.
    if "pdf_ocr_mode" not in loaded:
        settings["pdf_ocr_mode"] = (
            "missing_text" if loaded.get("use_ocr_enabled", False) else "none"
        )
    if settings["pdf_ocr_mode"] not in {"none", "missing_text", "all_pages"}:
        settings["pdf_ocr_mode"] = "none"
    # Older settings files may contain the former notification-area option.
    # It is deliberately ignored: ordinary minimising must never hide ScanBox.
    settings["minimize_to_notification_area"] = False
    # Retired experimental models may remain in an existing settings file.
    # Fall back cleanly instead of leaving the interface and inference path
    # referring to a model ScanBox no longer offers.
    if settings["vision_model"] == "qwen2_vl_2b":
        settings["vision_model"] = "qwen3_vl_2b"
    elif settings["vision_model"] not in VISION_MODELS:
        settings["vision_model"] = DEFAULT_VISION_MODEL_ID
    return settings


def _coerce_setting_for_save(key, default, value):
    if isinstance(default, bool):
        return bool(value)
    if isinstance(default, str):
        return value if isinstance(value, str) else default
    return value


def write_app_settings(settings):
    data = {
        key: _coerce_setting_for_save(key, default, settings.get(key, default))
        for key, default in DEFAULT_APP_SETTINGS.items()
    }
    with open(APP_SETTINGS_CONFIG, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)


LOCAL_AI_SERVICES = (
    {
        "kind": "ollama",
        "name": "Ollama",
        "models_url": "http://127.0.0.1:11434/api/tags",
        "base_url": "http://127.0.0.1:11434",
    },
    {
        "kind": "openai",
        "name": "LM Studio",
        "models_url": "http://127.0.0.1:1234/v1/models",
        "base_url": "http://127.0.0.1:1234/v1",
    },
    {
        "kind": "openai",
        "name": "vLLM",
        "models_url": "http://127.0.0.1:8000/v1/models",
        "base_url": "http://127.0.0.1:8000/v1",
    },
)


def _normalise_local_ai_url(url):
    """Validate a loopback-only AI service URL and remove trailing slashes."""
    value = str(url or "").strip().rstrip("/")
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or parsed.hostname not in {
        "127.0.0.1", "localhost", "::1",
    }:
        return ""
    return value


def _local_ai_kind_for_url(url):
    try:
        return "ollama" if urlsplit(str(url).strip()).port == 11434 else "openai"
    except ValueError:
        return "openai"


def external_ai_config(settings=None):
    settings = settings or read_app_settings()
    if settings.get("ai_provider") != "external":
        return None
    url = _normalise_local_ai_url(settings.get("external_ai_url", ""))
    model = settings.get("external_ai_model", "").strip()
    if not url or not model:
        return None
    kind = settings.get("external_ai_kind", "openai") or "openai"
    if kind == "openai" and urlsplit(url).path in {"", "/"}:
        url += "/v1"
    return {
        "kind": kind,
        "name": settings.get("external_ai_label", "Local AI") or "Local AI",
        "url": url,
        "model": model,
    }


def _probe_local_ai_service(service, timeout=1.5):
    """Return every model advertised by one known local AI service."""
    try:
        request = urllib.request.Request(
            service["models_url"],
            headers={"User-Agent": f"ScanBox/{APP_VERSION}"},
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.load(response)
        if service["kind"] == "ollama":
            model_ids = [
                str(item.get("model") or item.get("name") or "").strip()
                for item in payload.get("models", [])
                if isinstance(item, dict)
            ]
        else:
            model_ids = [
                str(item.get("id") or "").strip()
                for item in payload.get("data", [])
                if isinstance(item, dict)
            ]
        return [
            {
                "kind": service["kind"],
                "name": service["name"],
                "url": service["base_url"],
                "model": model_id,
                "label": f"{service['name']}: {model_id}",
            }
            for model_id in model_ids
            if model_id
        ]
    except Exception:
        logger.debug(
            "Local AI service was not available: %s",
            service["name"],
            exc_info=True,
        )
        return []


def discover_local_ai_models(timeout=1.5):
    """Discover all models from common local services, preserving service order."""
    with ThreadPoolExecutor(max_workers=len(LOCAL_AI_SERVICES)) as executor:
        groups = list(
            executor.map(
                lambda service: _probe_local_ai_service(service, timeout),
                LOCAL_AI_SERVICES,
            )
        )
    found = []
    seen = set()
    for group in groups:
        for item in group:
            identity = (item["kind"], item["url"], item["model"])
            if identity not in seen:
                seen.add(identity)
                found.append(item)
    return found


def _external_ai_prompt(task, prompt_override=None):
    if prompt_override:
        return prompt_override
    if task == "transcribe":
        return (
            "Transcribe all visible text accurately in reading order. Preserve "
            "headings, lists, and tables using clear plain text or Markdown. Do "
            "not add information that is not visible in the document."
        )
    return (
        "Describe this image accurately and in useful detail for a blind person. "
        "Describe the overall scene, main subjects, actions, important objects, "
        "visible text, and background. Do not invent details."
    )


def run_external_ai_task(
    task, image_path, prompt_override=None, max_tokens=None,
    cancel_event=None, settings=None,
):
    """Send an acquired image to the selected loopback-only AI service."""
    config = external_ai_config(settings)
    if not config:
        return "Local AI service is not configured."
    if cancel_event is not None and cancel_event.is_set():
        return PHOTO_DESCRIPTION_CANCELLED
    try:
        if max_tokens is None:
            max_tokens = 8192 if task == "transcribe" else 512
        # Scanner TIFFs and imported images must use a format the services
        # support. Encode the actual pixels, without modifying the source.
        with Image.open(image_path) as source, io.BytesIO() as image_file:
            ImageOps.exif_transpose(source).convert("RGB").save(image_file, "PNG")
            encoded = base64.b64encode(image_file.getvalue()).decode("ascii")
        prompt = _external_ai_prompt(task, prompt_override)
        if config["kind"] == "ollama":
            endpoint = config["url"] + "/api/chat"
            payload = {
                "model": config["model"],
                "messages": [{
                    "role": "user",
                    "content": prompt,
                    "images": [encoded],
                }],
                "stream": False,
                "options": {"temperature": 0, "num_predict": max_tokens},
            }
        else:
            endpoint = config["url"] + "/chat/completions"
            mime = "image/png"
            payload = {
                "model": config["model"],
                "messages": [{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:{mime};base64,{encoded}"},
                        },
                    ],
                }],
                "max_tokens": max_tokens,
                "temperature": 0,
            }
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Content-Type": "application/json",
                "User-Agent": f"ScanBox/{APP_VERSION}",
            },
        )
        with urllib.request.urlopen(
            request, timeout=MTMD_TIMEOUT_SECONDS
        ) as response:
            result = json.load(response)
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        if config["kind"] == "ollama":
            if result.get("done_reason") == "length" or result.get("done") is False:
                return "Local AI service returned incomplete text; the result was not used."
            text = str(result.get("message", {}).get("content", "")).strip()
        else:
            if result.get("choices", [{}])[0].get("finish_reason") == "length":
                return "Local AI service returned incomplete text; the result was not used."
            text = str(
                result.get("choices", [{}])[0].get("message", {}).get("content", "")
            ).strip()
        return text or "Local AI service returned no text."
    except Exception as exc:
        logger.exception(
            "External local AI task failed service=%s model=%s",
            config["name"],
            config["model"],
        )
        return f"Local AI service could not run: {exc}"


def load_photo_library(path=PHOTO_LIBRARY_MANIFEST):
    """Load stored photo descriptions, tolerating a missing or damaged file."""
    if not os.path.exists(path):
        return []
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        photos = data.get("photos", [])
        return [entry for entry in photos if isinstance(entry, dict)]
    except (OSError, ValueError, TypeError):
        logger.exception("Could not read photo library: %s", path)
        return []


def save_photo_library(entries, path=PHOTO_LIBRARY_MANIFEST):
    """Atomically save the portable photo-description library."""
    directory = os.path.dirname(path)
    os.makedirs(directory, exist_ok=True)
    temporary = path + ".tmp"
    with open(temporary, "w", encoding="utf-8") as f:
        json.dump({"version": 1, "photos": entries}, f, indent=2, ensure_ascii=False)
    os.replace(temporary, path)


def stored_photo_path(photo_path):
    """Prefer a portable relative path for photos located under DATA_DIR."""
    absolute = os.path.abspath(photo_path)
    try:
        relative = os.path.relpath(absolute, DATA_DIR)
        if relative != os.pardir and not relative.startswith(os.pardir + os.sep):
            return relative.replace(os.sep, "/"), "relative"
    except ValueError:
        pass
    return absolute, "absolute"


def resolve_stored_photo_path(entry):
    path = entry.get("path", "")
    if entry.get("pathType") == "relative":
        return os.path.abspath(os.path.join(DATA_DIR, *path.split("/")))
    return os.path.abspath(path) if path else ""


def add_photo_description(photo_path, description, path=PHOTO_LIBRARY_MANIFEST):
    entries = load_photo_library(path)
    stored_path, path_type = stored_photo_path(photo_path)
    entries.append(
        {
            "id": str(uuid.uuid4()),
            "path": stored_path,
            "pathType": path_type,
            "description": description,
            "createdAt": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        }
    )
    save_photo_library(entries, path)
    return entries[-1]


def update_photo_library_path(old_path, new_path, path=PHOTO_LIBRARY_MANIFEST):
    """Point the newest matching entry at a photo's final saved location."""
    entries = load_photo_library(path)
    old_absolute = os.path.normcase(os.path.abspath(old_path))
    replacement, path_type = stored_photo_path(new_path)
    for entry in reversed(entries):
        resolved = resolve_stored_photo_path(entry)
        if resolved and os.path.normcase(os.path.abspath(resolved)) == old_absolute:
            entry["path"] = replacement
            entry["pathType"] = path_type
            save_photo_library(entries, path)
            return True
    return False


def _trim_repetitive_model_output(output):
    """Stop obvious generation loops while retaining the useful first result.

    Catches both plain repeated lines (AAAA) and short alternating/cyclic
    loops (ABABAB, ABCABCABC) that small vision models sometimes fall into,
    e.g. a receipt transcription that keeps bouncing between two timestamps.
    """
    cleaned = []
    history = []  # normalized, non-empty lines seen so far, in order
    for line in output.splitlines():
        normalized = line.strip().casefold()
        if normalized:
            history.append(normalized)
            for period in (1, 2, 3, 4):
                needed = period * 3
                if (
                    len(history) >= needed
                    and history[-period:] == history[-2 * period : -period]
                    and history[-2 * period : -period]
                    == history[-3 * period : -2 * period]
                ):
                    cleaned.append("[Repeated local AI output stopped.]")
                    return "\n".join(cleaned).strip()
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def _trim_repetitive_phrases(text):
    """Second pass: catches hallucinated loops that repeat a short phrase
    within a single line rather than across separate lines, e.g. a
    headstone reading "Loved by all of us, Loved by all of us, ..." fifty
    times over, comma-separated rather than newline-separated. The
    line-level pass above can't see this because it never sees a repeated
    *line* - it's all one line.

    Only phrases of 3+ words qualify, deliberately excluding common 2-word
    idioms ("on and on", "again and again", "round and round") that are
    legitimate English, not model loops.
    """
    words = text.split(" ")
    if len(words) < 9:
        return text
    kept = []
    for word in words:
        kept.append(word)
        for period in range(3, 25):
            needed = period * 3
            if (
                len(kept) >= needed
                and kept[-period:] == kept[-2 * period : -period]
                and kept[-2 * period : -period] == kept[-3 * period : -2 * period]
            ):
                surviving = kept[: len(kept) - 2 * period]
                return (
                    " ".join(surviving).rstrip(" ,.;:")
                    + " [Repeated local AI output stopped.]"
                )
    return text


def _remove_duplicate_sentences(text):
    """Remove exact sentence loops while retaining each distinct observation.

    Qwen can alternate two otherwise complete sentences until its token limit,
    which is not caught by a short repeated-phrase detector. Descriptive prose
    has no useful reason to state the exact same complete sentence twice.
    """
    parts = re.split(r"(?<=[.!?])(?=\s+[A-Z])", text)
    seen = set()
    kept = []
    for part in parts:
        normalized = re.sub(r"\s+", " ", part).strip().casefold()
        if normalized and normalized in seen:
            continue
        if normalized:
            seen.add(normalized)
        kept.append(part)
    return "".join(kept).strip()


# Words legitimately doubled in English idioms ("on and on", "again and
# again", "round and round", "back and forth" - handled separately since the
# words differ). Never trimmed even when the pattern below matches.
_DUPLICATE_WORD_IDIOM_WHITELIST = {"on", "again", "round", "over"}

_duplicate_word_pair_re = re.compile(
    r"\b(\w+)( (?:and|or) )\1\b(?=[\s.,;:!?]|$)", re.IGNORECASE
)


def _trim_duplicate_word_pairs(text):
    """Catches a narrower stutter than the loop-detectors above: greedy
    decoding sometimes lands on the same word twice in a row across a short
    gap, e.g. "the overall mood is somber and somber" - not a repeated
    n-gram (the surrounding words differ each time, so no_repeat_ngram_size
    doesn't block it), just the same word chosen again a couple of tokens
    later. Collapses "X and X" / "X or X" down to a single "X", except for
    idioms where doubling is normal English.
    """

    def _replace(match):
        word = match.group(1)
        if word.lower() in _DUPLICATE_WORD_IDIOM_WHITELIST:
            return match.group(0)
        return word

    return _duplicate_word_pair_re.sub(_replace, text)


# --- Florence-2 (ONNX Runtime) vision engine -------------------------------
# Florence-2 Base runs locally through this in-process code path.

FLORENCE_IMAGE_SIZE = 768
FLORENCE_IMAGE_MEAN = (0.485, 0.456, 0.406)
FLORENCE_IMAGE_STD = (0.229, 0.224, 0.225)
# Derive the decoder layer count from the KV-cache tensors rather than
# hardcoding it, keeping the ONNX runner tolerant of compatible exports.
FLORENCE_MAX_NEW_TOKENS = 1024
FLORENCE_EOS_TOKEN_ID = 2
# Florence-2 uses its trained task prompts rather than arbitrary instructions.
FLORENCE_TASK_PROMPTS = {
    "describe": "Describe with a paragraph what is shown in the image.",
    "transcribe": "What is the text in the image?",
}
# Florence-2's own generation_config.json specifies this (alongside beam
# search, which greedy decoding here doesn't use). Without it, greedy
# decoding is prone to short immediate stutters like "rectangular
# rectangular frame" or "a cemetery or a cemetery" - a single word or short
# phrase repeated exactly once, right next to itself. That's a different,
# milder failure mode than the multi-repeat loops _trim_repetitive_phrases
# catches, and blocking it during generation is more reliable than trying
# to clean it up after the fact.
FLORENCE_NO_REPEAT_NGRAM_SIZE = 3

_florence_cache = {}


def _pick_next_token_no_repeat(logits_row, generated_tokens, ngram_size):
    """Greedy next-token pick that refuses to complete an n-gram already
    generated, mirroring HuggingFace's no_repeat_ngram_size logits
    processor. logits_row is shape (1, vocab); generated_tokens is the
    plain-Python list of token ids so far (not yet including this pick)."""
    logits_row = np.array(logits_row, copy=True)
    if ngram_size > 0 and len(generated_tokens) >= ngram_size - 1:
        prefix = tuple(generated_tokens[-(ngram_size - 1):]) if ngram_size > 1 else ()
        banned = set()
        for i in range(len(generated_tokens) - ngram_size + 1):
            if tuple(generated_tokens[i : i + ngram_size - 1]) == prefix:
                banned.add(generated_tokens[i + ngram_size - 1])
        for token_id in banned:
            logits_row[0, token_id] = -np.inf
    return int(np.argmax(logits_row, axis=-1)[0])


_FLORENCE_FILE_PATTERNS = {
    "vision_encoder": lambda n: n.startswith("vision_encoder") and n.endswith(".onnx"),
    "embed_tokens": lambda n: n.startswith("embed_tokens") and n.endswith(".onnx"),
    "encoder_model": lambda n: n.startswith("encoder_model") and n.endswith(".onnx"),
    "decoder_model": lambda n: (
        n.startswith("decoder_model")
        and "merged" not in n
        and "with_past" not in n
        and n.endswith(".onnx")
    ),
    "decoder_model_merged": lambda n: n.startswith("decoder_model_merged") and n.endswith(".onnx"),
    "decoder_with_past_model": lambda n: (
        n.startswith("decoder_with_past_model") and n.endswith(".onnx")
    ),
    "tokenizer": lambda n: n == "tokenizer.json",
}


def _find_florence_files_in(dirpath):
    """Look for the 5 Florence-2 ONNX pieces plus tokenizer.json directly
    inside dirpath (not recursively - each pack's files sit flat in their own
    folder). Returns a dict of paths, or None if anything is missing."""
    if not os.path.isdir(dirpath):
        return None
    candidates = {key: [] for key in _FLORENCE_FILE_PATTERNS}
    try:
        names = os.listdir(dirpath)
    except OSError:
        return None
    for name in names:
        full = os.path.join(dirpath, name)
        if not os.path.isfile(full):
            continue
        low = name.lower()
        for key, predicate in _FLORENCE_FILE_PATTERNS.items():
            if predicate(low):
                candidates[key].append(full)
    # INT8-quantized conv layers use the ConvInteger op, which has no kernel
    # in ONNX Runtime's Apple Silicon CPU provider (NOT_IMPLEMENTED at
    # session-creation time) - so macOS cannot use the same INT8 pack
    # Windows does. It prefers FP16 instead: full float math, no quantized
    # ops, so it loads correctly and stays much closer to full quality than
    # INT8 or the original Q4F16 pack. This also means a folder that still
    # has files left over from an earlier attempt (Q4F16, or the INT8 pack
    # that downloaded fine but failed to load) won't get picked by mistake -
    # each platform only prefers its own working suffix; the sorted(paths)[0]
    # fallback below only kicks in when nothing matching is present at all.
    preferred_suffix = "_fp16.onnx" if sys.platform == "darwin" else "_int8.onnx"
    found = {}
    for key, paths in candidates.items():
        if not paths:
            continue
        preferred = [path for path in paths if path.lower().endswith(preferred_suffix)]
        found[key] = preferred[0] if preferred else sorted(paths)[0]
    required = {
        "vision_encoder", "embed_tokens", "encoder_model", "decoder_model", "tokenizer"
    }
    # decoder_model_merged_*.onnx (the community pack's single graph that
    # switches between first-pass and cached decoding via a use_cache_branch
    # flag) cannot be loaded by ONNX Runtime on Apple Silicon - it appears to
    # be exported with a fixed token count for browser/WebGPU use. macOS
    # instead requires the ordinary first-pass decoder and looks for the
    # separate decoder_with_past_model_*.onnx graph (a plain, dynamically
    # shaped cached decoder) as an optional speed upgrade; see
    # _get_florence_engine and run_florence_task below.
    has_decoder = "decoder_model" in found if sys.platform == "darwin" else "decoder_model_merged" in found
    if not required.issubset(found) or not has_decoder:
        return None
    return found


_florence_migration_done = False


def _migrate_legacy_flat_florence_install():
    """One-time cleanup: ScanBox used to install Florence-2 straight into
    engines\\vision with no size subfolder. If that's what's sitting there,
    move a Base pack into its current subdirectory. A legacy Large pack is
    deliberately left untouched and is no longer loaded. Never overwrites an
    existing install. Safe to call repeatedly - a no-op once migrated."""
    found = _find_florence_files_in(VISION_DIR)
    if not found:
        return
    try:
        encoder_size = os.path.getsize(found["vision_encoder"])
    except OSError:
        return
    if encoder_size >= 200 * 1024 * 1024:
        return
    target_dir = os.path.join(VISION_DIR, FLORENCE_SUBDIR)
    if _find_florence_files_in(target_dir):
        # Something already installed there - leave the flat files alone
        # rather than risk clobbering a working install.
        return
    try:
        os.makedirs(target_dir, exist_ok=True)
        for path in found.values():
            os.replace(path, os.path.join(target_dir, os.path.basename(path)))
        logger.info("Migrated flat Florence-2 Base install to %s", target_dir)
    except OSError:
        logger.exception("Could not migrate legacy flat Florence-2 install")


def _find_florence_files():
    """Locate the Florence-2 Base pack, including legacy flat installs."""
    global _florence_migration_done
    if not _florence_migration_done:
        _florence_migration_done = True
        try:
            _migrate_legacy_flat_florence_install()
        except Exception:
            logger.exception("Florence-2 legacy install migration failed")

    search_dirs = [
        os.path.join(base, FLORENCE_SUBDIR) for base in _vision_search_dirs()
    ]
    search_dirs.extend(_vision_search_dirs())
    for dirpath in search_dirs:
        found = _find_florence_files_in(dirpath)
        if found:
            return found
    return None


def _get_florence_engine():
    """Lazily load and cache the Florence-2 Base ONNX engine."""
    if "base" in _florence_cache:
        return _florence_cache["base"]

    engine = None
    if ort is not None and _HFTokenizer is not None and np is not None:
        files = _find_florence_files()
        if files:
            try:
                session_options = ort.SessionOptions()
                session_options.intra_op_num_threads = VISION_THREADS
                providers = ["CPUExecutionProvider"]
                required_keys = (
                    "vision_encoder", "embed_tokens", "encoder_model", "decoder_model"
                )
                sessions = {
                    key: ort.InferenceSession(
                        files[key], sess_options=session_options, providers=providers
                    )
                    for key in required_keys
                }
                # decoder_model_merged / decoder_with_past_model are an
                # optional speed upgrade only (real KV-cache decoding
                # instead of replaying the whole sequence every token).
                # Different community-published quantizations of these
                # graphs have turned out to have real, platform-specific
                # problems - a missing ONNX Runtime kernel on Apple Silicon
                # for one quantization, a statically-shaped export unusable
                # for one-token generation for another - so each is loaded
                # in its own try/except: a broken cache-decoder file
                # degrades to the slower but always-correct fallback in
                # run_florence_task instead of disabling local AI entirely.
                for key in ("decoder_model_merged", "decoder_with_past_model"):
                    path = files.get(key)
                    if not path:
                        continue
                    try:
                        sessions[key] = ort.InferenceSession(
                            path, sess_options=session_options, providers=providers
                        )
                    except Exception:
                        logger.exception(
                            "Could not load Florence-2 cached decoder %s; "
                            "falling back to the slower per-token decoder.",
                            key,
                        )
                tokenizer = _HFTokenizer.from_file(files["tokenizer"])
                engine = {"sessions": sessions, "tokenizer": tokenizer}
            except Exception:
                logger.exception("Could not load Florence-2 Base ONNX engine")
                engine = None

    _florence_cache["base"] = engine
    return engine


def florence_ready():
    return _get_florence_engine() is not None


_ONNX_TYPE_TO_NUMPY = {
    "tensor(float)": np.float32 if np is not None else None,
    "tensor(float16)": np.float16 if np is not None else None,
    "tensor(double)": np.float64 if np is not None else None,
}


def _onnx_session_input_dtype(session, input_name, default):
    """Look up the numpy dtype an ONNX Runtime session actually expects for
    a given input. The FP16 Florence-2 pack ScanBox downloads on macOS may
    or may not keep pixel_values as float32 at the graph boundary (exports
    vary); reading it from the session avoids a second guess-and-fail
    round-trip like the one that hit ConvInteger on the INT8 pack."""
    for node in session.get_inputs():
        if node.name == input_name:
            return _ONNX_TYPE_TO_NUMPY.get(node.type, default)
    return default


def _florence_preprocess_image(image_path):
    with Image.open(image_path) as img:
        img = ImageOps.exif_transpose(img).convert("RGB")
        img = img.resize((FLORENCE_IMAGE_SIZE, FLORENCE_IMAGE_SIZE), Image.BICUBIC)
        arr = np.asarray(img).astype(np.float32) / 255.0
        mean = np.array(FLORENCE_IMAGE_MEAN, dtype=np.float32)
        std = np.array(FLORENCE_IMAGE_STD, dtype=np.float32)
        arr = (arr - mean) / std
        arr = arr.transpose(2, 0, 1)  # HWC -> CHW
        return arr[np.newaxis, ...].astype(np.float32)


def run_florence_task(task, image_path, cancel_event=None):
    """Describe or transcribe an image using a locally installed Florence-2
    ONNX pack. Returns plain text on success, or a string starting with
    "Local vision" on failure - the same convention run_vision_task uses,
    so callers don't need to know which engine answered."""
    engine = _get_florence_engine()
    if engine is None:
        return "Local vision is not configured. No Florence-2 pack is installed."

    prompt_text = FLORENCE_TASK_PROMPTS.get(task, FLORENCE_TASK_PROMPTS["describe"])
    sessions = engine["sessions"]
    tokenizer = engine["tokenizer"]

    try:
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        pixel_values = _florence_preprocess_image(image_path)
        vision_input_dtype = _onnx_session_input_dtype(
            sessions["vision_encoder"], "pixel_values", pixel_values.dtype
        )
        if pixel_values.dtype != vision_input_dtype:
            pixel_values = pixel_values.astype(vision_input_dtype)
        input_ids = np.array([tokenizer.encode(prompt_text).ids], dtype=np.int64)

        image_features = sessions["vision_encoder"].run(None, {"pixel_values": pixel_values})[0]
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        prompt_embeds = sessions["embed_tokens"].run(None, {"input_ids": input_ids})[0]

        batch_size, image_token_length = image_features.shape[:2]
        image_attention_mask = np.ones((batch_size, image_token_length), dtype=np.int64)
        prompt_attention_mask = np.ones((batch_size, prompt_embeds.shape[1]), dtype=np.int64)

        inputs_embeds = np.concatenate([image_features, prompt_embeds], axis=1)
        attention_mask = np.concatenate([image_attention_mask, prompt_attention_mask], axis=1)

        encoder_hidden_states = sessions["encoder_model"].run(
            None, {"inputs_embeds": inputs_embeds, "attention_mask": attention_mask}
        )[0]

        # First decoder pass: no KV cache yet.
        decoder_outs = sessions["decoder_model"].run(
            None,
            {
                "inputs_embeds": inputs_embeds[:, -1:],
                "encoder_hidden_states": encoder_hidden_states,
                "encoder_attention_mask": attention_mask,
            },
        )
        # Cross-attention (encoder) K/V are fixed for the whole generation,
        # computed once here. Self-attention (decoder) K/V grow every step,
        # re-read fresh from decoder_outs each iteration below.
        encoder_kv = decoder_outs[1:]
        # 4 tensors per layer (decoder key/value, encoder key/value).
        num_layers = len(encoder_kv) // 4
        decoder_kv = []
        for layer in range(num_layers):
            decoder_kv.extend(encoder_kv[layer * 4 : layer * 4 + 2])

        generated_tokens = []
        max_new_tokens = 256 if task == "describe" else FLORENCE_MAX_NEW_TOKENS
        # Prefer whichever cached decoder the installed pack actually has,
        # rather than branching on platform. decoder_model_merged (Windows
        # pack) combines the first-pass and cached-pass graphs behind a
        # use_cache_branch switch; decoder_with_past_model (Mac pack, also
        # usable anywhere) is a plain graph dedicated to the cached pass and
        # needs no such switch. Both give real KV-cache decoding, unlike the
        # "replay the whole sequence every token" fallback used only when
        # neither is present (e.g. an older Q4F16-only Mac install that has
        # not been re-downloaded yet).
        if "decoder_model_merged" in sessions:
            cached_decoder_key = "decoder_model_merged"
        elif "decoder_with_past_model" in sessions:
            cached_decoder_key = "decoder_with_past_model"
        else:
            cached_decoder_key = None
        decoder_input_embeds = inputs_embeds[:, -1:]
        for _ in range(max_new_tokens):
            if cancel_event is not None and cancel_event.is_set():
                return PHOTO_DESCRIPTION_CANCELLED
            logits = decoder_outs[0]
            next_token = _pick_next_token_no_repeat(
                logits[:, -1, :], generated_tokens, FLORENCE_NO_REPEAT_NGRAM_SIZE
            )
            if next_token == FLORENCE_EOS_TOKEN_ID:
                break
            generated_tokens.append(next_token)

            next_embeds = sessions["embed_tokens"].run(
                None, {"input_ids": np.array([[next_token]], dtype=np.int64)}
            )[0]
            # Kept up to date every iteration (cheap - one concat) even
            # while a cached decoder is in use, so the slow fallback below
            # always has the full sequence ready if the cached decoder
            # fails partway through a generation.
            decoder_input_embeds = np.concatenate(
                [decoder_input_embeds, next_embeds], axis=1
            )

            if cached_decoder_key is not None:
                feed = {
                    "inputs_embeds": next_embeds,
                    "encoder_hidden_states": encoder_hidden_states,
                    "encoder_attention_mask": attention_mask,
                }
                if cached_decoder_key == "decoder_model_merged":
                    feed["use_cache_branch"] = np.array([True], dtype=np.bool_)
                for layer in range(num_layers):
                    decoder_base = layer * 2
                    encoder_base = layer * 4
                    feed[f"past_key_values.{layer}.decoder.key"] = decoder_kv[decoder_base]
                    feed[f"past_key_values.{layer}.decoder.value"] = decoder_kv[decoder_base + 1]
                    feed[f"past_key_values.{layer}.encoder.key"] = encoder_kv[encoder_base + 2]
                    feed[f"past_key_values.{layer}.encoder.value"] = encoder_kv[encoder_base + 3]
                try:
                    decoder_outs = sessions[cached_decoder_key].run(None, feed)
                except Exception:
                    # A cached-decoder ONNX export that loads fine can still
                    # turn out to be unusable for this shape/pattern at run
                    # time (this is exactly what happened with the
                    # statically-shaped decoder_with_past_model export).
                    # Disable it for the rest of this description and fall
                    # through to the slow-but-correct path below for the
                    # current token too, rather than failing the whole
                    # request.
                    logger.exception(
                        "Florence cached decoder %s failed at run time; "
                        "falling back to the slower per-token decoder for "
                        "the rest of this description.",
                        cached_decoder_key,
                    )
                    cached_decoder_key = None
                else:
                    returned_kv = decoder_outs[1:]
                    if len(returned_kv) == num_layers * 2:
                        decoder_kv = returned_kv
                    else:
                        decoder_kv = []
                        for layer in range(num_layers):
                            decoder_kv.extend(returned_kv[layer * 4 : layer * 4 + 2])
                    continue

            decoder_outs = sessions["decoder_model"].run(
                None,
                {
                    "inputs_embeds": decoder_input_embeds,
                    "encoder_hidden_states": encoder_hidden_states,
                    "encoder_attention_mask": attention_mask,
                },
            )

        text = tokenizer.decode(generated_tokens, skip_special_tokens=False)
        text = text.replace("<s>", "").replace("</s>", "").strip()
        if task == "describe" and len(generated_tokens) >= max_new_tokens:
            # Never expose a half sentence when the model reaches its safety
            # ceiling. The recognised-text section can still follow normally.
            endings = [text.rfind(mark) for mark in (".", "?", "!")]
            last_sentence = max(endings)
            if last_sentence >= len(text) // 2:
                text = text[: last_sentence + 1]
            logger.info("Florence description reached its %d-token limit", max_new_tokens)
        return _trim_duplicate_word_pairs(
            _trim_repetitive_phrases(_trim_repetitive_model_output(text))
        )
    except Exception as exc:
        logger.exception("Florence-2 inference failed for %s", image_path)
        return f"Local vision command could not run: {exc}"


def vision_ready(task="describe"):
    if external_ai_config() is not None:
        return True
    if task != "describe":
        return florence_ready()
    model_id = read_app_settings().get("vision_model", DEFAULT_VISION_MODEL_ID)
    model = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
    if model["runner"] == "florence":
        return florence_ready()
    return _find_mtmd_model_files(model_id) is not None and _find_mtmd_runner() is not None


def _find_mtmd_runner():
    runners = _find_mtmd_runners()
    return runners[0] if runners else None


def _mtmd_runtime_needs_repair(platform_name=None):
    """Return whether an installed Qwen pack lacks its preferred runtime."""
    platform_name = platform_name or sys.platform
    runners = _find_mtmd_runners()
    if platform_name == "win32":
        # Older ScanBox versions could leave a working CPU runner behind.
        # Repair that installation too so Windows can use Vulkan-capable GPUs.
        return not any("vulkan" in path.lower() for path in runners)
    return platform_name == "darwin" and not runners


def _find_mtmd_runners():
    names = ("llama-mtmd-cli.exe", "llama-mtmd-cli")
    roots = (MTMD_RUNTIME_DIR, os.path.join(RESOURCE_BASE, "engines", "vision", "runtime"))
    found = []
    for root in roots:
        if not os.path.isdir(root):
            continue
        for dirpath, _, filenames in os.walk(root):
            for name in names:
                if name in filenames:
                    candidate = os.path.join(dirpath, name)
                    if candidate not in found:
                        found.append(candidate)
    external = shutil.which("llama-mtmd-cli") or shutil.which("mtmd-cli")
    if external and external not in found:
        found.append(external)
    # Prefer a broadly compatible GPU build, then the native macOS build,
    # while retaining the CPU runner as an automatic fallback.
    def priority(path):
        lowered = path.lower()
        if "vulkan" in lowered:
            return 0
        if sys.platform == "darwin":
            return 1
        if "cpu" in lowered or os.path.dirname(path) == MTMD_RUNTIME_DIR:
            return 3
        return 2
    return sorted(found, key=priority)


def _mtmd_backend_label(runner, accelerated):
    """Return the user-facing processor backend requested from llama.cpp."""
    if not accelerated:
        return "CPU"
    if "vulkan" in runner.lower():
        return "GPU via Vulkan"
    if sys.platform == "darwin":
        return "GPU via Metal"
    return "GPU"


def _parse_mtmd_devices(output):
    """Parse llama.cpp device rows into names, labels, and free memory."""
    devices = []
    pattern = re.compile(
        r"^\s*(Vulkan\d+):\s*(.*?)\s*\((\d+)\s+MiB,\s*(\d+)\s+MiB free\)\s*$",
        re.IGNORECASE,
    )
    for line in output.splitlines():
        match = pattern.match(line)
        if match:
            devices.append({
                "id": match.group(1),
                "label": match.group(2).strip(),
                "memory_mib": int(match.group(3)),
                "free_mib": int(match.group(4)),
            })
    return devices


def _preferred_mtmd_device(runner):
    """Prefer a likely discrete Vulkan GPU, then the strongest available device."""
    if "vulkan" not in runner.lower():
        return None
    try:
        result = subprocess.run(
            [runner, "--list-devices"], check=False, capture_output=True,
            text=True, encoding="utf-8", errors="replace", timeout=10,
            **NO_WINDOW,
        )
        devices = _parse_mtmd_devices(result.stdout + "\n" + result.stderr)
    except Exception:
        logger.exception("Could not inspect Vulkan devices")
        return None
    if not devices:
        return None

    def preference(device):
        label = device["label"].lower()
        if any(term in label for term in ("geforce", "quadro", "tesla", "rtx", "gtx")):
            tier = 4
        elif any(term in label for term in ("radeon rx", "radeon pro")):
            tier = 4
        elif "nvidia" in label:
            tier = 3
        elif "amd" in label or "radeon" in label:
            tier = 2
        elif "intel" in label:
            tier = 1
        else:
            tier = 0
        return tier, device["free_mib"], device["memory_mib"]

    return max(devices, key=preference)


def local_ai_acceleration_report(model_id):
    """Name the selected model and processor used for real inference."""
    external = external_ai_config()
    if external:
        return (
            f"Active model: {external['model']} via {external['name']}\n"
            "Graphics processor used: managed by the selected local AI service."
        )
    model = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
    if model["runner"] == "florence":
        return (
            f"Active model: {model['name']}\n"
            "Graphics processor used: none; this model uses the CPU."
        )
    if _find_mtmd_model_files(model_id) is None:
        return f"Active model: {model['name']}\nGraphics processor used: model not installed."
    runners = _find_mtmd_runners()
    if not runners:
        return f"Active model: {model['name']}\nGraphics processor used: runner not installed."
    runner = runners[0]
    actual_backend = _mtmd_active_backend or _mtmd_last_backend
    if actual_backend and actual_backend.startswith("GPU"):
        selected = _preferred_mtmd_device(runner)
        processor = selected["label"] if selected else actual_backend
    elif actual_backend == "CPU":
        processor = "none; the most recent description used the CPU"
    else:
        processor = "not yet known; describe an image first"
    return (
        f"Active model: {model['name']}\n"
        f"Graphics processor used: {processor}."
    )


def _find_mtmd_model_files(model_id):
    model = VISION_MODELS.get(model_id)
    if not model or model["runner"] != "mtmd":
        return None
    for root in _vision_search_dirs():
        folder = os.path.join(root, model["subdir"])
        if not os.path.isdir(folder):
            continue
        names = [name for name in os.listdir(folder) if name.lower().endswith(".gguf")]
        language = next((name for name in names if not name.lower().startswith("mmproj")), None)
        projector = next((name for name in names if name.lower().startswith("mmproj")), None)
        if language and projector:
            return os.path.join(folder, language), os.path.join(folder, projector)
    return None


def _total_physical_memory():
    try:
        if sys.platform == "win32":
            class MemoryStatus(ctypes.Structure):
                _fields_ = [
                    ("length", wintypes.DWORD),
                    ("memory_load", wintypes.DWORD),
                    ("total_physical", ctypes.c_ulonglong),
                    ("available_physical", ctypes.c_ulonglong),
                    ("total_page_file", ctypes.c_ulonglong),
                    ("available_page_file", ctypes.c_ulonglong),
                    ("total_virtual", ctypes.c_ulonglong),
                    ("available_virtual", ctypes.c_ulonglong),
                    ("available_extended_virtual", ctypes.c_ulonglong),
                ]
            status = MemoryStatus()
            status.length = ctypes.sizeof(status)
            if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status)):
                return int(status.total_physical)
        if sys.platform == "darwin":
            return int(subprocess.check_output(["sysctl", "-n", "hw.memsize"], text=True).strip())
        page_size = os.sysconf("SC_PAGE_SIZE")
        pages = os.sysconf("SC_PHYS_PAGES")
        return int(page_size * pages)
    except Exception:
        logger.exception("Could not determine physical memory")
        return 0


def _find_server_for_runner(runner):
    name = "llama-server.exe" if sys.platform == "win32" else "llama-server"
    candidate = os.path.join(os.path.dirname(runner), name)
    return candidate if os.path.isfile(candidate) else None


def _bind_process_lifetime_to_scanbox(process):
    """On Windows, make the local server die even if ScanBox crashes."""
    if sys.platform != "win32":
        return None
    try:
        class IoCounters(ctypes.Structure):
            _fields_ = [
                ("read_operations", ctypes.c_ulonglong),
                ("write_operations", ctypes.c_ulonglong),
                ("other_operations", ctypes.c_ulonglong),
                ("read_bytes", ctypes.c_ulonglong),
                ("write_bytes", ctypes.c_ulonglong),
                ("other_bytes", ctypes.c_ulonglong),
            ]

        class BasicLimitInformation(ctypes.Structure):
            _fields_ = [
                ("per_process_user_time", ctypes.c_longlong),
                ("per_job_user_time", ctypes.c_longlong),
                ("limit_flags", wintypes.DWORD),
                ("minimum_working_set", ctypes.c_size_t),
                ("maximum_working_set", ctypes.c_size_t),
                ("active_process_limit", wintypes.DWORD),
                ("affinity", ctypes.c_size_t),
                ("priority_class", wintypes.DWORD),
                ("scheduling_class", wintypes.DWORD),
            ]

        class ExtendedLimitInformation(ctypes.Structure):
            _fields_ = [
                ("basic", BasicLimitInformation),
                ("io", IoCounters),
                ("process_memory_limit", ctypes.c_size_t),
                ("job_memory_limit", ctypes.c_size_t),
                ("peak_process_memory", ctypes.c_size_t),
                ("peak_job_memory", ctypes.c_size_t),
            ]

        kernel32 = ctypes.windll.kernel32
        kernel32.CreateJobObjectW.restype = wintypes.HANDLE
        kernel32.CreateJobObjectW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR]
        kernel32.SetInformationJobObject.argtypes = [
            wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD
        ]
        kernel32.AssignProcessToJobObject.argtypes = [
            wintypes.HANDLE, wintypes.HANDLE
        ]
        kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        job = kernel32.CreateJobObjectW(None, None)
        if not job:
            return None
        information = ExtendedLimitInformation()
        information.basic.limit_flags = 0x00002000  # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        if not kernel32.SetInformationJobObject(
            job, 9, ctypes.byref(information), ctypes.sizeof(information)
        ) or not kernel32.AssignProcessToJobObject(job, int(process._handle)):
            kernel32.CloseHandle(job)
            logger.warning("Windows could not attach crash cleanup to the local image server")
            return None
        logger.info("Local image server crash cleanup attached")
        return job
    except Exception:
        logger.exception("Could not bind the local image server to ScanBox")
        return None


def stop_mtmd_server():
    global _mtmd_server_process, _mtmd_server_port, _mtmd_server_model_id
    global _mtmd_server_loading, _mtmd_server_job, _mtmd_server_backend
    with _mtmd_server_lock:
        process = _mtmd_server_process
        _mtmd_server_process = None
        _mtmd_server_port = None
        _mtmd_server_model_id = None
        _mtmd_server_backend = None
        loading = _mtmd_server_loading
        _mtmd_server_loading = None
        job = _mtmd_server_job
        _mtmd_server_job = None
    if loading is not None:
        loading.set()
    if process is not None and process.poll() is None:
        try:
            logger.info("Stopping persistent local image-model service pid=%s", process.pid)
            if sys.platform == "win32":
                process.terminate()
            else:
                os.killpg(process.pid, signal.SIGTERM)
            process.wait(timeout=5)
        except Exception:
            try:
                if sys.platform != "win32":
                    os.killpg(process.pid, signal.SIGKILL)
                process.kill()
            except Exception:
                pass
    if job is not None and sys.platform == "win32":
        try:
            ctypes.windll.kernel32.CloseHandle(job)
        except Exception:
            pass


def start_mtmd_server(model_id="qwen3_vl_2b"):
    """Load Qwen once and keep it available for later descriptions."""
    global _mtmd_server_process, _mtmd_server_port, _mtmd_server_model_id
    global _mtmd_server_loading, _mtmd_server_job, _mtmd_server_backend
    if model_id != "qwen3_vl_2b" or _total_physical_memory() < QWEN_PRELOAD_MINIMUM_RAM:
        return False
    files = _find_mtmd_model_files(model_id)
    runners = _find_mtmd_runners()
    if not files or not runners:
        return False
    runner = runners[0]
    server = _find_server_for_runner(runner)
    if not server:
        return False
    with _mtmd_server_lock:
        if (
            _mtmd_server_process is not None
            and _mtmd_server_process.poll() is None
            and _mtmd_server_model_id == model_id
        ):
            return True
        if _mtmd_server_loading is not None and not _mtmd_server_loading.is_set():
            # A launch-time preload is already doing the work. Callers can
            # wait for that same service instead of starting a second copy.
            return True
        loading = threading.Event()
        _mtmd_server_loading = loading
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    accelerated = (
        "vulkan" in runner.lower() or sys.platform == "darwin"
    ) and not _mtmd_gpu_out_of_memory
    selected_device = _preferred_mtmd_device(runner) if accelerated else None
    backend = _mtmd_backend_label(runner, accelerated)
    logger.info(
        "Starting persistent %s service; processor=%s device=%s runner=%s",
        model_id,
        backend,
        selected_device["label"] if selected_device else "default",
        runner,
    )
    command = [
        server, "-m", files[0], "--mmproj", files[1],
        "--host", "127.0.0.1", "--port", str(port), "--no-webui",
        "--ctx-size", "4096", "--parallel", "1",
        "-ngl", "99" if accelerated else "0",
        "--threads", str(VISION_THREADS), "--threads-batch", str(VISION_THREADS),
        "--image-min-tokens", str(MTMD_IMAGE_TOKENS),
        "--image-max-tokens", str(MTMD_IMAGE_TOKENS),
    ]
    if selected_device:
        command.extend(["--device", selected_device["id"]])
    creationflags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    try:
        process = subprocess.Popen(
            command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=creationflags,
            start_new_session=sys.platform != "win32",
        )
        job = _bind_process_lifetime_to_scanbox(process)
        with _mtmd_server_lock:
            _mtmd_server_process = process
            _mtmd_server_port = port
            _mtmd_server_model_id = model_id
            _mtmd_server_job = job
        deadline = time.monotonic() + 180
        health_url = f"http://127.0.0.1:{port}/health"
        while time.monotonic() < deadline and process.poll() is None:
            try:
                with urllib.request.urlopen(health_url, timeout=1) as response:
                    if response.status == 200:
                        logger.info(
                            "Persistent %s service ready on local port %d; processor=%s",
                            model_id,
                            port,
                            backend,
                        )
                        _mtmd_server_backend = backend
                        loading.set()
                        return True
            except (OSError, urllib.error.URLError):
                pass
            time.sleep(0.25)
    except Exception:
        logger.exception("Could not start persistent local image model")
    loading.set()
    stop_mtmd_server()
    return False


def _run_mtmd_server_task(prompt, image_path, max_tokens, cancel_event=None):
    global _mtmd_active_backend, _mtmd_last_backend
    with _mtmd_server_lock:
        process = _mtmd_server_process
        port = _mtmd_server_port
        loading = _mtmd_server_loading
    # The preload thread publishes its loading event just before it launches
    # the process. A photo request arriving in that tiny window should wait
    # for the preload, not start a separate one-shot model load.
    while process is None and loading is not None and not loading.wait(0.2):
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        with _mtmd_server_lock:
            process = _mtmd_server_process
            port = _mtmd_server_port
    if process is None or port is None or process.poll() is not None:
        return None
    while loading is not None and not loading.wait(0.2):
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        if process.poll() is not None:
            return None
    if process.poll() is not None:
        return None
    mime = "image/png" if image_path.lower().endswith(".png") else "image/jpeg"
    with open(image_path, "rb") as image_file:
        encoded = base64.b64encode(image_file.read()).decode("ascii")
    payload = json.dumps({
        "messages": [{
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}},
            ],
        }],
        "max_tokens": max_tokens,
        "temperature": 0,
        "repeat_penalty": 1.1,
    }).encode("utf-8")
    request = urllib.request.Request(
        f"http://127.0.0.1:{port}/v1/chat/completions",
        data=payload,
        headers={"Content-Type": "application/json"},
    )
    cancelled = threading.Event()

    def watch_cancel():
        if cancel_event is not None:
            while not cancelled.wait(0.2):
                if cancel_event.is_set():
                    stop_mtmd_server()
                    return

    watcher = threading.Thread(target=watch_cancel, daemon=True)
    watcher.start()
    try:
        backend = _mtmd_server_backend or "CPU"
        _mtmd_active_backend = backend
        with urllib.request.urlopen(request, timeout=MTMD_TIMEOUT_SECONDS) as response:
            result = json.load(response)
        cancelled.set()
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        text = result["choices"][0]["message"]["content"].strip()
        if text:
            _mtmd_last_backend = backend
        return text
    except Exception:
        cancelled.set()
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        logger.exception("Persistent local image-model request failed")
        # Never start the one-shot fallback beside a failed persistent server.
        # Even an unhealthy server may still own the model's full RAM and GPU
        # allocation, which would otherwise load Qwen twice.
        stop_mtmd_server()
        return None
    finally:
        _mtmd_active_backend = None


class DownloadCancelled(Exception):
    """Raised inside a download worker when the user chooses Cancel."""


def _check_download_cancelled(cancel_event):
    if cancel_event is not None and cancel_event.is_set():
        raise DownloadCancelled()


def _download_file(url, destination, name, status_callback=None, cancel_event=None):
    """Download in cancellable chunks instead of blocking in urlretrieve."""
    request = urllib.request.Request(
        url,
        headers={"User-Agent": f"ScanBox/{APP_VERSION}"},
    )
    last_percentage = -1
    with urllib.request.urlopen(request, timeout=30) as response, open(
        destination, "wb"
    ) as output:
        total_size = int(response.headers.get("Content-Length") or 0)
        downloaded = 0
        while True:
            _check_download_cancelled(cancel_event)
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            output.write(chunk)
            downloaded += len(chunk)
            if status_callback and total_size:
                percentage = min(100, int(downloaded * 100 / total_size))
                if percentage != last_percentage:
                    status_callback(
                        f"{name}: {percentage}% of {total_size / (1024 * 1024):.0f} MB"
                    )
                    last_percentage = percentage
    _check_download_cancelled(cancel_event)


def _install_mtmd_runtime(status_callback=None, cancel_event=None):
    existing = _find_mtmd_runners()
    if existing and (sys.platform != "win32" or any("vulkan" in p.lower() for p in existing)):
        return None
    if status_callback:
        status_callback("Downloading the local image model runner...")
    try:
        request = urllib.request.Request(
            MTMD_RUNTIME_RELEASE_URL,
            headers={"User-Agent": f"ScanBox/{APP_VERSION}"},
        )
        with urllib.request.urlopen(request, timeout=30) as response:
            release = json.load(response)
        _check_download_cancelled(cancel_event)
        machine = platform.machine().lower()
        if sys.platform == "win32":
            wanted_groups = (
                ("vulkan", ("win-vulkan-x64.zip",)),
                ("cpu", ("win-cpu-x64.zip",)),
            )
        elif sys.platform == "darwin" and machine in {"arm64", "aarch64"}:
            wanted_groups = (("native", ("macos-arm64.zip", "macos-arm64.tar.gz")),)
        elif sys.platform == "darwin":
            wanted_groups = (("native", ("macos-x64.zip", "macos-x64.tar.gz")),)
        else:
            return "No automatic local image-model runner is available for this system."
        for runtime_kind, wanted in wanted_groups:
            if runtime_kind == "cpu" and existing:
                continue
            if any(runtime_kind in path.lower() for path in existing):
                continue
            asset = next(
                (a for a in release.get("assets", []) if any(a.get("name", "").endswith(x) for x in wanted)),
                None,
            )
            if not asset or not asset.get("browser_download_url"):
                return (
                    f"The {runtime_kind} local image-model runner could not be "
                    "found in ScanBox's tested runtime release."
                )
            if status_callback:
                status_callback(f"Downloading the {runtime_kind} image-model runner...")
            archive = os.path.join(TEMP_DIR, asset["name"])
            try:
                _download_file(
                    asset["browser_download_url"], archive, asset["name"],
                    status_callback, cancel_event,
                )
                _check_download_cancelled(cancel_event)
            except DownloadCancelled:
                try:
                    os.remove(archive)
                except OSError:
                    pass
                raise
            runtime_target = os.path.join(MTMD_RUNTIME_DIR, runtime_kind)
            os.makedirs(runtime_target, exist_ok=True)
            if asset["name"].endswith(".zip"):
                with zipfile.ZipFile(archive) as bundle:
                    bundle.extractall(runtime_target)
            else:
                with tarfile.open(archive, "r:gz") as bundle:
                    runtime_root = os.path.abspath(runtime_target)
                    for member in bundle.getmembers():
                        target = os.path.abspath(os.path.join(runtime_root, member.name))
                        if os.path.commonpath([runtime_root, target]) != runtime_root:
                            raise ValueError("The runner archive contains an unsafe path.")
                    bundle.extractall(runtime_target)
            os.remove(archive)
        installed_runners = _find_mtmd_runners()
        if sys.platform == "win32" and not any(
            "vulkan" in path.lower() for path in installed_runners
        ):
            return (
                "The Vulkan image-model runner downloaded but its program "
                "was not found."
            )
        if not installed_runners:
            return "The image-model runner downloaded but its program was not found."
        return None
    except DownloadCancelled:
        return "AI download cancelled."
    except Exception as exc:
        logger.exception("Could not install multimodal runner")
        return f"The local image-model runner could not be installed: {exc}"


PHOTO_DESCRIPTION_CANCELLED = "Photo description cancelled."


def run_mtmd_task(
    task, image_path, model_id, prompt_override=None, max_tokens=512,
    cancel_event=None,
):
    global _mtmd_gpu_out_of_memory, _mtmd_active_backend, _mtmd_last_backend
    runners = _find_mtmd_runners()
    files = _find_mtmd_model_files(model_id)
    if not runners or not files:
        return "Local vision is not configured. The selected model is not installed."
    if prompt_override:
        prompt = prompt_override
    elif task == "describe":
        prompt = "Describe this image accurately and in useful detail. Do not invent details."
    else:
        prompt = "Transcribe all visible text exactly, preserving reading order and line breaks."
    failures = []
    text = None
    if model_id == "qwen3_vl_2b" and _total_physical_memory() >= QWEN_PRELOAD_MINIMUM_RAM:
        start_mtmd_server(model_id)
        text = _run_mtmd_server_task(
            prompt, image_path, max_tokens, cancel_event
        )
        if text == PHOTO_DESCRIPTION_CANCELLED:
            return text
        if text:
            logger.info("Persistent local multimodal service used")

    creationflags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0

    def run_once(runner, accelerated, selected_device=None):
        logger.info(
            "Starting local multimodal inference; processor=%s runner=%s",
            _mtmd_backend_label(runner, accelerated),
            runner,
        )
        command = [
            runner, "-m", files[0], "--mmproj", files[1], "--image", image_path,
            "-p", prompt, "-n", str(max_tokens), "--temp", "0",
            "--repeat-penalty", "1.1", "--repeat-last-n", "256",
            "-ngl", "99" if accelerated else "0",
            "--threads", str(VISION_THREADS), "--threads-batch", str(VISION_THREADS),
            "--image-min-tokens", str(MTMD_IMAGE_TOKENS),
            "--image-max-tokens", str(MTMD_IMAGE_TOKENS),
        ]
        if selected_device:
            command.extend(["--device", selected_device["id"]])
        process = subprocess.Popen(
            command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, encoding="utf-8", errors="replace",
            creationflags=creationflags,
        )
        started = time.monotonic()
        while True:
            if cancel_event is not None and cancel_event.is_set():
                process.terminate()
                try:
                    process.communicate(timeout=3)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.communicate()
                return None, True
            if time.monotonic() - started >= MTMD_TIMEOUT_SECONDS:
                process.kill()
                process.communicate()
                raise subprocess.TimeoutExpired(command, MTMD_TIMEOUT_SECONDS)
            try:
                stdout, stderr = process.communicate(timeout=0.2)
                break
            except subprocess.TimeoutExpired:
                continue
        return subprocess.CompletedProcess(command, process.returncode, stdout, stderr), False

    for runner in runners if not text else ():
        accelerated = (
            "vulkan" in runner.lower() or sys.platform == "darwin"
        ) and not _mtmd_gpu_out_of_memory
        backend = _mtmd_backend_label(runner, accelerated)
        selected_device = _preferred_mtmd_device(runner) if accelerated else None
        try:
            _mtmd_active_backend = backend
            try:
                result, cancelled = run_once(runner, accelerated, selected_device)
            finally:
                _mtmd_active_backend = None
            if cancelled:
                return PHOTO_DESCRIPTION_CANCELLED
        except subprocess.TimeoutExpired:
            logger.exception("Local multimodal runner timed out: %s", runner)
            failures.append("Image description took too long.")
            continue
        except Exception as exc:
            logger.exception("Local multimodal inference failed: %s", runner)
            failures.append(str(exc))
            continue
        if result.stderr:
            logger.debug("Local multimodal runner diagnostics (%s):\n%s", runner, result.stderr.strip())
        # The GPU genuinely running out of memory (as opposed to a real
        # configuration problem) is worth one automatic CPU-only retry
        # rather than surfacing raw diagnostic text as the photo's
        # "description". Remember it process-wide so later photos and the
        # persistent server skip straight to CPU instead of repeating the
        # same failed GPU attempt each time.
        if accelerated and result.returncode != 0 and _mtmd_diagnostics_show_gpu_oom(result.stderr):
            _mtmd_gpu_out_of_memory = True
            logger.warning(
                "GPU ran out of memory running %s; retrying on CPU.", runner
            )
            backend = _mtmd_backend_label(runner, False)
            try:
                _mtmd_active_backend = backend
                try:
                    result, cancelled = run_once(runner, False)
                finally:
                    _mtmd_active_backend = None
                if cancelled:
                    return PHOTO_DESCRIPTION_CANCELLED
            except subprocess.TimeoutExpired:
                logger.exception("Local multimodal runner timed out: %s", runner)
                failures.append("Image description took too long.")
                continue
            except Exception as exc:
                logger.exception("Local multimodal inference failed: %s", runner)
                failures.append(str(exc))
                continue
            if result.stderr:
                logger.debug("Local multimodal runner diagnostics (%s, CPU retry):\n%s", runner, result.stderr.strip())
        text = (result.stdout or "").strip()
        if result.returncode == 0 and text:
            _mtmd_last_backend = backend
            logger.info("Local multimodal runner used: %s", runner)
            break
        failures.append((result.stderr or text or "No result was returned.").strip())
        text = ""
    if not text:
        detail = failures[-1] if failures else "No result was returned."
        return f"Local vision command could not run: {detail[-800:]}"
    text = _remove_duplicate_sentences(
        _trim_duplicate_word_pairs(
            _trim_repetitive_phrases(_trim_repetitive_model_output(text))
        )
    )
    text = re.sub(
        r"\s*\[Repeated local AI output stopped\.?\]?\s*$", "", text
    ).strip()
    # Some compact GGUF models emit an end token immediately after beginning
    # another sentence. Do not expose that dangling fragment to a screen reader.
    last_sentence = max(text.rfind(mark) for mark in (".", "?", "!"))
    if 0 <= last_sentence < len(text) - 1 and last_sentence >= len(text) // 2:
        text = text[: last_sentence + 1]
    return text


_ORIENTATION_COMMON_WORDS = {
    "a", "about", "all", "an", "and", "are", "as", "at", "be", "been",
    "but", "by", "can", "do", "for", "from", "had", "has", "have", "if",
    "in", "is", "it", "may", "not", "of", "on", "one", "or", "our",
    "page", "so", "that", "the", "their", "there", "they", "this", "to",
    "was", "we", "were", "which", "will", "with", "you", "your",
}


def _orientation_transcript_score(text):
    """Score whether rotation-specific OCR resembles readable English text."""
    words = re.findall(r"[A-Za-z]+(?:['’-][A-Za-z]+)?", text.lower())
    if len(words) < 4:
        return -10000.0
    common = sum(word in _ORIENTATION_COMMON_WORDS for word in words)
    plausible = sum(
        2 <= len(word) <= 20 and any(letter in "aeiouy" for letter in word)
        for word in words
    )
    strange = sum(
        len(word) > 24 or not any(letter in "aeiouy" for letter in word)
        for word in words
    )
    return common * 30 + plausible * 3 + len(text) * 0.02 - strange * 8


def detect_document_rotation(image_path):
    """Return a strongly supported clockwise text rotation, or None."""
    if sys.platform == "darwin":
        return macos_document_rotation(image_path)
    temporary_paths = []
    transcripts = {}
    native_ocr = (
        windows_ocr if sys.platform == "win32"
        else macos_ocr if sys.platform == "darwin"
        else None
    )
    candidates = (
        ("A", 0, None),
        ("B", 90, Image.Transpose.ROTATE_270),
        ("C", 180, Image.Transpose.ROTATE_180),
        ("D", 270, Image.Transpose.ROTATE_90),
    )
    try:
        if native_ocr is not None:
            with Image.open(image_path) as opened:
                source = ImageOps.exif_transpose(opened).convert("RGB")
                for label, _degrees, transpose in candidates:
                    candidate_path = os.path.join(
                        TEMP_DIR,
                        f"orientation_candidate_{uuid.uuid4().hex}.png",
                    )
                    candidate = (
                        source if transpose is None
                        else source.transpose(transpose)
                    )
                    candidate.save(candidate_path, "PNG")
                    temporary_paths.append(candidate_path)
                    transcripts[label] = native_ocr(candidate_path).strip()[:1000]
        scores = {
            label: _orientation_transcript_score(transcripts.get(label, ""))
            for label, _degrees, _transpose in candidates
        }
    finally:
        _remove_quietly(*temporary_paths)
    ranked = sorted(scores, key=scores.get, reverse=True)
    best, second = ranked[:2]
    margin = scores[best] - scores[second]
    logger.info(
        "Orientation transcript_scores=%s best=%s margin=%.1f",
        {label: round(score, 1) for label, score in scores.items()},
        best,
        margin,
    )
    if scores[best] < 20 or margin < 40:
        logger.warning(
            "Document orientation is uncertain for %s: best=%s margin=%.1f",
            image_path,
            best,
            margin,
        )
        return None
    return {"A": 0, "B": 90, "C": 180, "D": 270}[best]


def load_vision_pack(model_id=DEFAULT_VISION_MODEL_ID):
    """Return the install manifest for a supported local model."""
    model = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
    manifest_name = (
        model.get("mac_manifest", model["manifest"])
        if sys.platform == "darwin"
        else model["manifest"]
    )
    builtin_path = os.path.join(RESOURCE_BASE, "config", manifest_name)
    if os.path.exists(builtin_path):
        with open(builtin_path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def _vision_search_dirs():
    dirs = [VISION_DIR]
    if BUNDLED_VISION_DIR != VISION_DIR and os.path.isdir(BUNDLED_VISION_DIR):
        dirs.append(BUNDLED_VISION_DIR)
    return dirs


def install_local_ai_pack(status_callback=None, model_id=DEFAULT_VISION_MODEL_ID, cancel_event=None):
    model = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
    manifest = load_vision_pack(model_id)
    files = manifest.get("files", [])
    if not files:
        return "The local AI pack lists no files to download."

    subdir = manifest.get("subdir", model["subdir"])
    target_dir = os.path.join(VISION_DIR, subdir)
    os.makedirs(target_dir, exist_ok=True)
    downloads_dir = os.path.join(TEMP_DIR, "downloads")
    os.makedirs(downloads_dir, exist_ok=True)
    installed_this_attempt = []

    def cancelled_result():
        for path in installed_this_attempt:
            try:
                os.remove(path)
            except OSError:
                pass
        return "AI download cancelled."

    for item in files:
        try:
            _check_download_cancelled(cancel_event)
        except DownloadCancelled:
            return cancelled_result()
        url = item.get("url", "").strip()
        name = item.get("name", "").strip()
        if not url or not name:
            return "A local AI download entry is missing a URL or filename."
        if "example.com" in url:
            return (
                f"The download URL for {name} is not configured correctly."
            )

        installed_path = os.path.join(target_dir, name)
        if os.path.exists(installed_path):
            if status_callback:
                status_callback(f"{name} is already installed.")
            continue
        if status_callback:
            status_callback(f"Downloading {name}...")
        download_path = os.path.join(downloads_dir, name)
        try:
            _download_file(
                url, download_path, name, status_callback, cancel_event
            )
            _check_download_cancelled(cancel_event)
        except DownloadCancelled:
            try:
                os.remove(download_path)
            except OSError:
                pass
            return cancelled_result()
        except Exception as exc:
            return f"Download failed for {name}: {exc}"

        os.replace(download_path, installed_path)
        installed_this_attempt.append(installed_path)

    if model["runner"] == "mtmd":
        runtime_error = _install_mtmd_runtime(status_callback, cancel_event)
        if runtime_error:
            if runtime_error == "AI download cancelled.":
                return cancelled_result()
            return runtime_error
        if _find_mtmd_model_files(model_id):
            return f"Local AI pack installed and configured. ScanBox will use {model['name']}."
    else:
        _florence_cache.pop("base", None)
    if model["runner"] == "florence" and florence_ready():
        return (
            "Local AI pack installed and configured. ScanBox will use "
            "Florence-2 Base for local image processing."
        )

    return "The selected local AI model downloaded but could not be loaded."


def run_vision_task(
    task, image_path, cancel_event=None, prompt_override=None, max_tokens=None,
):
    """Run the local model selected in Settings."""
    settings = read_app_settings()
    external = external_ai_config(settings)
    if external:
        started = time.perf_counter()
        text = run_external_ai_task(
            task,
            image_path,
            prompt_override,
            max_tokens if max_tokens is not None else (384 if task == "question" else None),
            cancel_event,
            settings,
        )
        elapsed = time.perf_counter() - started
        logger.info(
            "%s model=%s %s completed in %.1f seconds",
            external["name"], external["model"], task, elapsed,
        )
        if (
            task in {"describe", "question"}
            and text != PHOTO_DESCRIPTION_CANCELLED
            and not text.startswith("Local AI service")
        ):
            text += (
                f"\n\nModel: {external['model']} via {external['name']}. "
                f"Processing time: {elapsed:.1f} seconds."
            )
        return text
    if task != "describe":
        return run_florence_task(task, image_path)
    model_id = settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
    model = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
    started = time.perf_counter()
    if model["runner"] == "florence":
        text = run_florence_task(task, image_path, cancel_event)
    elif model_id == "qwen3_vl_2b":
        prompt = (
            "Describe this image accurately and in useful detail for a blind person. "
            "Begin with the overall scene, then describe the main subjects, their "
            "appearance, actions and positions, important objects, and the background. "
            "Describe only details visibly supported by the image, and distinguish actual "
            "subjects from representations when relevant. Do not list objects merely to "
            "say they are absent. Be decisive about clear details. When a detail is small, "
            "blurry, stylised, or partly hidden, describe it more broadly instead of guessing "
            "an exact identity, name, brand, count, or attribute. Do not invent details or "
            "infer thoughts, emotions, stories, or symbolism. State uncertainty when necessary "
            "and avoid repetition. Before finishing, check the full background and image edges "
            "for distinct visible objects that would otherwise be easy to overlook."
        )
        text = run_mtmd_task(
            task, image_path, model_id, prompt, 512, cancel_event
        )
    else:
        text = run_mtmd_task(task, image_path, model_id)
    elapsed = time.perf_counter() - started
    logger.info("%s %s completed in %.1f seconds", model["name"], task, elapsed)
    if text == PHOTO_DESCRIPTION_CANCELLED:
        return text
    if task == "describe" and not text.startswith("Local vision"):
        text += f"\n\nModel: {model['name']}. Processing time: {elapsed:.1f} seconds."
    return text


PAGE_CROP_MIN_AREA_RATIO = 0.15
PAGE_CROP_MAX_AREA_RATIO = 0.97
PAGE_CROP_MAX_DIMENSION = 1000


def _order_page_points(pts):
    """Sort four (x, y) points into top-left, top-right, bottom-right,
    bottom-left order, regardless of the order the contour returned them."""
    rect = np.zeros((4, 2), dtype="float32")
    total = pts.sum(axis=1)
    rect[0] = pts[np.argmin(total)]
    rect[2] = pts[np.argmax(total)]
    diff = np.diff(pts, axis=1)
    rect[1] = pts[np.argmin(diff)]
    rect[3] = pts[np.argmax(diff)]
    return rect


def detect_and_crop_page(image_path):
    """Find a page- or photo-sized rectangle in a camera capture (an
    overhead document camera, or a handheld/desktop scanner such as a
    Pearl/IRIScan-style device that photographs the page rather than
    scanning it through a calibrated driver) and perspective-correct +
    crop it, discarding the desk/background around it.

    Safe no-op on flatbed/WIA scans (already edge-to-edge, so no separate
    page contour exists) and on ordinary photos with no single dominant
    rectangular subject. Returns True if the image was cropped in place,
    False if left unchanged for any reason (including OpenCV not being
    available)."""
    if cv2 is None or np is None:
        return False
    try:
        with Image.open(image_path) as img:
            img = ImageOps.exif_transpose(img).convert("RGB")
            original = np.array(img)
    except Exception:
        logger.exception("Could not open %s for page detection", image_path)
        return False

    height, width = original.shape[:2]
    if height < 50 or width < 50:
        return False

    scale = min(PAGE_CROP_MAX_DIMENSION / float(max(height, width)), 1.0)
    small = cv2.resize(original, (max(int(width * scale), 1), max(int(height * scale), 1)))

    try:
        gray = cv2.cvtColor(small, cv2.COLOR_RGB2GRAY)
        blurred = cv2.GaussianBlur(gray, (5, 5), 0)
        edges = cv2.dilate(cv2.Canny(blurred, 50, 150), None, iterations=2)
        contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    except Exception:
        logger.exception("Page detection edge-finding failed for %s", image_path)
        return False

    if not contours:
        return False

    small_area = small.shape[0] * small.shape[1]
    best_points = None
    best_area = 0
    for contour in sorted(contours, key=cv2.contourArea, reverse=True)[:10]:
        area = cv2.contourArea(contour)
        ratio = area / small_area
        if ratio < PAGE_CROP_MIN_AREA_RATIO or ratio > PAGE_CROP_MAX_AREA_RATIO:
            continue
        perimeter = cv2.arcLength(contour, True)
        approx = cv2.approxPolyDP(contour, 0.02 * perimeter, True)
        if len(approx) == 4 and cv2.isContourConvex(approx) and area > best_area:
            best_points = approx.reshape(4, 2).astype("float32")
            best_area = area

    if best_points is None:
        return False

    rect = _order_page_points(best_points / scale)
    (top_left, top_right, bottom_right, bottom_left) = rect
    max_width = int(max(
        np.linalg.norm(bottom_right - bottom_left),
        np.linalg.norm(top_right - top_left),
    ))
    max_height = int(max(
        np.linalg.norm(top_right - bottom_right),
        np.linalg.norm(top_left - bottom_left),
    ))
    if max_width < 50 or max_height < 50:
        return False

    destination = np.array(
        [
            [0, 0],
            [max_width - 1, 0],
            [max_width - 1, max_height - 1],
            [0, max_height - 1],
        ],
        dtype="float32",
    )
    try:
        matrix = cv2.getPerspectiveTransform(rect, destination)
        warped = cv2.warpPerspective(original, matrix, (max_width, max_height))
        save_kwargs = {}
        if os.path.splitext(image_path)[1].lower() in (".jpg", ".jpeg"):
            save_kwargs["quality"] = 95
        Image.fromarray(warped).save(image_path, **save_kwargs)
    except Exception:
        logger.exception("Could not warp/save cropped page for %s", image_path)
        return False

    logger.info(
        "Auto-cropped page in %s to %sx%s (was %sx%s)",
        image_path,
        max_width,
        max_height,
        width,
        height,
    )
    return True


def detect_and_crop_screen_document(image_path):
    """Conservatively isolate a document displayed inside an application."""
    if cv2 is None or np is None:
        return False

    try:
        with Image.open(image_path) as img:
            rgb = np.array(ImageOps.exif_transpose(img).convert("RGB"))
    except Exception:
        logger.exception("Could not open %s for screen document detection", image_path)
        return False

    height, width = rgb.shape[:2]
    if height < 100 or width < 100:
        return False

    scale = min(PAGE_CROP_MAX_DIMENSION / float(max(height, width)), 1.0)
    small = cv2.resize(
        rgb,
        (max(int(width * scale), 1), max(int(height * scale), 1)),
    )
    hsv = cv2.cvtColor(small, cv2.COLOR_RGB2HSV)
    # Documents displayed on websites may be either light paper or a dark menu
    # with pale lettering. Build candidates from both neutral ranges.
    masks = [
        cv2.inRange(hsv, np.array((0, 0, 150)), np.array((179, 85, 255))),
        cv2.inRange(hsv, np.array((0, 0, 0)), np.array((179, 105, 115))),
    ]
    kernel_size = max(5, int(min(small.shape[:2]) * 0.018))
    if kernel_size % 2 == 0:
        kernel_size += 1
    kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT, (kernel_size, kernel_size)
    )
    contours = []
    for mask in masks:
        closed = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel, iterations=2)
        found, _ = cv2.findContours(
            closed, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        contours.extend(found)
    image_area = float(small.shape[0] * small.shape[1])
    image_center = np.array((small.shape[1] / 2.0, small.shape[0] / 2.0))
    best = None
    best_score = 0.0

    for contour in contours:
        x, y, candidate_width, candidate_height = cv2.boundingRect(contour)
        box_area = candidate_width * candidate_height
        area_ratio = box_area / image_area
        aspect = candidate_width / float(candidate_height)
        if (
            area_ratio < 0.14
            or area_ratio > 0.92
            or candidate_height < small.shape[0] * 0.34
            or candidate_width < small.shape[1] * 0.25
            or not 0.42 <= aspect <= 1.9
        ):
            continue

        contour_fill = cv2.contourArea(contour) / float(max(box_area, 1))
        if contour_fill < 0.42:
            continue
        center = np.array((x + candidate_width / 2.0, y + candidate_height / 2.0))
        center_distance = np.linalg.norm(
            (center - image_center)
            / np.array((small.shape[1], small.shape[0]))
        )
        center_score = max(0.0, 1.0 - center_distance * 2.2)
        score = area_ratio * 2.0 + contour_fill + center_score
        if score > best_score:
            best = (x, y, candidate_width, candidate_height)
            best_score = score

    if best is None:
        # The page may have no visible boundary, but application chrome should
        # still not be treated as document text. Remove the title/tab/address
        # area and narrow window borders before OCRing the content viewport.
        top = min(max(int(height * 0.12), 48), 180)
        left = max(int(width * 0.01), 4)
        right = min(width, width - left)
        bottom = min(height, height - max(int(height * 0.025), 4))
        if right - left < 100 or bottom - top < 100:
            logger.info("Screen OCR did not find a usable content region")
            return False
        Image.fromarray(rgb[top:bottom, left:right]).save(image_path)
        logger.info(
            "Screen OCR used application-content fallback left=%d top=%d "
            "right=%d bottom=%d",
            left,
            top,
            right,
            bottom,
        )
        return True

    x, y, candidate_width, candidate_height = best
    left = max(0, int(x / scale))
    top = max(0, int(y / scale))
    right = min(width, int((x + candidate_width) / scale))
    bottom = min(height, int((y + candidate_height) / scale))
    if right - left < 100 or bottom - top < 100:
        return False

    Image.fromarray(rgb[top:bottom, left:right]).save(image_path)
    logger.info(
        "Screen OCR cropped document region left=%d top=%d right=%d bottom=%d",
        left,
        top,
        right,
        bottom,
    )
    return True


def _remove_quietly(*paths):
    for path in paths:
        if path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass


def clear_temp_directory():
    """Remove session working files on shutdown, but retain the temp folder."""
    temp_root = os.path.abspath(TEMP_DIR)
    if not os.path.isdir(temp_root):
        return

    removed_files = 0
    for dirpath, dirnames, filenames in os.walk(temp_root, topdown=False):
        # Never follow or traverse directory links during cleanup.
        for dirname in list(dirnames):
            directory = os.path.abspath(os.path.join(dirpath, dirname))
            if os.path.islink(directory):
                try:
                    os.unlink(directory)
                except OSError:
                    pass
                continue
            try:
                if os.path.commonpath([temp_root, directory]) == temp_root:
                    os.rmdir(directory)
            except OSError:
                pass

        for filename in filenames:
            target = os.path.abspath(os.path.join(dirpath, filename))
            try:
                if os.path.commonpath([temp_root, target]) == temp_root:
                    os.remove(target)
                    removed_files += 1
            except OSError:
                pass

    logger.info("Shutdown temp cleanup removed %s files from %s", removed_files, temp_root)


def looks_like_ocr_repetition_garbage(text):
    """Detect a stuck-loop OCR/vision-model failure: hallucinated
    near-duplicate words (e.g. engraved/stylized text like a headstone photo
    producing "LOVING DANIEL LOVINGS DANIELS LOVIENS LOVINIROLO..." or
    "ORGINLIGHTNING...ORGANICATION...ORGINALION.ORIGINAL.ORGIANT..."). Each
    hallucinated word is typically a distinct spelling variant rather than
    an exact repeat, and can run together without spaces after punctuation,
    so tokenizing needs to split on any non-letter run, not just spaces.
    What's diagnostic is many words sharing the same short prefix (a "stuck"
    word stem), well beyond what normal prose produces."""
    words = [w.upper() for w in re.findall(r"[A-Za-z]+", text) if len(w) >= 3]
    if len(words) < 10:
        return False
    stem_counts = Counter(w[:3] for w in words)
    _, top_count = stem_counts.most_common(1)[0]
    return (top_count / len(words)) > 0.3


def clean_screen_ocr_text(text):
    """Remove browser metadata and implausibly long OCR tokens."""
    cleaned_lines = []
    for line in text.splitlines():
        kept = []
        for token in line.split():
            lower = token.lower()
            letters_and_digits = re.sub(r"[^a-z0-9]", "", lower)
            if (
                "http://" in lower
                or "https://" in lower
                or "www." in lower
                or (".pdf" in lower and len(token) > 20)
                or len(letters_and_digits) > 80
            ):
                # Browser title/filename text normally precedes the URL on the
                # same OCR line. Drop that prefix as well as the contaminated
                # token, while retaining readable document words that follow.
                kept.clear()
                continue
            kept.append(token)
        cleaned = " ".join(kept).strip()
        if cleaned:
            cleaned_lines.append(cleaned)
    return "\n".join(cleaned_lines).strip()


def windows_ocr(path):
    """Recognise an image with Windows.Media.Ocr, returning an empty string
    when the Windows runtime or a suitable recognition language is unavailable."""
    if sys.platform != "win32":
        return ""

    prepared_path = path
    try:
        with Image.open(path) as source:
            width, height = source.size
            longest = max(width, height)
            if 0 < longest < 2200:
                scale = min(2.0, 2400.0 / longest)
                prepared_path = os.path.splitext(path)[0] + "_windows_ocr.png"
                source.resize(
                    (max(1, int(width * scale)), max(1, int(height * scale))),
                    Image.Resampling.LANCZOS,
                ).save(prepared_path, "PNG")
    except Exception:
        prepared_path = path

    async def recognise():
        from winsdk.windows.graphics.imaging import BitmapDecoder
        from winsdk.windows.media.ocr import OcrEngine
        from winsdk.windows.storage import FileAccessMode, StorageFile

        source = await StorageFile.get_file_from_path_async(
            os.path.abspath(prepared_path)
        )
        stream = await source.open_async(FileAccessMode.READ)
        try:
            decoder = await BitmapDecoder.create_async(stream)
            bitmap = await decoder.get_software_bitmap_async()
            engine = OcrEngine.try_create_from_user_profile_languages()
            if engine is None:
                return ""
            result = await engine.recognize_async(bitmap)
            lines = [line.text.strip() for line in result.lines if line.text.strip()]
            return "\n".join(lines)
        finally:
            stream.close()

    try:
        started = time.perf_counter()
        text = asyncio.run(recognise()).strip()
        logger.info(
            "Windows OCR finished elapsed=%.2fs image=%s",
            time.perf_counter() - started,
            path,
        )
        return text
    except Exception:
        logger.exception("Windows OCR failed for %s", path)
        return ""
    finally:
        if prepared_path != path:
            _remove_quietly(prepared_path)


def macos_ocr(path):
    """Recognise image text with Apple's on-device Vision framework."""
    if sys.platform != "darwin" or not os.path.isfile(MACOS_CAPTURE_HELPER):
        return ""
    try:
        started = time.perf_counter()
        result = subprocess.run(
            [MACOS_CAPTURE_HELPER, "ocr", os.path.abspath(path)],
            capture_output=True,
            text=True,
            timeout=60,
            **NO_WINDOW,
        )
        if result.returncode != 0:
            logger.warning("macOS Vision OCR failed: %s", result.stderr.strip())
            return ""
        logger.info(
            "macOS Vision OCR finished elapsed=%.2fs image=%s",
            time.perf_counter() - started,
            path,
        )
        return result.stdout.strip()
    except Exception:
        logger.exception("macOS Vision OCR failed for %s", path)
        return ""


def macos_document_rotation(path):
    """Infer physical text direction using Vision character geometry."""
    if sys.platform != "darwin" or not os.path.isfile(MACOS_CAPTURE_HELPER):
        return None
    try:
        started = time.perf_counter()
        result = subprocess.run(
            [MACOS_CAPTURE_HELPER, "orientation", os.path.abspath(path)],
            capture_output=True,
            text=True,
            timeout=60,
            **NO_WINDOW,
        )
        value = result.stdout.strip()
        rotation = int(value) if value in {"0", "90", "180", "270"} else None
        logger.info(
            "macOS Vision orientation finished elapsed=%.2fs image=%s result=%s",
            time.perf_counter() - started,
            path,
            rotation if rotation is not None else "unknown",
        )
        if result.returncode != 0:
            logger.warning(
                "macOS Vision orientation failed: %s", result.stderr.strip()
            )
            return None
        return rotation
    except Exception:
        logger.exception("macOS Vision orientation failed for %s", path)
        return None


def macos_choose_files_via_osascript(
    title, initial_directory, multiple=False, extensions=()
):
    """Show Apple's Open dialog via AppleScript's `choose file`, run
    through `osascript` as a genuinely separate OS process with its own
    AppKit event loop.

    Two earlier approaches - wx.FileDialog, and NSOpenPanel called
    in-process via PyObjC - looked different in code but behaved
    identically for VoiceOver: no arrow-key or type-to-select navigation,
    no working search field, poor tab order. Both ran inside ScanBox's own
    process, and wxWidgets owns the single NSApplication instance and main
    event loop for that whole process, so neither could actually escape
    wx's own keyboard/event handling. `choose file` is a long-established
    system utility that opens its own window with its own, real event
    loop, entirely outside wx's control.

    Returns a list of selected paths (empty if the user cancelled), or
    None if osascript itself is unavailable/failed unexpectedly, so the
    caller can fall back to wx.FileDialog.
    """
    if sys.platform != "darwin":
        return None
    lines = [f'set thePrompt to "{_applescript_quote(title)}"']
    choose_clause = "choose file with prompt thePrompt"
    if extensions:
        type_list = ", ".join(f'"{ext}"' for ext in extensions)
        lines.append(f"set theTypes to {{{type_list}}}")
        choose_clause += " of type theTypes"
    if initial_directory and os.path.isdir(initial_directory):
        quoted_dir = _applescript_quote(initial_directory)
        lines.append(f'set theLocation to POSIX file "{quoted_dir}"')
        choose_clause += " default location theLocation"
    if multiple:
        choose_clause += " with multiple selections allowed"
    lines.append(f"set theResult to {choose_clause}")
    lines.append(
        "if (class of theResult) is list then\n"
        "    set posixPaths to {}\n"
        "    repeat with anItem in theResult\n"
        "        set end of posixPaths to POSIX path of anItem\n"
        "    end repeat\n"
        "else\n"
        "    set posixPaths to {POSIX path of theResult}\n"
        "end if\n"
        "set AppleScript's text item delimiters to linefeed\n"
        "return posixPaths as text"
    )
    script = "\n".join(lines)
    try:
        result = subprocess.run(
            ["osascript", "-e", script],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if result.returncode != 0:
            stderr = (result.stderr or "").strip()
            if "-128" in stderr:
                # The user cancelled the dialog.
                return []
            logger.warning("osascript file chooser failed: %s", stderr)
            return None
        return [line for line in result.stdout.splitlines() if line]
    except Exception:
        logger.exception("osascript file chooser failed")
        return None


def _applescript_quote(text):
    """Escape a plain string for safe embedding inside a double-quoted
    AppleScript string literal."""
    return text.replace("\\", "\\\\").replace('"', '\\"')


def choose_macos_files_async(parent, title, start_dir, completed, multiple=False, extensions=()):
    """Run macos_choose_files_via_osascript() on a background thread and
    deliver the result back to the UI thread via wx.CallAfter.

    osascript blocks for as long as its own window is open, which may be
    a long time while the user browses. Waiting for it on wx's own thread
    would stop wx's event loop from pumping for that whole time, which
    macOS would report as ScanBox being unresponsive even though a
    separate, genuinely working window is showing. A background thread
    avoids that.
    """

    def choose():
        paths = macos_choose_files_via_osascript(
            title, start_dir, multiple=multiple, extensions=extensions
        )
        if paths is None:
            wx.CallAfter(
                wx.MessageBox,
                "ScanBox could not open the native file chooser. Please "
                "check the log and try again.",
                "Import failed",
            )
            return
        wx.CallAfter(completed, paths)

    threading.Thread(
        target=choose, name="ScanBox macOS file chooser", daemon=True
    ).start()


def macos_capture_window(window_id, output_path):
    """Capture a macOS window from the main ScanBox app process."""
    try:
        from AppKit import NSBitmapImageFileTypePNG, NSBitmapImageRep
        from Quartz import (
            CGRectNull,
            CGWindowListCreateImage,
            kCGWindowImageBoundsIgnoreFraming,
            kCGWindowImageDefault,
            kCGWindowListOptionIncludingWindow,
            kCGWindowListOptionOnScreenOnly,
        )

        image = CGWindowListCreateImage(
            CGRectNull,
            kCGWindowListOptionIncludingWindow,
            int(window_id),
            kCGWindowImageBoundsIgnoreFraming,
        )
        if image is None:
            logger.warning("Quartz window capture returned no image for window %s", window_id)
            image = CGWindowListCreateImage(
                CGRectNull,
                kCGWindowListOptionOnScreenOnly,
                0,
                kCGWindowImageDefault,
            )
        if image is None:
            raise RuntimeError("Screen Recording permission was not granted.")
        rep = NSBitmapImageRep.alloc().initWithCGImage_(image)
        data = rep.representationUsingType_properties_(NSBitmapImageFileTypePNG, {})
        if data is None or not data.writeToFile_atomically_(output_path, True):
            raise RuntimeError("The captured screen image could not be saved.")
        logger.info("macOS Quartz captured window/display to %s", output_path)
        return output_path
    except ImportError:
        logger.exception("Quartz capture support is not installed")
    except Exception:
        logger.exception("macOS Quartz capture failed")

    command = ["/usr/sbin/screencapture", "-x", "-l", str(window_id), output_path]
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=30,
    )
    if result.returncode == 0 and os.path.isfile(output_path):
        logger.info("macOS captured window %s to %s", window_id, output_path)
        return output_path
    detail = (result.stderr or result.stdout or "").strip()
    logger.warning("macOS window capture failed: %s", detail or result.returncode)
    # Some windows cannot be captured by id even when Screen Recording is
    # granted. Fall back to the visible display so the global workflow still
    # works; this is less precise, but it is far better than a dead shortcut.
    result = subprocess.run(
        ["/usr/sbin/screencapture", "-x", output_path],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=30,
    )
    if result.returncode != 0 or not os.path.isfile(output_path):
        detail = (result.stderr or result.stdout or "").strip()
        raise RuntimeError(detail or "Screen Recording permission was not granted.")
    logger.info("macOS captured full screen to %s after window capture fallback", output_path)
    return output_path


_INVALID_XML_TEXT = re.compile(
    r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]"
)


def xml_safe_text(text):
    """Remove characters that cannot be represented in DOCX XML.

    PDF extraction and OCR can occasionally return embedded NUL bytes or
    other control characters. They are not visible to the user, but lxml
    rejects them when python-docx creates a paragraph.
    """
    return _INVALID_XML_TEXT.sub("", str(text))


def suggested_document_filename(text, extension="txt"):
    """Return a conservative, editable filename derived from local OCR."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in str(text).splitlines()]
    lines = [line for line in lines if line]
    lower = "\n".join(lines).lower()
    kind = next((label for marker, label in (
        ("receipt", "Receipt"), ("invoice", "Invoice"),
        ("statement", "Statement"), ("bill", "Bill"), ("dear ", "Letter"),
    ) if marker in lower), "Scanned document")
    issuer = next((line for line in lines[:6] if 3 <= len(line) <= 60 and not
                   re.search(r"\d{2,}|receipt|invoice|statement|bill", line, re.I)), "")
    body = "\n".join(lines)
    date_match = re.search(
        r"\b(?:\d{1,2}[ /.-](?:\d{1,2}|[A-Za-z]{3,9})[ /.-]\d{2,4}|"
        r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}|[A-Za-z]{3,9}\s+\d{4})\b",
        body, re.I,
    )
    total_match = re.search(r"\btotal\b[^\d$]{0,12}(\$?\s*\d+(?:[.,]\d{2})?)", body, re.I)
    parts = [kind]
    if kind != "Scanned document" and issuer and issuer.lower() != kind.lower():
        parts.append(issuer)
    if date_match:
        parts.append(date_match.group(0))
    if kind in {"Receipt", "Invoice", "Bill"} and total_match:
        parts.append(total_match.group(1).replace(" ", ""))
    stem = " - ".join(parts)
    stem = "".join(c if c not in '<>:"/\\|?*' else "_" for c in stem).rstrip(" .")
    stem = stem[:120] or "Scanned document"
    extension = extension.lstrip(".")
    return f"{stem}.{extension}" if extension else stem


def save_text_output(text, path, fmt):
    if fmt == "txt":
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
    elif fmt == "docx":
        doc = docx.Document()
        for block in xml_safe_text(text).split("\n\n"):
            doc.add_paragraph(block)
        doc.save(path)
    elif fmt == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        safe = text.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 10, safe)
        pdf.output(path)


class PdfNeedsOcrError(Exception):
    """Raised when imported-document OCR is off but a PDF has no selectable text.

    PDF2Word can't invent text that isn't there without OCR, so converting
    anyway would just produce and open an empty (or near-empty) Word
    document with no indication of why. Caught separately from a normal
    processing failure so the user gets a specific, actionable message
    instead."""


class ScanBoxTaskBarIcon(wx.adv.TaskBarIcon):
    """Notification-area access while the main window is hidden."""

    def __init__(self, frame):
        super().__init__()
        self.frame = frame
        self.open_id = wx.NewIdRef()
        self.settings_id = wx.NewIdRef()
        self.exit_id = wx.NewIdRef()
        self.Bind(wx.adv.EVT_TASKBAR_LEFT_DCLICK, self._open)
        self.Bind(wx.EVT_MENU, self._open, id=self.open_id)
        self.Bind(wx.EVT_MENU, self._settings, id=self.settings_id)
        self.Bind(wx.EVT_MENU, self._exit, id=self.exit_id)
        icon = wx.ArtProvider.GetIcon(
            wx.ART_INFORMATION, wx.ART_OTHER, (16, 16)
        )
        self.SetIcon(icon, f"{APP_NAME} {APP_VERSION}")

    def CreatePopupMenu(self):
        menu = wx.Menu()
        open_label = "&Open ScanBox"
        if sys.platform == "win32":
            open_label += "\tCtrl+Alt+\\"
        menu.Append(self.open_id, open_label)
        menu.Append(self.settings_id, "&Settings")
        menu.AppendSeparator()
        menu.Append(self.exit_id, "E&xit ScanBox")
        return menu

    def _open(self, event=None):
        self.frame.restore_from_notification_area()

    def _settings(self, event=None):
        self.frame.restore_from_notification_area(self.frame.open_settings)

    def _exit(self, event=None):
        self.frame.Close()


class ScanBox(wx.Frame):
    TAB_SCAN = 0
    TAB_IMPORT = 1
    TAB_LIBRARY = 2

    def __init__(self):
        super().__init__(None, title="ScanBox", size=(940, 780))

        self.session_files = []
        self.session_file_sources = []
        self.session_file_modes = []
        self.session_pages = []
        self.session_page_modes = []
        self.session_photo_descriptions = {}
        self.page_counter = 0
        self.installing = False
        self._installing_model_id = None
        self.install_cancel_event = None
        self._close_when_install_stops = False
        self.photo_cancel_event = None
        self.shutdown_event = threading.Event()
        self._skip_exit_file_cleanup = False
        self.photo_cancel_dialog = None
        self.photo_cancel_button = None
        self.last_active_mode = "document"
        self.busy = False
        self.screen_capture_busy = False
        self._screen_result_placeholder = None
        self.busy_status_visible = False
        self.busy_status_token = 0
        self.app_settings = read_app_settings()
        force_logging = os.environ.get("SCANBOX_FORCE_LOGGING") == "1"
        configure_logging(
            force_logging
            or self.app_settings.get("diagnostic_logging", False)
        )
        if force_logging:
            logger.info("macOS test logging forced; log file: %s", LOG_FILE)
        self.created_output_files = []
        self.tray_icon = None
        self._restore_focus = None
        self._mac_helper_process = None
        self.camera_capture_timer = wx.Timer(self)
        self.Bind(
            wx.EVT_TIMER,
            self.on_camera_countdown,
            self.camera_capture_timer,
        )
        self.camera_capture_active = False
        self.camera_capture_stopping = False
        self.camera_capture_remaining = 0
        self.camera_capture_completed = 0
        self.camera_capture_target = 1
        self.camera_capture_interval = 5
        self.camera_capture_index = None
        self.camera_capture_device = None
        self.camera_capture_photo_mode = False
        self.camera_alignment_active = False
        self.camera_alignment_stop_event = None
        self.facealign_dialog = None
        self.facealign_status = None
        self.facealign_restore_focus = None
        self.image_export_active = False
        self.image_export_multiple = False
        self.image_export_paths = []
        self.image_export_rotations = {}
        self.browser_url_executor = ThreadPoolExecutor(
            max_workers=1,
            thread_name_prefix="ScanBox browser accessibility",
        )
        self.browser_url_executor.submit(self._warm_browser_url_reader)

        menu_bar = wx.MenuBar()
        help_menu = wx.Menu()
        manual_item = help_menu.Append(wx.ID_HELP, "&User Manual\tF1")
        permissions_item = None
        if sys.platform == "darwin":
            permissions_item = help_menu.Append(wx.ID_ANY, "Mac &Permissions")
        acceleration_item = help_menu.Append(wx.ID_ANY, "Local AI &Acceleration...")
        self.check_updates_item = help_menu.Append(wx.ID_ANY, "Check for &Updates")
        donate_item = help_menu.Append(wx.ID_ANY, "&Donate to Project")
        license_item = help_menu.Append(wx.ID_ANY, "&License")
        help_menu.AppendSeparator()
        about_item = help_menu.Append(wx.ID_ABOUT, "&About ScanBox...")
        menu_bar.Append(help_menu, "&Help")
        quit_item = None
        if sys.platform == "darwin":
            # wxWidgets moves a wx.ID_EXIT menu item into the application
            # menu and wires Command+Q to it automatically on macOS -  but
            # only if one exists somewhere in the menu bar. ScanBox had no
            # File/Quit menu at all, so Command+Q silently did nothing; the
            # only way to close the app was Cmd+Tab to something else and
            # quit from there, or force-quit. "Ctrl+Q" in a wx accelerator
            # is the standard way to request Command+Q specifically on Mac
            # (wx maps Ctrl to Command in accelerators on this platform).
            quit_item = help_menu.Append(wx.ID_EXIT, "&Quit ScanBox\tCtrl+Q")
        self.SetMenuBar(menu_bar)
        self.Bind(wx.EVT_MENU, self.open_user_manual, manual_item)
        if permissions_item is not None:
            self.Bind(wx.EVT_MENU, self.show_macos_permissions, permissions_item)
        self.Bind(wx.EVT_MENU, self.show_local_ai_acceleration, acceleration_item)
        self.Bind(wx.EVT_MENU, self.check_for_updates, self.check_updates_item)
        self.Bind(wx.EVT_MENU, self.open_donation_page, donate_item)
        self.Bind(wx.EVT_MENU, self.open_license, license_item)
        self.Bind(wx.EVT_MENU, self.show_about, about_item)
        if quit_item is not None:
            self.Bind(wx.EVT_MENU, lambda event: self.Close(), quit_item)

        panel = wx.Panel(self)
        root = wx.BoxSizer(wx.VERTICAL)

        self.mode_tabs = wx.Notebook(panel)

        # Scan tab - capture straight from a camera/scanner.
        self.scan_panel = wx.Panel(self.mode_tabs, style=wx.TAB_TRAVERSAL)
        _set_named_page_accessible(self.scan_panel, "Scan")
        scan_sizer = wx.WrapSizer(wx.HORIZONTAL)
        self.document_scan_btn = wx.Button(
            self.scan_panel, label="Scan and Read Document"
        )
        self.scan_save_images_btn = wx.Button(
            self.scan_panel, label="Scan and Save Images"
        )
        self.document_camera_btn = wx.Button(
            self.scan_panel, label="OCR using Camera"
        )
        self.photo_scan_btn = wx.Button(self.scan_panel, label="Scan Photo")
        self.photo_camera_btn = wx.Button(self.scan_panel, label="Take Photo")
        self.document_scan_btn.SetToolTip(
            "Scan a document and read its text."
        )
        self.scan_save_images_btn.SetToolTip(
            "Scan without reading and save one JPEG image or one multi-page "
            "PDF. ScanBox checks text orientation before saving."
        )
        self.document_camera_btn.SetToolTip(
            "Capture, straighten and recognise a page using a camera."
        )
        self.photo_scan_btn.SetToolTip(
            "Scan a photograph and describe it."
        )
        self.photo_camera_btn.SetToolTip(
            "Take a photograph with the selected camera and describe it."
        )
        self.camera_alignment_btn = wx.Button(
            self.scan_panel, label="FaceAlign"
        )
        # Keep the two document workflows adjacent in the native keyboard
        # navigation order. Cocoa can otherwise derive a different order
        # from the wrapped visual layout after controls are added or resized.
        self.scan_save_images_btn.MoveAfterInTabOrder(self.document_scan_btn)
        self.document_camera_btn.MoveAfterInTabOrder(self.scan_save_images_btn)
        scan_sizer.Add(self.document_scan_btn, 0, wx.ALL, 6)
        scan_sizer.Add(self.scan_save_images_btn, 0, wx.ALL, 6)
        scan_sizer.Add(self.document_camera_btn, 0, wx.ALL, 6)
        scan_sizer.Add(self.photo_scan_btn, 0, wx.ALL, 6)
        scan_sizer.Add(self.photo_camera_btn, 0, wx.ALL, 6)
        scan_sizer.Add(self.camera_alignment_btn, 0, wx.ALL, 6)
        self.scan_panel.SetSizer(scan_sizer)
        self.mode_tabs.AddPage(self.scan_panel, "Scan")

        # Import tab - bring in existing files. PDF2Word cannot detect tables
        # while OCR is running, so this document-specific group makes the
        # conversion behavior explicit.
        self.import_panel = wx.Panel(self.mode_tabs)
        _set_named_page_accessible(self.import_panel, "Import")
        import_sizer = wx.BoxSizer(wx.HORIZONTAL)
        self.pdf_ocr_modes = ("none", "missing_text", "all_pages")
        self.pdf_ocr_radio = wx.RadioBox(
            self.import_panel,
            label="PDF OCR",
            choices=[
                "No OCR.",
                "OCR only pages with no selectable text.",
                "OCR each page, including pages that contain selectable text.",
            ],
            majorDimension=1,
            style=wx.RA_SPECIFY_COLS,
        )
        self.pdf_ocr_radio.SetName("PDF OCR")
        current_ocr_mode = self.app_settings.get("pdf_ocr_mode", "none")
        self.pdf_ocr_radio.SetSelection(
            self.pdf_ocr_modes.index(current_ocr_mode)
            if current_ocr_mode in self.pdf_ocr_modes
            else 0
        )
        self.pdf_ocr_radio.SetToolTip(
            "Choose whether PDF OCR is disabled, used only on pages without "
            "selectable text, or used on every page. Table-aware Word "
            "conversion is available with No OCR."
        )
        self.pdf_ocr_radio.Bind(wx.EVT_RADIOBOX, self.on_pdf_ocr_mode_change)
        self.document_import_btn = wx.Button(self.import_panel, label="Import Document")
        self.document_import_btn.Bind(wx.EVT_BUTTON, self.on_document_import)
        self.photo_import_btn = wx.Button(self.import_panel, label="Import Photo")
        self.photo_batch_import_btn = wx.Button(
            self.import_panel, label="Batch Import Photos"
        )
        import_sizer.Add(
            self.pdf_ocr_radio, 0, wx.ALL | wx.ALIGN_CENTER_VERTICAL, 6
        )
        import_sizer.Add(self.document_import_btn, 0, wx.ALL, 6)
        import_sizer.Add(self.photo_import_btn, 0, wx.ALL, 6)
        import_sizer.Add(self.photo_batch_import_btn, 0, wx.ALL, 6)
        self.import_panel.SetSizer(import_sizer)
        self.mode_tabs.AddPage(self.import_panel, "Import")

        self.library_panel = wx.Panel(self.mode_tabs)
        _set_named_page_accessible(self.library_panel, "Photo Library")
        library_sizer = wx.BoxSizer(wx.VERTICAL)
        self.library_list = wx.ListCtrl(
            self.library_panel,
            style=wx.LC_REPORT | wx.LC_SINGLE_SEL,
        )
        self.library_list.EnableCheckBoxes(True)
        # Put the meaningful description first for screen readers. The actual
        # source filename is the final column, after availability status.
        self.library_list.InsertColumn(0, "Description", width=430)
        self.library_list.InsertColumn(1, "Status", width=100)
        self.library_list.InsertColumn(2, "Photo", width=230)
        self.library_description = wx.TextCtrl(
            self.library_panel,
            # Plain multiline read-only edit control, not TE_RICH2. Rich
            # Edit (MSFTEDIT) has a different accessibility layer that does
            # not reliably report restored focus to screen readers after
            # Alt+Tab. The plain control behaves consistently here.
            style=wx.TE_MULTILINE | wx.TE_READONLY | wx.TE_DONTWRAP,
        )
        self.library_description.SetName("Stored photo description")
        # Floor height so this stays usable with a screen reader even in a
        # restored (non-maximized) window - see also self.Maximize(True)
        # below, which makes maximized the default state on launch.
        self.library_description.SetMinSize((-1, 180))
        library_buttons = wx.BoxSizer(wx.HORIZONTAL)
        self.library_open_btn = wx.Button(self.library_panel, label="Open Photo")
        self.library_folder_btn = wx.Button(
            self.library_panel, label="Open Containing Folder"
        )
        self.library_copy_btn = wx.Button(self.library_panel, label="Copy Path")
        self.library_ask_btn = wx.Button(
            self.library_panel, label="Ask About Photo"
        )
        self.library_remove_btn = wx.Button(
            self.library_panel, label="Remove Stored Description"
        )
        self.library_select_all_btn = wx.Button(
            self.library_panel, label="Select All"
        )
        for button in (
            self.library_open_btn,
            self.library_folder_btn,
            self.library_copy_btn,
            self.library_ask_btn,
            self.library_remove_btn,
            self.library_select_all_btn,
        ):
            library_buttons.Add(button, 0, wx.ALL, 3)
        library_sizer.Add(self.library_list, 1, wx.ALL | wx.EXPAND, 6)
        library_sizer.Add(self.library_description, 1, wx.ALL | wx.EXPAND, 6)
        library_sizer.Add(library_buttons, 0, wx.ALL | wx.EXPAND, 3)
        self.library_panel.SetSizer(library_sizer)
        self.mode_tabs.AddPage(self.library_panel, "Photo Library")
        self.mode_tabs.SetSelection(self.TAB_SCAN)
        self.mode_tabs.Bind(wx.EVT_NOTEBOOK_PAGE_CHANGED, self.on_mode_change)

        self.output_box = wx.TextCtrl(
            panel,
            # See the matching accessibility note on
            # self.library_description above: use a plain edit control,
            # rather than TE_RICH2, for reliable restored-focus events.
            style=wx.TE_MULTILINE | wx.TE_READONLY | wx.TE_DONTWRAP,
        )
        self.output_box.SetName("Results")
        # Stay out of the tab order until there is something to read.
        self.output_box.Disable()
        # Floor height so results stay usable with a screen reader even in a
        # restored (non-maximized) window.
        self.output_box.SetMinSize((-1, 260))

        self.save_text_btn = wx.Button(panel, label="Save Text")
        self.save_img_btn = wx.Button(panel, label="Save Images")
        self.ask_about_photo_btn = wx.Button(panel, label="Ask About Photo")
        self.choose_another_photo_btn = wx.Button(
            panel, label="Choose Another Photo"
        )
        self.discard_btn = wx.Button(panel, label="Discard Session")
        self.cancel_batch_btn = wx.Button(panel, label="Cancel Batch Import")
        self.cancel_batch_btn.Hide()
        self.stop_camera_btn = wx.Button(panel, label="Stop Camera Capture")
        self.stop_camera_btn.Hide()
        self.settings_btn = wx.Button(panel, label="Settings")
        self.cancel_ai_download_btn = wx.Button(panel, label="Cancel AI Download")
        self.cancel_ai_download_btn.SetName("Cancel AI Download")
        self.cancel_ai_download_btn.Hide()
        self.install_gauge = wx.Gauge(
            panel,
            range=100,
            style=wx.GA_HORIZONTAL,
        )
        self.install_gauge.SetName("Local AI download progress")
        self.install_gauge.SetToolTip("Local AI model download progress")
        self.install_gauge.Hide()

        self.document_scan_btn.Bind(
            wx.EVT_BUTTON, lambda event: self.on_scan(event, False)
        )
        self.scan_save_images_btn.Bind(
            wx.EVT_BUTTON, self.on_scan_and_save_images
        )
        self.document_camera_btn.Bind(
            wx.EVT_BUTTON, lambda event: self.on_camera_capture(event, False)
        )
        self.photo_scan_btn.Bind(
            wx.EVT_BUTTON, lambda event: self.on_scan(event, True)
        )
        self.photo_camera_btn.Bind(
            wx.EVT_BUTTON, lambda event: self.on_camera_capture(event, True)
        )
        self.camera_alignment_btn.Bind(
            wx.EVT_BUTTON, self.start_camera_alignment
        )
        self.photo_import_btn.Bind(
            wx.EVT_BUTTON, lambda event: self.on_import_image(event, True)
        )
        self.photo_batch_import_btn.Bind(wx.EVT_BUTTON, self.batch_import_photos)
        self.library_list.Bind(wx.EVT_LIST_ITEM_SELECTED, self.on_library_selected)
        self.library_list.Bind(wx.EVT_LIST_ITEM_ACTIVATED, self.open_library_photo)
        self.library_list.Bind(wx.EVT_CONTEXT_MENU, self.on_library_context_menu)
        self.library_list.Bind(wx.EVT_CHAR_HOOK, self.on_library_key)
        self.library_open_btn.Bind(wx.EVT_BUTTON, self.open_library_photo)
        self.library_folder_btn.Bind(wx.EVT_BUTTON, self.open_library_folder)
        self.library_copy_btn.Bind(wx.EVT_BUTTON, self.copy_library_path)
        self.library_ask_btn.Bind(wx.EVT_BUTTON, self.ask_about_photo)
        self.library_remove_btn.Bind(wx.EVT_BUTTON, self.remove_library_entry)
        self.library_select_all_btn.Bind(wx.EVT_BUTTON, self.select_all_library_entries)
        self.save_text_btn.Bind(wx.EVT_BUTTON, self.save_text)
        self.save_img_btn.Bind(wx.EVT_BUTTON, self.save_images)
        self.ask_about_photo_btn.Bind(wx.EVT_BUTTON, self.ask_about_photo)
        self.choose_another_photo_btn.Bind(
            wx.EVT_BUTTON, self.choose_another_photo
        )
        self.discard_btn.Bind(wx.EVT_BUTTON, self.close_or_reset_session)
        self.cancel_batch_btn.Bind(wx.EVT_BUTTON, self.cancel_batch_import)
        self.stop_camera_btn.Bind(wx.EVT_BUTTON, self.stop_camera_workflow)
        self.settings_btn.Bind(wx.EVT_BUTTON, self.open_settings)
        self.cancel_ai_download_btn.Bind(wx.EVT_BUTTON, self.cancel_ai_download)
        self.Bind(wx.EVT_CLOSE, self.on_close)
        self.Bind(wx.EVT_ICONIZE, self.on_iconize)
        self.Bind(wx.EVT_ACTIVATE, self.on_activate)
        self.screen_hotkey_id = wx.NewIdRef()
        if sys.platform == "win32" and not self.RegisterHotKey(
            self.screen_hotkey_id, wx.MOD_CONTROL, 0xDC
        ):
            logger.warning("Could not register global Ctrl+\\ hotkey")
        self.Bind(wx.EVT_HOTKEY, self.on_screen_hotkey, id=self.screen_hotkey_id)
        self.screen_ocr_hotkey_id = wx.NewIdRef()
        if sys.platform == "win32" and not self.RegisterHotKey(
            self.screen_ocr_hotkey_id, wx.MOD_CONTROL | wx.MOD_SHIFT, 0xDC
        ):
            logger.warning("Could not register global Ctrl+Shift+\\ hotkey")
        self.Bind(
            wx.EVT_HOTKEY,
            self.on_screen_ocr_hotkey,
            id=self.screen_ocr_hotkey_id,
        )
        self.screen_question_hotkey_id = wx.NewIdRef()
        if sys.platform == "win32" and not self.RegisterHotKey(
            self.screen_question_hotkey_id,
            wx.MOD_CONTROL | wx.MOD_SHIFT,
            0xBF,
        ):
            logger.warning("Could not register global Ctrl+Shift+/ hotkey")
        self.Bind(
            wx.EVT_HOTKEY,
            self.on_screen_question_hotkey,
            id=self.screen_question_hotkey_id,
        )
        self.browser_pdf_hotkey_id = wx.NewIdRef()
        if sys.platform == "win32" and not self.RegisterHotKey(
            self.browser_pdf_hotkey_id,
            wx.MOD_CONTROL | wx.MOD_SHIFT,
            0xBD,  # VK_OEM_MINUS
        ):
            logger.warning("Could not register global Ctrl+Shift+hyphen hotkey")
        self.Bind(
            wx.EVT_HOTKEY,
            self.on_browser_pdf_hotkey,
            id=self.browser_pdf_hotkey_id,
        )
        self.toggle_window_hotkey_id = wx.NewIdRef()
        if sys.platform == "win32" and not self.RegisterHotKey(
            self.toggle_window_hotkey_id,
            wx.MOD_CONTROL | wx.MOD_ALT,
            0xDC,
        ):
            logger.warning("Could not register global Ctrl+Alt+\\ hotkey")
        self.Bind(
            wx.EVT_HOTKEY,
            self.toggle_window_visibility,
            id=self.toggle_window_hotkey_id,
        )
        if sys.platform == "darwin":
            pass
        elif sys.platform != "win32":
            logger.info("Global screen shortcuts are unavailable on this platform.")
        settings_id = wx.NewIdRef()
        self.Bind(wx.EVT_MENU, self.open_settings, id=settings_id)
        accelerators = [
            (
                wx.ACCEL_CMD if sys.platform == "darwin" else wx.ACCEL_CTRL,
                ord(","),
                settings_id,
            )
        ]
        if sys.platform == "darwin":
            accelerators.append((wx.ACCEL_CMD, ord("Q"), wx.ID_EXIT))
        self.SetAcceleratorTable(
            wx.AcceleratorTable(accelerators)
        )

        self.button_sizer = wx.WrapSizer(wx.HORIZONTAL)
        for btn in (
            self.save_text_btn,
            self.save_img_btn,
            self.ask_about_photo_btn,
            self.choose_another_photo_btn,
            self.discard_btn,
            self.cancel_batch_btn,
            self.stop_camera_btn,
            self.cancel_ai_download_btn,
            self.settings_btn,
        ):
            self.button_sizer.Add(btn, 0, wx.ALL, 3)

        self.mode_tabs_sizer_item = root.Add(
            self.mode_tabs, 0, wx.ALL | wx.EXPAND, 10
        )
        root.Add(self.output_box, 1, wx.ALL | wx.EXPAND, 10)
        root.Add(self.install_gauge, 0, wx.LEFT | wx.RIGHT | wx.EXPAND, 10)
        root.Add(self.button_sizer, 0, wx.ALL | wx.EXPAND, 10)

        # Page count lives in the status bar so it stays out of the tab order.
        self.CreateStatusBar()
        self.SetStatusText("No active session")

        panel.SetSizer(root)
        self.photo_library = []
        self.refresh_photo_library()
        self.update_controls()
        self.Centre()
        # Open maximized by default so the results boxes get their full
        # expandable height right away, rather than the cramped height a
        # restored 940x780 window leaves them - previously this meant
        # manually maximizing was needed before NVDA had a comfortable
        # amount of visible text to work with.
        self.Maximize(True)
        update_failure_log = os.environ.pop("SCANBOX_UPDATE_FAILED", "")
        if update_failure_log:
            wx.CallAfter(
                wx.MessageBox,
                "The update could not be completed. ScanBox attempted to "
                "restore the previous version.\n\nDetails: " + update_failure_log,
                "ScanBox update failed", wx.OK | wx.ICON_ERROR, self,
            )
        # On macOS, privacy prompts from helpers can appear behind other
        # windows if several background tasks start at once. Sequence the
        # permission prompts before the hotkey helper and model preload.
        if sys.platform == "darwin":
            wx.CallLater(800, self._start_macos_startup_sequence)
        else:
            threading.Thread(
                target=self._preload_local_ai,
                name="ScanBox AI preload",
                daemon=True,
            ).start()
        if self.app_settings.get("check_for_updates_on_startup", True):
            # Wait until the frame is visible before a newer-release prompt
            # can appear. A startup check stays silent when current or offline.
            wx.CallLater(1500, self._check_for_updates_on_startup)

    def _preload_local_ai(self):
        try:
            if external_ai_config(self.app_settings) is not None:
                logger.info("Built-in AI preload skipped while another local AI is selected.")
                return
            model_id = self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
            if (
                model_id == "qwen3_vl_2b"
                and _find_mtmd_model_files(model_id) is not None
                and _mtmd_runtime_needs_repair()
            ):
                logger.info(
                    "Installed Qwen model is missing its preferred runtime; "
                    "starting automatic repair."
                )
                wx.CallAfter(self.start_install, model_id)
                return
            if model_id == "qwen3_vl_2b" and vision_ready():
                total_memory = _total_physical_memory()
                if total_memory >= QWEN_PRELOAD_MINIMUM_RAM:
                    if start_mtmd_server(model_id):
                        logger.info("Qwen3-VL 2B preloaded and kept ready in background.")
                else:
                    logger.info(
                        "Qwen preload skipped: %.1f GB physical memory is below 8 GB.",
                        total_memory / (1024 ** 3),
                    )
            elif model_id == "florence_base" and florence_ready():
                logger.info("Florence-2 Base model preloaded in background.")
        except Exception:
            logger.exception("Background local AI preload failed")

    def _start_macos_startup_sequence(self):
        if sys.platform == "darwin" and not self.app_settings.get(
            "mac_permissions_prompted", False
        ):
            self.show_macos_permissions()
        threading.Thread(
            target=self._macos_startup_worker,
            name="ScanBox macOS startup permissions",
            daemon=True,
        ).start()

    def _macos_startup_worker(self):
        wx.CallAfter(self._start_macos_helper)
        time.sleep(12)
        self._preload_local_ai()

    def show_macos_permissions(self, event=None):
        if sys.platform != "darwin":
            return
        dialog = wx.Dialog(self, title="Mac Permissions")
        reset_performed = False
        root = wx.BoxSizer(wx.VERTICAL)
        intro = wx.StaticText(
            dialog,
            label=(
                "ScanBox needs Screen Recording for global screen description "
                "and OCR, and Accessibility for Mac global shortcuts. Use the "
                "buttons below, grant the permission in System Settings, then "
                "restart ScanBox."
            ),
        )
        intro.Wrap(560)
        root.Add(intro, 0, wx.ALL | wx.EXPAND, 12)

        rows = wx.FlexGridSizer(cols=2, hgap=8, vgap=8)
        rows.AddGrowableCol(1, 1)

        def add_row(label, button_label, handler):
            button = wx.Button(dialog, label=button_label)
            button.Bind(wx.EVT_BUTTON, handler)
            text = wx.StaticText(dialog, label=label)
            text.Wrap(390)
            rows.Add(button, 0, wx.EXPAND)
            rows.Add(text, 0, wx.EXPAND | wx.ALIGN_CENTER_VERTICAL)

        add_row(
            "Required for Control Backslash and Control Shift Backslash.",
            "Open Screen Recording",
            lambda event: open_macos_privacy_pane("Privacy_ScreenCapture"),
        )
        add_row(
            "Required for global shortcuts and locating the active window.",
            "Open Accessibility",
            lambda event: open_macos_privacy_pane("Privacy_Accessibility"),
        )

        def reset_permissions(event):
            nonlocal reset_performed
            if wx.MessageBox(
                "Reset ScanBox permissions so macOS will ask again? "
                "You must quit and reopen ScanBox afterwards.",
                "Reset Permissions",
                wx.YES_NO | wx.ICON_WARNING,
                dialog,
            ) != wx.YES:
                return
            try:
                reset_macos_scanbox_permissions()
                reset_performed = True
                self.app_settings["mac_permissions_prompted"] = False
                write_app_settings(self.app_settings)
                wx.MessageBox(
                    "ScanBox permissions were reset. Quit and reopen ScanBox, "
                    "then grant the permissions again.",
                    "Permissions Reset",
                    wx.OK | wx.ICON_INFORMATION,
                    dialog,
                )
            except Exception as exc:
                wx.MessageBox(
                    f"ScanBox could not reset permissions.\n\n{exc}",
                    "Reset Failed",
                    wx.OK | wx.ICON_ERROR,
                    dialog,
                )

        add_row(
            "Use this if macOS says permission is on but capture still fails.",
            "Reset ScanBox Permissions",
            reset_permissions,
        )

        root.Add(rows, 0, wx.ALL | wx.EXPAND, 12)
        dont_show_again = wx.CheckBox(
            dialog,
            label="Don't show this permissions guide again at startup",
        )
        dont_show_again.SetValue(True)
        root.Add(dont_show_again, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 12)
        buttons = dialog.CreateButtonSizer(wx.OK)
        root.Add(buttons, 0, wx.ALL | wx.ALIGN_RIGHT, 12)
        dialog.SetSizerAndFit(root)
        dialog.SetMinSize((680, -1))
        dialog.CentreOnParent()
        suppress_startup_guide = True
        try:
            dialog.ShowModal()
            suppress_startup_guide = dont_show_again.GetValue()
        finally:
            dialog.Destroy()
        # Reset deliberately clears the startup choice so the guide appears
        # again after relaunch. Do not immediately overwrite that choice when
        # this dialog closes.
        if not reset_performed:
            self.app_settings["mac_permissions_prompted"] = (
                suppress_startup_guide
            )
            write_app_settings(self.app_settings)

    def open_user_manual(self, event=None):
        """Open the bundled user manual in the default web browser."""
        if not os.path.isfile(MANUAL_PATH):
            wx.MessageBox(
                f"The ScanBox user manual could not be found:\n{MANUAL_PATH}",
                "User Manual",
                wx.OK | wx.ICON_ERROR,
                self,
            )
            return
        try:
            manual_url = Path(MANUAL_PATH).resolve().as_uri()
            if not wx.LaunchDefaultBrowser(manual_url):
                raise OSError("The default browser did not open the manual.")
        except Exception as exc:
            logger.exception("Could not open the user manual")
            wx.MessageBox(
                f"ScanBox could not open the user manual.\n\n{exc}",
                "User Manual",
                wx.OK | wx.ICON_ERROR,
                self,
            )

    def open_donation_page(self, event=None):
        if not wx.LaunchDefaultBrowser(DONATE_URL):
            wx.MessageBox(
                "ScanBox could not open the donation page.",
                "Donate to Project",
                wx.OK | wx.ICON_ERROR,
                self,
            )

    def open_license(self, event=None):
        if not os.path.isfile(LICENSE_PATH):
            wx.MessageBox(
                "The ScanBox license file could not be found.",
                "License",
                wx.OK | wx.ICON_ERROR,
                self,
            )
            return
        try:
            open_with_default_application(LICENSE_PATH)
        except Exception as exc:
            wx.MessageBox(
                f"ScanBox could not open the license.\n\n{exc}",
                "License",
                wx.OK | wx.ICON_ERROR,
                self,
            )

    def show_about(self, event=None):
        """Show application identity and version information."""
        wx.MessageBox(
            f"{APP_NAME}\nVersion {APP_VERSION}\n"
            "Copyright © 2026 Sam Taylor\n"
            "Licensed under the GNU Affero General Public License, "
            "version 3 or later.\n\n"
            "A privacy-first scanning, recognition, and photo-description "
            "program for Windows and macOS.\n\n"
            "Scanning, recognition, and image descriptions are processed "
            "on this computer.",
            f"About {APP_NAME}",
            wx.OK | wx.ICON_INFORMATION,
            self,
        )

    @staticmethod
    def _version_parts(version):
        """Return a tuple suitable for comparing calendar-style versions."""
        match = re.fullmatch(r"\s*v?(\d+(?:\.\d+)*)\s*", str(version))
        if not match:
            raise ValueError(f"Invalid version number: {version}")
        return tuple(int(part) for part in match.group(1).split("."))

    def check_for_updates(self, event=None):
        """Fetch release information without blocking the application window."""
        self._start_update_check(show_current=True)

    def _check_for_updates_on_startup(self):
        self._start_update_check(show_current=False)

    def _start_update_check(self, show_current):
        manifest_url = (
            self.app_settings.get("update_manifest_url", "").strip()
            or UPDATE_MANIFEST_URL
        )
        if not manifest_url:
            if show_current:
                wx.MessageBox(
                    f"{APP_NAME} {APP_VERSION} is installed.\n\n"
                    "Automatic update information is not available for this build.",
                    "Check for Updates",
                    wx.OK | wx.ICON_INFORMATION,
                    self,
                )
            return

        self.check_updates_item.Enable(False)
        if show_current:
            self.SetStatusText("Checking for ScanBox updates...")
        threading.Thread(
            target=self._check_for_updates_worker,
            args=(manifest_url, show_current),
            name="ScanBox update check",
            daemon=True,
        ).start()

    def _check_for_updates_worker(self, manifest_url, show_current):
        try:
            request = urllib.request.Request(
                manifest_url,
                headers={"User-Agent": f"{APP_NAME}/{APP_VERSION}"},
            )
            with urllib.request.urlopen(request, timeout=15) as response:
                manifest = json.loads(response.read().decode("utf-8"))
            # Accept both ScanBox's small update-manifest format and GitHub's
            # latest-release response.
            latest_version = str(
                manifest.get("version") or manifest["tag_name"]
            ).strip()
            self._version_parts(latest_version)
            download_url = _release_asset_url(manifest)
            notes = str(
                manifest.get("notes") or manifest.get("body", "")
            ).strip()
            result = (latest_version, download_url, notes, None)
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                logger.info("No published ScanBox release was found")
                result = (APP_VERSION, "", "", None)
            else:
                logger.exception("Update check failed")
                result = (None, "", "", str(exc))
        except Exception as exc:
            logger.exception("Update check failed")
            result = (None, "", "", str(exc))
        wx.CallAfter(self._finish_update_check, *result, show_current)

    def _finish_update_check(
        self, latest_version, download_url, notes, error, show_current
    ):
        self.check_updates_item.Enable(True)
        if show_current:
            self.SetStatusText("Ready")
        if error:
            if show_current:
                wx.MessageBox(
                    "ScanBox could not check for updates.\n\n"
                    "Check your internet connection and try again.\n\n"
                    f"Details: {error}",
                    "Check for Updates",
                    wx.OK | wx.ICON_ERROR,
                    self,
                )
            return

        if self._version_parts(latest_version) <= self._version_parts(APP_VERSION):
            if show_current:
                wx.MessageBox(
                    f"{APP_NAME} {APP_VERSION} is up to date.",
                    "Check for Updates",
                    wx.OK | wx.ICON_INFORMATION,
                    self,
                )
            return

        message = (
            f"ScanBox {latest_version} is available.\n"
            f"You currently have version {APP_VERSION}."
        )
        if notes:
            message += f"\n\n{notes}"
        style = wx.YES_NO | wx.ICON_INFORMATION if download_url else wx.OK | wx.ICON_INFORMATION
        if download_url:
            if getattr(sys, "frozen", False) and download_url.lower().endswith(".zip"):
                message += "\n\nWould you like ScanBox to download and install it now?"
            else:
                message += "\n\nWould you like to download it now?"
        answer = wx.MessageBox(message, "ScanBox Update Available", style, self)
        if download_url and answer == wx.YES:
            if getattr(sys, "frozen", False) and download_url.lower().endswith(".zip"):
                self._start_update_download(latest_version, download_url)
            else:
                wx.LaunchDefaultBrowser(download_url)

    def _start_update_download(self, latest_version, download_url):
        """Download and stage a portable update without blocking the UI."""
        self.check_updates_item.Enable(False)
        self.SetStatusText(f"Downloading ScanBox {latest_version}...")
        threading.Thread(
            target=self._update_download_worker,
            args=(latest_version, download_url),
            name="ScanBox application update",
            daemon=True,
        ).start()

    def _update_download_worker(self, latest_version, download_url):
        archive_path = os.path.join(
            tempfile.gettempdir(),
            f"ScanBox-{latest_version.strip().lstrip('v')}-{uuid.uuid4().hex}.zip",
        )
        try:
            _download_file(
                download_url,
                archive_path,
                "ScanBox update",
                lambda message: wx.CallAfter(self.SetStatusText, message),
            )
            payload_path = _prepare_update_payload(archive_path)
            _launch_portable_updater(payload_path)
            result = (payload_path, None)
        except Exception as exc:
            logger.exception("Could not download or prepare ScanBox update")
            result = (None, str(exc))
        wx.CallAfter(self._finish_update_download, *result)

    def _finish_update_download(self, payload_path, error):
        self.check_updates_item.Enable(True)
        if error:
            self.SetStatusText("ScanBox update failed.")
            wx.MessageBox(
                f"ScanBox could not install the update.\n\n{error}",
                "Update Failed",
                wx.OK | wx.ICON_ERROR,
                self,
            )
            return
        self.SetStatusText("Update ready. Restarting ScanBox...")
        self._close_after_announcement = True
        self.Close()

    def on_mode_change(self, event):
        selected_tab = self.mode_tabs.GetSelection()
        if selected_tab == self.TAB_LIBRARY:
            self.refresh_photo_library()
        self.update_controls()
        event.Skip()

    def focus_scan_tab_on_startup(self):
        """Give macOS a stable, screen-reader-friendly entry point."""
        self.mode_tabs.SetSelection(self.TAB_SCAN)
        self.mode_tabs.SetFocus()

    def on_activate(self, event):
        # During a backslash screen workflow the previously focused action
        # control is disabled. If the user Alt+Tabs back before completion,
        # explicitly land in the enabled Results box instead of leaving the
        # frame with no keyboard focus and no working Tab traversal.
        if event.GetActive() and self.screen_capture_busy:
            wx.CallAfter(self.output_box.SetFocusFromKbd)
        event.Skip()

    def _ensure_tray_icon(self):
        if self.tray_icon is None:
            self.tray_icon = ScanBoxTaskBarIcon(self)

    def on_iconize(self, event):
        # An ordinary minimise remains a normal taskbar/Dock window and must
        # stay available to Alt+Tab. Complete hiding belongs exclusively to
        # Control+Alt+Backslash in toggle_window_visibility().
        event.Skip()

    def restore_from_notification_area(self, after_restore=None):
        self._complete_restore(after_restore)

    def _complete_restore(self, after_restore=None):
        self.Show(True)
        if self.IsIconized():
            self.Iconize(False)
        self.Raise()
        try:
            ctypes.windll.user32.SetForegroundWindow(self.GetHandle())
        except Exception:
            pass
        focus = self._restore_focus
        if focus is not None:
            try:
                focus.SetFocus()
            except Exception:
                self.mode_tabs.SetFocus()
        else:
            self.mode_tabs.SetFocus()
        announce("restored.")
        if after_restore is not None:
            wx.CallAfter(after_restore)

    def toggle_window_visibility(self, event=None):
        is_foreground = self.IsActive()
        if sys.platform == "win32":
            try:
                is_foreground = (
                    ctypes.windll.user32.GetForegroundWindow() == self.GetHandle()
                )
            except Exception:
                pass
        if self.IsShown() and not self.IsIconized() and is_foreground:
            focus = wx.Window.FindFocus()
            if focus is not None:
                self._restore_focus = focus
            if sys.platform == "win32":
                self._ensure_tray_icon()
                # Minimising first makes Windows transfer keyboard and
                # accessibility focus to another application. Hiding a frame
                # directly can leave a screen reader focused on its now
                # invisible controls. The deferred hide then gives this
                # explicit shortcut its intended removal from Alt+Tab.
                self.Iconize(True)
                wx.CallAfter(self.Hide)
            else:
                self.Hide()
            logger.info("ScanBox hidden by the global window shortcut")
            announce("minimised.")
        else:
            logger.info("ScanBox restored by the global window shortcut")
            self.restore_from_notification_area()

    def _begin_screen_processing(self):
        """Keep a focusable placeholder available during screen processing."""
        self.busy = True
        self.screen_capture_busy = True
        if not self.app_settings.get("append_text_to_buffer", False):
            self.prepare_output_buffer()
        self.output_box.Enable()
        start = self.output_box.GetLastPosition()
        separator = "\n\n" if start else ""
        placeholder = separator + "Processing..."
        self.output_box.AppendText(placeholder)
        self.output_box.SetInsertionPoint(start + len(separator))
        self.output_box.ShowPosition(start + len(separator))
        self._screen_result_placeholder = (start, start + len(placeholder), separator)
        self.update_controls()
        self.SetStatusText("Processing...")

    def on_screen_hotkey(self, event=None):
        if self.busy:
            return
        if self.photo_mode_blocked(True):
            return
        self.last_active_mode = "photo"
        try:
            image = grab_foreground_window()
            path = os.path.join(TEMP_DIR, f"screen_{uuid.uuid4().hex}.png")
            image.save(path, "PNG")
            play_shutter_sound()
        except Exception as exc:
            logger.exception("Could not capture the screen")
            wx.MessageBox(f"Could not capture the screen: {exc}", "Screen capture failed")
            return
        self._begin_screen_processing()
        threading.Thread(
            target=self._screen_description_worker,
            args=(path,),
            name="ScanBox screen description",
            daemon=True,
        ).start()

    def on_screen_ocr_hotkey(self, event=None):
        if self.busy:
            return
        self.last_active_mode = "document"
        try:
            image = grab_foreground_window()
            path = os.path.join(TEMP_DIR, f"screen_ocr_{uuid.uuid4().hex}.png")
            image.save(path, "PNG")
            play_shutter_sound()
        except Exception as exc:
            logger.exception("Could not capture the screen for OCR")
            wx.MessageBox(f"Could not capture the screen: {exc}", "Screen capture failed")
            return
        self._begin_screen_processing()
        threading.Thread(
            target=self._screen_ocr_worker,
            args=(path,),
            name="ScanBox screen OCR",
            daemon=True,
        ).start()

    def _restore_after_screen_question(self, restore_window=None, restore_app=""):
        if sys.platform == "win32" and restore_window:
            try:
                ctypes.windll.user32.SetForegroundWindow(restore_window)
            except Exception:
                pass
        elif sys.platform == "darwin" and restore_app:
            try:
                subprocess.Popen(
                    ["open", "-b", restore_app],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
            except Exception:
                pass

    def _ask_screen_question(self, path, restore_window=None, restore_app=""):
        model_id = self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
        external_available = external_ai_config(self.app_settings) is not None
        qwen_installed = (
            _find_mtmd_model_files("qwen3_vl_2b") is not None
            and _find_mtmd_runner() is not None
        )
        logger.info(
            "Ask shortcut captured window; selected_model=%s qwen_installed=%s",
            model_id,
            qwen_installed,
        )
        if not external_available and model_id != "qwen3_vl_2b" and qwen_installed:
            if wx.MessageBox(
                "Qwen3-VL 2B is required to answer questions. Switch to Qwen now?",
                "Ask ScanBox a question",
                wx.YES_NO | wx.YES_DEFAULT | wx.ICON_QUESTION,
                None,
            ) != wx.YES:
                _remove_quietly(path)
                self._restore_after_screen_question(restore_window, restore_app)
                return
            self.app_settings["vision_model"] = "qwen3_vl_2b"
            write_app_settings(self.app_settings)
            threading.Thread(
                target=start_mtmd_server,
                args=("qwen3_vl_2b",),
                name="ScanBox AI preload",
                daemon=True,
            ).start()
            logger.info("Ask ScanBox switched the selected model to Qwen3-VL 2B")
        elif not external_available and not qwen_installed:
            _remove_quietly(path)
            install = wx.MessageBox(
                "Qwen3-VL 2B is required to answer questions and is not installed. "
                "Install it now?",
                "Ask ScanBox a question",
                wx.YES_NO | wx.YES_DEFAULT | wx.ICON_QUESTION,
                None,
            )
            if install == wx.YES:
                self.app_settings["vision_model"] = "qwen3_vl_2b"
                write_app_settings(self.app_settings)
                self.start_install("qwen3_vl_2b")
            self._restore_after_screen_question(restore_window, restore_app)
            return
        # This is intentionally not owned by the main ScanBox frame. The
        # question briefly needs focus, but the main window must remain in the
        # background like the other global workflows.
        dialog = wx.Dialog(
            None,
            title="Ask ScanBox a question",
            style=wx.DEFAULT_DIALOG_STYLE | wx.STAY_ON_TOP,
        )
        root = wx.BoxSizer(wx.VERTICAL)
        label = wx.StaticText(dialog, label="Question")
        question = wx.TextCtrl(dialog, style=wx.TE_PROCESS_ENTER)
        question.SetName("Question")
        buttons = dialog.CreateButtonSizer(wx.OK | wx.CANCEL)
        root.Add(label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 12)
        root.Add(question, 0, wx.ALL | wx.EXPAND, 12)
        root.Add(buttons, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM | wx.ALIGN_RIGHT, 12)
        dialog.SetSizerAndFit(root)
        dialog.SetMinSize((480, -1))
        dialog.CentreOnScreen()
        question.Bind(
            wx.EVT_TEXT_ENTER,
            lambda event: dialog.EndModal(wx.ID_OK) if question.GetValue().strip() else None,
        )
        def focus_question_dialog():
            dialog.Raise()
            if sys.platform == "win32":
                try:
                    ctypes.windll.user32.SetForegroundWindow(dialog.GetHandle())
                except Exception:
                    pass
            question.SetFocus()

        wx.CallAfter(focus_question_dialog)
        logger.info("Showing Ask ScanBox question dialog")
        result = dialog.ShowModal()
        query = question.GetValue().strip()
        dialog.Destroy()
        self._restore_after_screen_question(restore_window, restore_app)
        if result != wx.ID_OK or not query:
            logger.info("Ask ScanBox question dialog cancelled")
            _remove_quietly(path)
            return
        logger.info("Ask ScanBox question submitted")
        play_shutter_sound()
        self.last_active_mode = "photo"
        self.busy = True
        self.update_controls()
        threading.Thread(
            target=self._screen_question_worker,
            args=(path, query),
            name="ScanBox image question",
            daemon=True,
        ).start()

    def on_screen_question_hotkey(self, event=None):
        if self.busy:
            return
        restore_window = None
        if sys.platform == "win32":
            try:
                restore_window = ctypes.windll.user32.GetForegroundWindow()
            except Exception:
                pass
        try:
            image = grab_foreground_window()
            path = os.path.join(TEMP_DIR, f"screen_question_{uuid.uuid4().hex}.png")
            image.save(path, "PNG")
        except Exception as exc:
            logger.exception("Could not capture the screen for a question")
            wx.MessageBox(f"Could not capture the screen: {exc}", "Screen capture failed")
            return
        self._ask_screen_question(path, restore_window=restore_window)

    def on_browser_pdf_hotkey(self, event=None):
        """Import a public PDF exposed by the foreground browser tab."""
        if self.busy:
            return
        self.busy = True
        self.update_controls()
        self.SetStatusText("Reading the current browser tab...")
        announce("Please wait. Reading the current browser tab.")
        self.browser_url_executor.submit(self._browser_pdf_url_worker)

    @staticmethod
    def _warm_browser_url_reader():
        started = time.perf_counter()
        try:
            warm_up_active_tab_url_reader()
            logger.info(
                "Browser accessibility initialized in %.2f seconds",
                time.perf_counter() - started,
            )
        except Exception:
            logger.exception("Could not initialize browser accessibility")

    def _browser_pdf_url_worker(self):
        """Read browser UIA/AX state away from wx's existing COM apartment."""
        started = time.perf_counter()
        try:
            url = get_active_tab_url()
            if not url.lower().startswith("https://"):
                raise ActiveTabUrlError(
                    "Only HTTPS PDF addresses can be imported from a browser."
                )
        except Exception as exc:
            logger.exception("Could not read the active browser URL")
            wx.CallAfter(self._finish_browser_pdf_error, str(exc))
            return
        logger.info(
            "Read active browser URL in %.2f seconds",
            time.perf_counter() - started,
        )
        wx.CallAfter(
            self.SetStatusText,
            "Downloading PDF from the current browser tab...",
        )
        wx.CallAfter(
            announce,
            "Downloading PDF from the current browser tab. Please wait.",
        )
        self._download_browser_pdf_worker(url)

    def _download_browser_pdf_worker(self, url):
        path = ""
        started = time.perf_counter()

        def save_through_active_browser():
            nonlocal path
            path = os.path.join(
                TEMP_DIR, f"Browser PDF {uuid.uuid4().hex}.pdf"
            )
            logger.info("Saving session-bound PDF through the active browser")
            try:
                download_active_browser_document_windows(path, timeout=30)
            except Exception:
                logger.exception(
                    "Silent browser download was unavailable; using Save As"
                )
                save_active_browser_document_windows(path, timeout=30)
            with open(path, "rb") as saved_pdf:
                first = saved_pdf.read(8)
                saved_pdf.seek(0, os.SEEK_END)
                total = saved_pdf.tell()
            if total > 100 * 1024 * 1024:
                raise ValueError("The PDF is larger than ScanBox's 100 MB limit.")
            if not first.startswith(b"%PDF-"):
                raise ValueError("The browser did not save a valid PDF file.")
            logger.info("Browser saved PDF bytes=%d", total)
            wx.CallAfter(self._open_downloaded_browser_pdf, path)

        try:
            if sys.platform == "win32" and is_session_bound_download_url(url):
                save_through_active_browser()
                return

            request = make_browser_download_request(url)
            try:
                response = urllib.request.urlopen(request, timeout=30)
            except urllib.error.HTTPError as exc:
                failed_url = exc.geturl()
                session_bound = (
                    is_session_bound_download_url(url)
                    or is_session_bound_download_url(failed_url)
                )
                if exc.code not in {401, 403, 404, 429}:
                    raise
                if sys.platform == "win32":
                    logger.info(
                        "Browser PDF request failed with HTTP %d; "
                        "using the active browser",
                        exc.code,
                    )
                    save_through_active_browser()
                    return
                if not session_bound:
                    raise
                logger.info("Retrying browser-session PDF with browser cookies")
                response = open_with_browser_session(url, timeout=30)
            with response:
                final_url = response.geturl()
                if not final_url.lower().startswith("https://"):
                    raise ValueError("The PDF redirected to a non-HTTPS address.")
                supplied_name = response.headers.get_filename() or unquote(
                    os.path.basename(urlsplit(final_url).path)
                )
                supplied_stem = os.path.splitext(supplied_name)[0]
                safe_stem = "".join(
                    character if character not in '<>:"/\\|?*' else "_"
                    for character in supplied_stem.strip()
                ).rstrip(" .")[:120] or "Browser PDF"
                path = os.path.join(TEMP_DIR, safe_stem + ".pdf")
                if os.path.exists(path):
                    for number in range(2, 1000):
                        candidate = os.path.join(
                            TEMP_DIR, f"{safe_stem} ({number}).pdf"
                        )
                        if not os.path.exists(candidate):
                            path = candidate
                            break
                    else:
                        path = os.path.join(
                            TEMP_DIR, f"{safe_stem} {uuid.uuid4().hex}.pdf"
                        )
                content_type = response.headers.get_content_type().lower()
                try:
                    expected_size = int(response.headers.get("Content-Length") or 0)
                except (TypeError, ValueError):
                    expected_size = 0
                next_progress_announcement = 25
                download_started = time.monotonic()
                total = 0
                first = b""
                with open(path, "wb") as output:
                    while True:
                        chunk = response.read(1024 * 1024)
                        if not chunk:
                            break
                        if not first:
                            first = chunk[:8]
                        total += len(chunk)
                        if total > 100 * 1024 * 1024:
                            raise ValueError("The PDF is larger than ScanBox's 100 MB limit.")
                        output.write(chunk)
                        if expected_size > 0:
                            percentage = min(99, int(total * 100 / expected_size))
                            if percentage >= next_progress_announcement:
                                wx.CallAfter(
                                    self._browser_pdf_download_progress,
                                    next_progress_announcement,
                                    time.monotonic() - download_started >= 2.5,
                                )
                                next_progress_announcement += 25
                if not first.startswith(b"%PDF-"):
                    raise ValueError(
                        f"The current tab did not return a PDF (content type: {content_type})."
                    )
            logger.info(
                "Downloaded browser PDF bytes=%d elapsed=%.2fs",
                total,
                time.perf_counter() - started,
            )
            wx.CallAfter(self._open_downloaded_browser_pdf, path)
        except Exception as exc:
            if path:
                _remove_quietly(path)
            logger.exception("Current browser PDF download failed")
            wx.CallAfter(self._finish_browser_pdf_error, str(exc))

    def _open_downloaded_browser_pdf(self, path):
        self.busy = False
        self.update_controls()
        try:
            opening_word = bool(
                sys.platform == "win32"
                and self.app_settings.get("open_word_after_pdf_conversion", False)
                and microsoft_word_available()
            )
            message = (
                "PDF downloaded. Converting and opening it in Microsoft Word. Please wait."
                if opening_word
                else "PDF downloaded. Reading it in ScanBox. Please wait."
            )
            self.SetStatusText(message)
            announce(message)
            self.process_pdf(path)
        except Exception as exc:
            _remove_quietly(path)
            self._finish_browser_pdf_error(str(exc))

    def _finish_browser_pdf_error(self, message):
        self.busy = False
        self.update_controls()
        self.SetStatusText("Current browser PDF could not be imported.")
        wx.MessageBox(message, "Current browser PDF unavailable", wx.OK | wx.ICON_ERROR)

    def _browser_pdf_download_progress(self, percentage, speak):
        message = f"Downloading current browser PDF: {percentage} percent."
        self.SetStatusText(message)
        if speak:
            announce(message)

    def _start_macos_helper(self):
        """Start the native hotkey and ScreenCaptureKit companion process."""
        global _mac_announce_process
        if not os.path.isfile(MACOS_CAPTURE_HELPER):
            logger.warning(
                "macOS helper is missing; global shortcuts are unavailable: %s",
                MACOS_CAPTURE_HELPER,
            )
            return
        # A helper from a previous, force-quit ScanBox session (rather than
        # a normal Cmd+Q, which terminates it in on_close) can be left
        # running. Carbon's global hotkeys are registered per key
        # combination system-wide, so a stray helper still holding those
        # registrations makes this session's own registration attempts
        # silently lose the combo to the orphaned process - the shortcuts
        # then work or not depending on which helper happens to still be
        # alive, which looks exactly like intermittent failure. Clearing
        # out any stray copies first guarantees only this session's helper
        # ever holds the registrations.
        try:
            subprocess.run(
                ["pkill", "-f", MACOS_CAPTURE_HELPER],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            time.sleep(0.3)
        except Exception:
            logger.exception("Could not clear stray macOS helper processes")
        try:
            self._mac_helper_process = subprocess.Popen(
                [MACOS_CAPTURE_HELPER, "listen", TEMP_DIR],
                stdin=subprocess.PIPE,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                bufsize=1,
            )
            _mac_announce_process = self._mac_helper_process
            threading.Thread(
                target=self._read_macos_helper_events,
                name="ScanBox macOS hotkey events",
                daemon=True,
            ).start()
            logger.info("Started native macOS global-shortcut helper.")
        except Exception:
            self._mac_helper_process = None
            _mac_announce_process = None
            logger.exception("Could not start the macOS global-shortcut helper")

    def _request_macos_screen_recording_permission(self):
        try:
            result = subprocess.run(
                [MACOS_CAPTURE_HELPER, "request-screen-recording"],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=30,
            )
            status = result.stdout.strip()
            logger.info(
                "macOS Screen Recording permission request returned: %s stderr=%r",
                status or f"exit {result.returncode}",
                result.stderr,
            )
            if status != "granted":
                wx.CallAfter(
                    self.SetStatusText,
                    "Screen Recording permission is needed for screen shortcuts.",
                )
        except Exception:
            logger.exception("Could not request macOS Screen Recording permission")

    def _read_macos_helper_events(self):
        process = self._mac_helper_process
        if process is None or process.stdout is None:
            return
        try:
            for raw_line in process.stdout:
                line = raw_line.rstrip("\r\n")
                if not line:
                    continue
                kind, _separator, value = line.partition("\t")
                if kind in {
                    "describe", "ocr", "ask", "toggle", "error", "fatal",
                    "browser-pdf", "accessibility", "ready", "hotkey",
                }:
                    wx.CallAfter(self._handle_macos_helper_event, kind, value)
                else:
                    logger.warning("macOS helper: %s", line)
        except Exception:
            logger.exception("macOS helper event reader failed")
        finally:
            return_code = process.poll()
            logger.warning("macOS helper event reader stopped; returncode=%s", return_code)

    def _handle_macos_helper_event(self, kind, value=""):
        if kind == "fatal":
            wx.MessageBox(
                "ScanBox could not register its macOS global shortcuts.\n\n"
                f"{value}\n\n"
                "Another application may already use one of these shortcuts.",
                "Global shortcuts unavailable",
                wx.OK | wx.ICON_ERROR,
                self,
            )
            return
        if kind == "ready":
            logger.info("macOS global-shortcut helper ready: %s", value)
            return
        if kind == "hotkey":
            logger.info("macOS global shortcut pressed: %s", value)
            return
        if kind == "toggle":
            self.toggle_window_visibility()
            return
        if kind == "browser-pdf":
            self.on_browser_pdf_hotkey()
            return
        if kind == "accessibility":
            logger.warning("macOS Accessibility reminder: %s", value)
            return
        if kind == "error":
            wx.MessageBox(
                "ScanBox could not capture the active window.\n\n"
                f"{value}\n\n"
                "Allow Screen Recording for ScanBox in System Settings > "
                "Privacy & Security > Screen Recording, then try again.",
                "Screen capture failed",
                wx.OK | wx.ICON_ERROR,
                self,
            )
            return

        restore_app = ""
        path = value
        if value.startswith("window\t"):
            fields = value.split("\t")
            if kind == "ask" and len(fields) >= 4:
                _tag, window_id, path, restore_app = fields[:4]
            elif len(fields) >= 3:
                _tag, window_id, path = fields[:3]
            else:
                logger.warning("Malformed macOS window capture event: %r", value)
                return
            try:
                macos_capture_window(window_id, path)
            except Exception as exc:
                logger.exception("macOS window capture failed")
                wx.MessageBox(
                    "ScanBox could not capture the active window.\n\n"
                    f"{exc}\n\n"
                    "Allow Screen Recording for ScanBox in System Settings > "
                    "Privacy & Security > Screen Recording, then try again.",
                    "Screen capture failed",
                    wx.OK | wx.ICON_ERROR,
                    self,
                )
                return
        elif kind == "ask" and "\t" in value:
            path, restore_app = value.split("\t", 1)
        if kind not in {"describe", "ocr", "ask"} or not os.path.isfile(path):
            return
        if self.busy:
            try:
                os.remove(path)
            except OSError:
                pass
            return
        if kind == "ask":
            self._ask_screen_question(path, restore_app=restore_app)
            return
        read_text = kind == "ocr"
        if not read_text and self.photo_mode_blocked(True):
            try:
                os.remove(path)
            except OSError:
                pass
            return

        play_shutter_sound()
        self.last_active_mode = "document" if read_text else "photo"
        self._begin_screen_processing()
        self.Show(True)
        if self.IsIconized():
            self.Iconize(False)
        self.Raise()
        worker = (
            self._screen_ocr_worker if read_text else self._screen_description_worker
        )
        threading.Thread(
            target=worker,
            args=(path,),
            name=(
                "ScanBox macOS active-window OCR"
                if read_text
                else "ScanBox macOS active-window description"
            ),
            daemon=True,
        ).start()

    def _screen_ocr_worker(self, path):
        try:
            detect_and_crop_screen_document(path)
            text = self._read_screen_text(
                path, use_vision_fallback=True, cancel_event=self.shutdown_event
            )
        except Exception as exc:
            logger.exception("Screen OCR failed")
            text = f"Processing failed: {exc}"
        wx.CallAfter(self._finish_screen_description, path, text)

    def _read_screen_text(self, path, use_vision_fallback=False, cancel_event=None):
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        external_attempted = (
            use_vision_fallback
            and external_ai_config(self.app_settings) is not None
        )
        if external_attempted:
            text = run_external_ai_task(
                "transcribe", path, cancel_event=cancel_event,
                settings=self.app_settings,
            )
            if text == PHOTO_DESCRIPTION_CANCELLED:
                return text
            if not text.startswith("Local AI service") and not looks_like_ocr_repetition_garbage(text):
                return clean_screen_ocr_text(text)
            logger.warning("External screen OCR failed; falling back to native OCR: %s", text)
        if sys.platform == "darwin":
            text = clean_screen_ocr_text(macos_ocr(path))
            native_ocr_name = "Apple Vision"
        else:
            text = clean_screen_ocr_text(windows_ocr(path))
            native_ocr_name = "Windows.Media.Ocr"
        if len(re.findall(r"[A-Za-z0-9]+", text)) >= 5:
            logger.info("Screen OCR used %s", native_ocr_name)
            return text
        if use_vision_fallback and not external_attempted and vision_ready("transcribe"):
            logger.info("Screen OCR used Florence-2 fallback")
            return clean_screen_ocr_text(
                run_vision_task("transcribe", path, cancel_event)
            )
        return ""

    def _screen_description_worker(self, path):
        try:
            text = self.process_photo(
                path, include_ocr=False, cancel_event=self.shutdown_event
            )
            detect_and_crop_screen_document(path)
            visible_text = self._read_screen_text(path)
            if visible_text:
                text += "\n\nRecognised text:\n" + visible_text
        except Exception as exc:
            logger.exception("Screen description failed")
            text = f"Processing failed: {exc}"
        wx.CallAfter(self._finish_screen_description, path, text)

    def _screen_question_worker(self, path, question):
        try:
            model_id = self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
            prompt = (
                "Answer the user's question using only evidence visible in the image. "
                "Be clear and concise, but include the detail needed to answer fully. "
                "If the answer cannot be determined from the image, say so. Do not invent "
                "details.\n\nQuestion: " + question
            )
            started = time.perf_counter()
            if external_ai_config(self.app_settings) is not None:
                text = run_external_ai_task(
                    "question", path, prompt, 384, self.shutdown_event,
                    self.app_settings,
                )
                model_label = self.app_settings.get("external_ai_model", "Local AI")
            else:
                text = run_mtmd_task(
                    "question", path, model_id, prompt, 384, self.shutdown_event
                )
                model_label = VISION_MODELS[model_id]["name"]
            elapsed = time.perf_counter() - started
            if not text.startswith(("Local vision", "Local AI service")):
                text += (
                    f"\n\nModel: {model_label}. "
                    f"Processing time: {elapsed:.1f} seconds."
                )
        except Exception as exc:
            logger.exception("Image question failed")
            text = f"Processing failed: {exc}"
        wx.CallAfter(self._finish_screen_description, path, text)

    def _prompt_photo_question(self):
        dialog = wx.Dialog(self, title="Ask About Photo")
        root = wx.BoxSizer(wx.VERTICAL)
        label = wx.StaticText(dialog, label="Question")
        question = wx.TextCtrl(dialog, style=wx.TE_PROCESS_ENTER)
        question.SetName("Question")
        buttons = dialog.CreateButtonSizer(wx.OK | wx.CANCEL)
        root.Add(label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 12)
        root.Add(question, 0, wx.ALL | wx.EXPAND, 12)
        root.Add(buttons, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM | wx.ALIGN_RIGHT, 12)
        dialog.SetSizerAndFit(root)
        dialog.SetMinSize((480, -1))
        question.Bind(
            wx.EVT_TEXT_ENTER,
            lambda event: dialog.EndModal(wx.ID_OK)
            if question.GetValue().strip()
            else None,
        )
        wx.CallAfter(question.SetFocus)
        result = dialog.ShowModal()
        value = question.GetValue().strip()
        dialog.Destroy()
        return value if result == wx.ID_OK else ""

    def ask_about_photo(self, event=None):
        if self.busy:
            return
        library_entry_id = None
        if self.mode_tabs.GetSelection() == self.TAB_LIBRARY:
            entry = self.selected_library_entry()
            path = resolve_stored_photo_path(entry) if entry else ""
            library_entry_id = entry.get("id") if entry else None
        else:
            path = ""
            for candidate, mode in reversed(
                list(zip(self.session_files, self.session_file_modes))
            ):
                if mode == "photo" and os.path.isfile(candidate):
                    path = candidate
                    break
        if not path or not os.path.isfile(path):
            wx.MessageBox(
                "The original photo could not be found.",
                "Ask About Photo",
                wx.OK | wx.ICON_INFORMATION,
                self,
            )
            return
        question = self._prompt_photo_question()
        if not question:
            return
        self.photo_cancel_event = threading.Event()
        self.begin_busy("", restore_focus=False)
        self.show_photo_cancel_dialog()
        threading.Thread(
            target=self._photo_question_worker,
            args=(path, question, library_entry_id, self.photo_cancel_event),
            name="ScanBox photo question",
            daemon=True,
        ).start()

    def _photo_question_worker(self, path, question, library_entry_id, cancel_event):
        prompt = (
            "Answer the user's question using only evidence visible in the image. "
            "Be clear and concise, but include the detail needed to answer fully. "
            "If the answer cannot be determined from the image, say so. Do not invent "
            "details.\n\nQuestion: " + question
        )
        started = time.perf_counter()
        try:
            if external_ai_config(self.app_settings) is not None:
                answer = run_external_ai_task(
                    "question", path, prompt, 384, cancel_event,
                    self.app_settings,
                )
                model_label = self.app_settings.get("external_ai_model", "Local AI")
            else:
                answer = run_mtmd_task(
                    "question", path, "qwen3_vl_2b", prompt, 384, cancel_event
                )
                model_label = "Qwen3-VL 2B"
            elapsed = time.perf_counter() - started
            if (
                answer != PHOTO_DESCRIPTION_CANCELLED
                and not answer.startswith(("Local vision", "Local AI service"))
            ):
                answer += (
                    f"\n\nModel: {model_label}. "
                    f"Processing time: {elapsed:.1f} seconds."
                )
        except Exception as exc:
            logger.exception("Photo question failed")
            answer = f"Processing failed: {exc}"
        finally:
            # Asking a question must not turn Qwen into the persistent model
            # when Florence remains the user's everyday selection.
            if (
                external_ai_config(self.app_settings) is None
                and self.app_settings.get("vision_model") != "qwen3_vl_2b"
            ):
                stop_mtmd_server()
        wx.CallAfter(
            self._finish_photo_question,
            question,
            answer,
            library_entry_id,
        )

    def _finish_photo_question(self, question, answer, library_entry_id):
        self.close_photo_cancel_dialog()
        self.photo_cancel_event = None
        self.end_busy()
        if answer == PHOTO_DESCRIPTION_CANCELLED:
            self.SetStatusText(PHOTO_DESCRIPTION_CANCELLED)
            return
        rendered = f"Question: {question}\n\n{answer}"
        if library_entry_id is not None:
            entry = self.selected_library_entry()
            if entry and entry.get("id") == library_entry_id:
                self.library_description.AppendText("\n\n" + rendered)
                self.library_description.SetFocusFromKbd()
        else:
            self.append_output("\n\n" + rendered + "\n\n")
            self.output_box.SetFocusFromKbd()
        self.update_controls()
        wx.CallLater(500, announce, answer)

    def _finish_screen_description(self, path, text):
        self.busy = False
        self.screen_capture_busy = False
        placeholder = self._screen_result_placeholder
        self._screen_result_placeholder = None
        if placeholder is not None:
            start, end, separator = placeholder
            self.output_box.Replace(start, end, separator + text + "\n\n")
            self.output_box.SetInsertionPoint(start + len(separator))
            self.output_box.ShowPosition(start + len(separator))
        else:
            self.append_output(text + "\n\n")
        self.update_controls()
        wx.CallLater(1000, announce, text)
        try:
            os.remove(path)
        except OSError:
            pass

    def on_pdf_ocr_mode_change(self, event):
        self.app_settings["pdf_ocr_mode"] = self.pdf_ocr_modes[
            self.pdf_ocr_radio.GetSelection()
        ]
        write_app_settings(self.app_settings)
        self.update_controls()

    def update_controls(self):
        has_session = bool(self.session_pages)
        camera_workflow_active = bool(
            self.camera_capture_active
            or self.camera_alignment_active
        )
        interaction_locked = bool(
            self.busy or camera_workflow_active or self.image_export_active
        )
        # Scan and Import are no longer split one-tab-per-mode - both handle
        # document and photo actions together - so which is "active" for
        # Save/status purposes now follows whichever action ran last rather
        # than which tab is selected. Photo Library is still its own tab.
        is_photo_mode = self.last_active_mode == "photo"
        is_library_mode = self.mode_tabs.GetSelection() == self.TAB_LIBRARY
        active_mode = "photo" if is_photo_mode else "document"
        has_scanned_images = is_photo_mode and any(
            source in {"scan", "camera"} and mode == active_mode
            for source, mode in zip(
                self.session_file_sources, self.session_file_modes
            )
        )
        has_document_text = any(
            mode == "document" for mode in self.session_page_modes
        )
        can_save_text = bool(
            has_document_text
            and not is_photo_mode
            and self.app_settings.get("append_text_to_buffer", False)
        )
        is_import_photo_session = bool(
            is_photo_mode
            and has_session
            and self.session_file_sources
            and all(
                source == "import" and mode == "photo"
                for source, mode in zip(
                    self.session_file_sources, self.session_file_modes
                )
            )
        )
        qwen_questions_available = bool(
            external_ai_config(self.app_settings) is not None
            or (
                _find_mtmd_model_files("qwen3_vl_2b") is not None
                and _find_mtmd_runner() is not None
            )
        )
        has_current_photo = any(
            mode == "photo" and os.path.isfile(path)
            for path, mode in zip(self.session_files, self.session_file_modes)
        )
        selected_library = self.selected_library_entry() if is_library_mode else None
        has_library_photo = bool(
            selected_library
            and os.path.isfile(resolve_stored_photo_path(selected_library))
        )

        self.save_text_btn.SetLabel("Save Text")
        self.save_img_btn.SetLabel(
            "Save Scanned Photo" if is_photo_mode else "Save Scanned Images"
        )
        self.discard_btn.SetLabel(
            "Close" if is_import_photo_session else "Discard Session"
        )

        self.output_box.Show(not is_library_mode)
        self.mode_tabs_sizer_item.SetProportion(1 if is_library_mode else 0)
        self.save_text_btn.Show(can_save_text and not is_library_mode)
        self.save_img_btn.Show(has_scanned_images and not is_library_mode)
        self.ask_about_photo_btn.Show(
            is_photo_mode
            and has_current_photo
            and qwen_questions_available
            and not is_library_mode
        )
        self.choose_another_photo_btn.Show(
            is_import_photo_session and not is_library_mode
        )
        self.discard_btn.Show(has_session and not is_library_mode)
        self.cancel_batch_btn.Show(
            getattr(self, "batch_cancel_event", None) is not None and self.busy
        )
        self.stop_camera_btn.Show(camera_workflow_active)
        self.stop_camera_btn.SetLabel("Stop Camera Capture")

        # Keep navigation and existing Results text reachable while the two
        # backslash screen workflows run. Their action controls remain
        # disabled below, preventing concurrent work.
        self.mode_tabs.Enable(not interaction_locked or self.screen_capture_busy)
        self.document_scan_btn.Enable(not interaction_locked)
        self.scan_save_images_btn.Enable(not interaction_locked)
        self.document_camera_btn.Enable(not interaction_locked)
        self.document_import_btn.Enable(not interaction_locked)
        self.pdf_ocr_radio.Enable(not interaction_locked)
        self.photo_scan_btn.Enable(not interaction_locked)
        if hasattr(self, "photo_camera_btn"):
            self.photo_camera_btn.Enable(not interaction_locked)
        self.camera_alignment_btn.Enable(not interaction_locked)
        self.photo_import_btn.Enable(not interaction_locked)
        self.photo_batch_import_btn.Enable(not interaction_locked)
        self.library_list.Enable(not interaction_locked)
        self.library_ask_btn.Enable(
            qwen_questions_available and has_library_photo and not interaction_locked
        )
        self.settings_btn.Enable(not interaction_locked)
        self.save_text_btn.Enable(
            can_save_text and not self.busy
        )
        self.save_img_btn.Enable(has_scanned_images and not self.busy)
        self.ask_about_photo_btn.Enable(
            has_current_photo and qwen_questions_available and not self.busy
        )
        self.choose_another_photo_btn.Enable(
            is_import_photo_session and not self.busy
        )
        self.discard_btn.Enable(has_session and not self.busy)
        self.cancel_batch_btn.Enable(self.busy)
        self.stop_camera_btn.Enable(camera_workflow_active)

        if self.busy and self.busy_status_visible:
            self.SetStatusText("Please wait. Processing...")
        elif is_library_mode:
            self.SetStatusText(f"Stored photo descriptions: {len(self.photo_library)}")
        elif has_session:
            noun = "Photos processed" if is_photo_mode else "Pages scanned"
            self.SetStatusText(f"{noun}: {self.page_counter}")
        else:
            self.SetStatusText("No active session")

        self.Layout()

    def selected_library_index(self):
        return self.library_list.GetFirstSelected()

    def selected_library_entry(self):
        index = self.selected_library_index()
        if 0 <= index < len(self.photo_library):
            return self.photo_library[index]
        return None

    def checked_library_indices(self):
        return [
            index
            for index in range(self.library_list.GetItemCount())
            if self.library_list.IsItemChecked(index)
        ]

    def select_all_library_entries(self, event=None):
        for index in range(self.library_list.GetItemCount()):
            self.library_list.CheckItem(index, True)
        announce("All photo descriptions selected.")

    def refresh_photo_library(self, selected_id=None):
        self.photo_library = load_photo_library()
        self.library_list.DeleteAllItems()
        selected_index = -1
        for index, entry in enumerate(self.photo_library):
            path = resolve_stored_photo_path(entry)
            description = entry.get("description", "").replace("\r", " ").replace("\n", " ")
            row = self.library_list.InsertItem(index, description[:160])
            self.library_list.SetItem(
                row, 1, "Available" if os.path.isfile(path) else "Missing"
            )
            self.library_list.SetItem(
                row, 2, os.path.basename(path) or "(No filename)"
            )
            if selected_id and entry.get("id") == selected_id:
                selected_index = index
        if selected_index < 0 and self.photo_library:
            selected_index = 0
        if selected_index >= 0:
            self.library_list.Select(selected_index)
            self.library_list.Focus(selected_index)
        else:
            self.library_description.SetValue("")
        if hasattr(self, "library_open_btn"):
            has_entry = bool(self.photo_library)
            self.library_open_btn.Enable(has_entry)
            self.library_folder_btn.Enable(has_entry)
            self.library_copy_btn.Enable(has_entry)
            qwen_questions_available = bool(
                _find_mtmd_model_files("qwen3_vl_2b") is not None
                and _find_mtmd_runner() is not None
            )
            selected_entry = self.selected_library_entry()
            self.library_ask_btn.Enable(
                bool(
                    selected_entry
                    and qwen_questions_available
                    and os.path.isfile(resolve_stored_photo_path(selected_entry))
                    and not self.busy
                )
            )
            self.library_remove_btn.Enable(has_entry)
            self.library_select_all_btn.Enable(has_entry and not self.busy)

    def on_library_selected(self, event):
        entry = self.selected_library_entry()
        if entry:
            path = resolve_stored_photo_path(entry)
            self.library_description.SetValue(
                f"{entry.get('description', '')}\n\nPath: {path}"
            )
            self.update_controls()
        event.Skip()

    def on_library_key(self, event):
        key = event.GetKeyCode()
        if key == wx.WXK_DELETE:
            self.remove_library_entry()
            return
        if key == wx.WXK_SPACE:
            index = self.library_list.GetFocusedItem()
            if index >= 0:
                self.library_list.CheckItem(
                    index, not self.library_list.IsItemChecked(index)
                )
            return
        if key in (wx.WXK_RETURN, wx.WXK_NUMPAD_ENTER):
            self.open_library_photo()
            return
        if key in (wx.WXK_LEFT, wx.WXK_RIGHT) and self.photo_library:
            current = max(0, self.selected_library_index())
            change = -1 if key == wx.WXK_LEFT else 1
            target = max(0, min(len(self.photo_library) - 1, current + change))
            self.library_list.Select(current, False)
            self.library_list.Select(target)
            self.library_list.Focus(target)
            return
        event.Skip()

    def on_library_context_menu(self, event):
        if not self.selected_library_entry():
            return
        menu = wx.Menu()
        open_item = menu.Append(wx.ID_OPEN, "Open Photo")
        folder_item = menu.Append(wx.ID_ANY, "Open Containing Folder")
        copy_item = menu.Append(wx.ID_COPY, "Copy Path")
        menu.AppendSeparator()
        remove_item = menu.Append(wx.ID_DELETE, "Remove Stored Description")
        menu.Bind(wx.EVT_MENU, self.open_library_photo, open_item)
        menu.Bind(wx.EVT_MENU, self.open_library_folder, folder_item)
        menu.Bind(wx.EVT_MENU, self.copy_library_path, copy_item)
        menu.Bind(wx.EVT_MENU, self.remove_library_entry, remove_item)
        self.PopupMenu(menu)
        menu.Destroy()

    def open_library_photo(self, event=None):
        entry = self.selected_library_entry()
        path = resolve_stored_photo_path(entry or {})
        if not path or not os.path.isfile(path):
            wx.MessageBox("The stored photo could not be found.", "Photo unavailable")
            return
        open_with_default_application(path)

    def open_library_folder(self, event=None):
        entry = self.selected_library_entry()
        path = resolve_stored_photo_path(entry or {})
        if not path or not os.path.exists(os.path.dirname(path)):
            wx.MessageBox("The containing folder could not be found.", "Folder unavailable")
            return
        reveal_in_file_manager(path)

    def copy_library_path(self, event=None):
        entry = self.selected_library_entry()
        path = resolve_stored_photo_path(entry or {})
        if not path:
            return
        if wx.TheClipboard.Open():
            try:
                wx.TheClipboard.SetData(wx.TextDataObject(path))
                wx.TheClipboard.Flush()
            finally:
                wx.TheClipboard.Close()
        announce("Photo path copied.")

    def remove_library_entry(self, event=None):
        checked_indices = self.checked_library_indices()
        if checked_indices:
            entries = [self.photo_library[index] for index in checked_indices]
        else:
            entry = self.selected_library_entry()
            entries = [entry] if entry else []
        if not entries:
            return
        count = len(entries)
        description = (
            "these stored descriptions" if count != 1 else "this stored description"
        )
        answer = wx.MessageBox(
            f"Remove {description}? The original photo will not be deleted.",
            "Remove stored description",
            wx.YES_NO | wx.NO_DEFAULT | wx.ICON_QUESTION,
        )
        if answer != wx.YES:
            return
        entry_ids = {entry.get("id") for entry in entries}
        remaining = [
            item for item in self.photo_library if item.get("id") not in entry_ids
        ]
        save_photo_library(remaining)
        self.refresh_photo_library()
        self.update_controls()
        announce(
            "Stored description%s removed. The photo%s not deleted."
            % ("s" if count != 1 else "", "s were" if count != 1 else " was")
        )

    def begin_busy(self, message, restore_focus=True):
        self.busy = True
        self.restore_focus_after_busy = restore_focus
        self.busy_status_visible = False
        self.busy_status_token += 1
        token = self.busy_status_token
        if restore_focus and not self.IsActive():
            # Bring ScanBox back to the front if focus drifted elsewhere
            # while a scan was underway (e.g. an external scanner dialog
            # took it). Skip this when ScanBox is already the active
            # window - as it always is right after Import, since the file
            # picker is owned by this same window and never hands focus
            # away. Re-asserting focus/activation when nothing actually
            # moved it fires a redundant focus-change announcement, which
            # was colliding with the screen reader announcing the result a
            # moment later on Import (reported as a clipped "ScanBox!" cut
            # off right before the description).
            self.Raise()
            self.RequestUserAttention()
        # Move focus off whatever's about to be disabled (typically the
        # button just clicked, e.g. Scan Photo) before update_controls()
        # disables the whole tab strip and its buttons. Otherwise Windows
        # drops focus out from under the disabled control with nowhere to
        # go, and NVDA reads out the disabled ancestor chain on the way
        # past it - "Tab control unavailable, Scan property page
        # unavailable, Scan Photo button unavailable" - instead of just
        # moving on quietly. The frame itself is never disabled, so parking
        # focus there is a safe, silent landing spot; whatever finishes the
        # action (e.g. output_box.SetFocusFromKbd()) moves focus on again
        # once there's a result to read.
        focused = wx.Window.FindFocus()
        if focused and focused is not self:
            self.SetFocus()
        self.update_controls()
        if message:
            wx.CallLater(1000, self.show_busy_status, token, message)
        wx.SafeYield(self, onlyIfNeeded=True)

    def show_busy_status(self, token, message):
        if not self.busy or token != self.busy_status_token:
            return
        self.busy_status_visible = True
        self.SetStatusText("Please wait. Processing...")
        announce(message)

    def end_busy(self):
        self.busy = False
        self.busy_status_visible = False
        self.busy_status_token += 1
        if getattr(self, "restore_focus_after_busy", True):
            self.Raise()
        self.update_controls()

    @staticmethod
    def _wia_scanner_name(device_info):
        try:
            return str(device_info.Properties("Name").Value).strip()
        except Exception:
            try:
                for index in range(1, device_info.Properties.Count + 1):
                    prop = device_info.Properties[index]
                    if str(getattr(prop, "Name", "")).lower() == "name":
                        return str(prop.Value).strip()
            except Exception:
                pass
        return "Unnamed scanner"

    def available_scanners(self):
        """Return stable identifiers and names for all detected scanners."""
        if sys.platform == "win32":
            if win32com is None:
                return []
            com_initialized = False
            try:
                if (
                    pythoncom is not None
                    and threading.current_thread() is not threading.main_thread()
                ):
                    pythoncom.CoInitialize()
                    com_initialized = True
                manager = win32com.client.Dispatch("WIA.DeviceManager")
                scanners = []
                seen = set()
                for index in range(1, manager.DeviceInfos.Count + 1):
                    info = manager.DeviceInfos[index]
                    identifier = str(info.DeviceID)
                    if info.Type == 1 and identifier not in seen:
                        seen.add(identifier)
                        scanners.append({
                            "id": identifier,
                            "name": self._wia_scanner_name(info),
                        })
                return scanners
            except Exception:
                logger.exception("Could not enumerate WIA scanners")
                return []
            finally:
                if com_initialized:
                    pythoncom.CoUninitialize()
        if sys.platform == "darwin" and os.path.isfile(MACOS_CAPTURE_HELPER):
            try:
                listed = subprocess.run(
                    [MACOS_CAPTURE_HELPER, "list-scanners"],
                    check=True,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=40,
                )
                scanners = []
                for line in listed.stdout.splitlines():
                    kind, separator, value = line.partition("\t")
                    identifier, separator2, name = value.partition("\t")
                    if kind == "scanner" and separator and separator2 and identifier:
                        scanners.append({
                            "id": identifier,
                            "name": name or "Unnamed scanner",
                        })
                return scanners
            except (OSError, subprocess.SubprocessError):
                logger.exception("Could not enumerate macOS scanners")
        return []

    def selected_scanner(self, scanners):
        """Resolve the saved scanner, prompting only for Ask me each time."""
        saved_id = self.app_settings.get("scanner_id", "")
        if saved_id:
            return next(
                (scanner for scanner in scanners if scanner["id"] == saved_id),
                None,
            )
        dialog = wx.SingleChoiceDialog(
            self,
            "Choose the scanner to use.",
            "Select scanner",
            [scanner["name"] for scanner in scanners],
        )
        try:
            if dialog.ShowModal() != wx.ID_OK:
                return None
            return scanners[dialog.GetSelection()]
        finally:
            dialog.Destroy()

    def scan_page(self):
        if sys.platform == "darwin":
            return self.scan_macos_scanner()
        if win32com is None:
            wx.MessageBox(
                "Windows scanner support is not available. Use Import Image or "
                "install the Windows scanner support package.",
                "Scanner unavailable",
            )
            return None

        scanners = self.available_scanners()
        if not scanners:
            wx.MessageBox(
                "No Windows scanner was found. To use a USB document camera "
                "instead, choose OCR using Camera on the Scan tab.",
                "Scanner not found",
            )
            return None
        scanner = self.selected_scanner(scanners)
        if scanner is None:
            if self.app_settings.get("scanner_id", ""):
                wx.MessageBox(
                    "The scanner selected in Settings is not available. Connect "
                    "it or choose another scanner in Settings.",
                    "Scanner not found",
                )
            return None

        cd = win32com.client.Dispatch("WIA.CommonDialog")
        # WIA device type 1 is a scanner. USB/UVC cameras are deliberately
        # handled by capture_camera_page so the two capture sources cannot
        # mask or take precedence over one another.
        device_manager = win32com.client.Dispatch("WIA.DeviceManager")
        info = next(
            (
                device_manager.DeviceInfos[index]
                for index in range(1, device_manager.DeviceInfos.Count + 1)
                if str(device_manager.DeviceInfos[index].DeviceID) == scanner["id"]
            ),
            None,
        )
        if info is None:
            return None
        dev = info.Connect()

        item = dev.Items[1]
        tiff_format = "{B96B3CB1-0728-11D3-9D7B-0000F81EF32E}"
        img = cd.ShowTransfer(item, tiff_format)

        path = os.path.join(TEMP_DIR, f"scan_{uuid.uuid4().hex}.tif")
        with open(path, "wb") as f:
            f.write(img.FileData.BinaryData)
        return path

    def scan_macos_scanner(self):
        """Scan from the configured ImageCaptureCore scanner on macOS."""
        if not os.path.isfile(MACOS_CAPTURE_HELPER):
            wx.MessageBox(
                "The native macOS scanner component is missing. Reinstall "
                "ScanBox and try again.",
                "Scanner unavailable",
            )
            return None
        scanners = self.available_scanners()
        if not scanners:
            logger.warning("macOS ImageCaptureCore reported no scanners")
            wx.MessageBox(
                "No scanner was found. Confirm that the scanner is switched "
                "on and available to macOS, then try again.",
                "Scanner not found",
            )
            return None

        scanner = self.selected_scanner(scanners)
        if scanner is None:
            if self.app_settings.get("scanner_id", ""):
                wx.MessageBox(
                    "The scanner selected in Settings is not available. Connect "
                    "it or choose another scanner in Settings.",
                    "Scanner not found",
                )
            return None
        identifier, name = scanner["id"], scanner["name"]
        announce(f"Scanning with {name}. Please wait.")
        try:
            result = subprocess.run(
                [MACOS_CAPTURE_HELPER, "scan", identifier, TEMP_DIR],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=600,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            logger.exception("macOS scanner acquisition failed")
            detail = getattr(exc, "stderr", "") or str(exc)
            logger.error("macOS scanner acquisition detail: %s", detail.strip())
            wx.MessageBox(
                f"ScanBox could not complete the scan.\n\n{detail.strip()}",
                "Scan failed",
            )
            return None
        for line in result.stdout.splitlines():
            kind, separator, value = line.partition("\t")
            if kind == "scan" and separator and os.path.isfile(value):
                return value
        wx.MessageBox(
            "The scanner finished without returning an image to ScanBox.",
            "Scan failed",
        )
        return None

    def camera_backend(self):
        return (
            getattr(cv2, "CAP_AVFOUNDATION", 0)
            if sys.platform == "darwin"
            else getattr(cv2, "CAP_DSHOW", 0)
        )

    def available_camera_indexes(self):
        if sys.platform == "darwin":
            return sorted(self.camera_device_names())
        if cv2 is None:
            return []
        backend = self.camera_backend()
        available = []
        for index in range(6):
            candidate = cv2.VideoCapture(index, backend)
            try:
                if candidate.isOpened():
                    available.append(index)
            finally:
                candidate.release()
        return available

    def camera_device_names(self):
        """Return capture-index names from the platform's camera API."""
        names = {}
        if sys.platform == "win32":
            if _DirectShowFilterGraph is not None:
                try:
                    for index, name in enumerate(
                        _DirectShowFilterGraph().get_input_devices()
                    ):
                        if name:
                            names[index] = str(name)
                except Exception:
                    logger.exception("Could not enumerate DirectShow camera names")
            if not names:
                command = (
                    "Get-CimInstance Win32_PnPEntity | "
                    "Where-Object { $_.PNPClass -eq 'Camera' -or "
                    "$_.Service -eq 'usbvideo' } | "
                    "Select-Object -ExpandProperty Name | ConvertTo-Json -Compress"
                )
                try:
                    result = subprocess.run(
                        ["powershell.exe", "-NoProfile", "-NonInteractive",
                         "-Command", command],
                        check=True,
                        capture_output=True,
                        text=True,
                        encoding="utf-8",
                        errors="replace",
                        timeout=10,
                        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                    )
                    found = json.loads(result.stdout or "[]")
                    if isinstance(found, str):
                        found = [found]
                    for index, name in enumerate(found):
                        if name:
                            names[index] = str(name)
                except Exception:
                    logger.exception("Could not enumerate Windows camera names")
        elif sys.platform == "darwin" and os.path.isfile(MACOS_CAPTURE_HELPER):
            try:
                result = subprocess.run(
                    [MACOS_CAPTURE_HELPER, "list-cameras"],
                    check=True,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=20,
                )
                for line in result.stdout.splitlines():
                    kind, separator, value = line.partition("\t")
                    index_text, separator2, name = value.partition("\t")
                    if kind == "camera" and separator and separator2:
                        names[int(index_text)] = name
            except Exception:
                logger.exception("Could not enumerate AVFoundation camera names")
        return names

    def choose_camera(self, front_facing_only=False):
        """Return a camera index selected from the detected devices."""
        if cv2 is None and sys.platform != "darwin":
            wx.MessageBox(
                "Camera capture requires OpenCV. Install the full ScanBox "
                "dependencies or use Import instead.",
                "Camera unavailable",
            )
            return None
        available = self.available_camera_indexes()
        if not available:
            wx.MessageBox(
                "ScanBox could not find an available camera. Connect the "
                "camera, allow camera access, and try again.",
                "Camera unavailable",
            )
            return None

        selected_index = available[0]
        if len(available) > 1:
            device_names = self.camera_device_names()
            raw_labels = [
                device_names.get(index, f"Camera {index + 1}")
                for index in available
            ]
            labels = [
                (
                    f"{name} (Camera {index + 1})"
                    if raw_labels.count(name) > 1
                    else name
                )
                for index, name in zip(available, raw_labels)
            ]
            dialog = wx.Dialog(self, title="Select Camera")
            root = wx.BoxSizer(wx.VERTICAL)
            explanation = wx.StaticText(
                dialog,
                label=(
                    "Choose the camera facing you."
                    if front_facing_only
                    else "Choose the camera to use."
                ),
            )
            camera_label = wx.StaticText(dialog, label="Camera")
            camera_choice = wx.ComboBox(
                dialog,
                choices=labels,
                style=wx.CB_READONLY,
            )
            camera_choice.SetName("Camera")
            camera_choice.SetSelection(0)
            root.Add(explanation, 0, wx.ALL | wx.EXPAND, 10)
            root.Add(camera_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
            root.Add(camera_choice, 0, wx.ALL | wx.EXPAND, 10)
            root.Add(
                dialog.CreateButtonSizer(wx.OK | wx.CANCEL),
                0,
                wx.ALL | wx.ALIGN_RIGHT,
                10,
            )
            dialog.SetSizerAndFit(root)
            try:
                if dialog.ShowModal() != wx.ID_OK:
                    return None
                selected_index = available[camera_choice.GetSelection()]
            finally:
                dialog.Destroy()
        return selected_index

    def open_camera_capture_device(self, camera_index):
        """Open and retain a camera so bridge devices remain initialised."""
        self.close_camera_capture_device()
        camera = cv2.VideoCapture(camera_index, self.camera_backend())
        self.camera_capture_device = camera
        try:
            buffer_property = getattr(cv2, "CAP_PROP_BUFFERSIZE", None)
            if buffer_property is not None:
                camera.set(buffer_property, 1)
        except Exception:
            logger.debug("Camera backend does not support buffer sizing", exc_info=True)
        logger.info(
            "Opened camera index %s; opened=%s",
            camera_index,
            camera.isOpened(),
        )
        return camera

    def close_camera_capture_device(self):
        camera = getattr(self, "camera_capture_device", None)
        self.camera_capture_device = None
        if camera is not None:
            try:
                camera.release()
            except Exception:
                logger.exception("Could not release camera")

    def capture_camera_page(self, camera_index=None):
        """Capture one still from the default USB or built-in camera.

        A repositionable USB document camera is suitable for documents. The
        existing pipeline performs page detection, perspective correction,
        and OCR after this method saves the frame.
        """
        if camera_index is None:
            camera_index = self.choose_camera()
            if camera_index is None:
                return None
        if sys.platform == "darwin":
            path = os.path.join(TEMP_DIR, f"camera_{uuid.uuid4().hex}.png")
            try:
                result = subprocess.run(
                    [
                        MACOS_CAPTURE_HELPER,
                        "capture-camera",
                        str(camera_index),
                        path,
                    ],
                    check=True,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=30,
                )
                if not os.path.isfile(path):
                    raise RuntimeError(
                        result.stderr.strip()
                        or "The camera did not return an image."
                    )
                play_shutter_sound()
                return path
            except Exception as exc:
                logger.exception("Native macOS camera capture failed")
                detail = getattr(exc, "stderr", "") or str(exc)
                wx.MessageBox(
                    "ScanBox could not capture an image from the camera.\n\n"
                    + detail.strip(),
                    "Camera capture failed",
                )
                _remove_quietly(path)
                return None
        camera = getattr(self, "camera_capture_device", None)
        owns_camera = camera is None
        if owns_camera:
            camera = cv2.VideoCapture(camera_index, self.camera_backend())
        try:
            if not camera.isOpened():
                permission_help = (
                    "Allow camera access in System Settings > Privacy & "
                    "Security > Camera, then try again."
                    if sys.platform == "darwin"
                    else "Check Windows camera privacy settings, then try again."
                )
                wx.MessageBox(
                    "ScanBox could not open the selected camera. Connect or "
                    f"reconnect the camera. {permission_help}",
                    "Camera unavailable",
                )
                return None

            # Read over real elapsed time. Freedom Scientific's PEARL
            # DirectShow bridge and some document cameras need time to start;
            # rapid fixed-count reads can save a blank startup image.
            frame, attempts = _read_settled_camera_frame(camera)
            logger.info(
                "Camera index %s capture completed after %s reads; frame=%s",
                camera_index,
                attempts,
                None if frame is None else getattr(frame, "shape", "unknown"),
            )
            if frame is None:
                wx.MessageBox(
                    "The camera opened but did not return a usable image. Close "
                    "other camera applications, reconnect the camera, and try again.",
                    "Camera capture failed",
                )
                return None

            path = os.path.join(TEMP_DIR, f"camera_{uuid.uuid4().hex}.png")
            if not cv2.imwrite(path, frame):
                raise OSError("The captured camera image could not be saved.")
            play_shutter_sound()
            return path
        except Exception as exc:
            logger.exception("Camera capture failed")
            wx.MessageBox(
                f"ScanBox could not capture an image from the camera.\n\n{exc}",
                "Camera capture failed",
            )
            return None
        finally:
            if owns_camera:
                camera.release()

    def photo_mode_blocked(self, is_photo_mode):
        """In photo mode the model is required, so check before capturing."""
        if not is_photo_mode or vision_ready():
            return False
        self.offer_install_vision(
            "Photo descriptions require the selected local image model."
        )
        return True

    def on_scan(self, event, is_photo_mode):
        if self.photo_mode_blocked(is_photo_mode):
            return
        self.last_active_mode = "photo" if is_photo_mode else "document"
        if not is_photo_mode:
            # A prior photograph description is not relevant while the user
            # is beginning a document-OCR task. Clear it before scanner
            # selection/acquisition, rather than leaving stale text on screen.
            self.prepare_output_buffer()
        path = self.scan_page()
        if path:
            self.process_image(path, is_photo_mode, "scan")

    def on_scan_and_save_images(self, event=None):
        """Scan without OCR and export one JPEG or one multi-page PDF."""
        if self.busy or self.image_export_active:
            return
        choices = [
            "Single page image — save as JPEG",
            "Multiple page document — save as one PDF",
        ]
        dialog = wx.SingleChoiceDialog(
            self,
            "What would you like to scan and save?",
            "Scan and Save Images",
            choices,
        )
        dialog.SetSelection(0)
        try:
            if dialog.ShowModal() != wx.ID_OK:
                return
            multiple = dialog.GetSelection() == 1
        finally:
            dialog.Destroy()

        self.image_export_active = True
        self.image_export_multiple = multiple
        self.image_export_paths = []
        self.image_export_rotations = {}
        self.update_controls()
        self._scan_image_export_page()

    def _scan_image_export_page(self):
        path = self.scan_page()
        if not path:
            if self.image_export_paths and self.image_export_multiple:
                self._prompt_for_another_export_page()
            else:
                self._finish_image_export()
            return
        self.image_export_paths.append(path)
        self.begin_busy("ScanBox is checking the page orientation.")
        threading.Thread(
            target=self._image_export_orientation_worker,
            args=(path,),
            name="ScanBox image orientation",
            daemon=True,
        ).start()

    def _image_export_orientation_worker(self, path):
        try:
            rotation = detect_document_rotation(path)
            logger.info(
                "Image export orientation result for %s: %s",
                path,
                "unknown" if rotation is None else f"{rotation} degrees clockwise",
            )
        except Exception:
            logger.exception("Image export orientation failed for %s", path)
            rotation = None
        wx.CallAfter(self._finish_image_export_orientation, path, rotation)

    def _finish_image_export_orientation(self, path, rotation):
        self.end_busy()
        page_number = self.image_export_paths.index(path) + 1
        correction = self._confirm_image_export_rotation(page_number, rotation)
        if correction is None:
            self._finish_image_export()
            return
        self.image_export_rotations[path] = correction
        if self.image_export_multiple:
            self._prompt_for_another_export_page()
        else:
            self._save_single_page_image()

    def _confirm_image_export_rotation(self, page_number, rotation):
        if rotation is None:
            wx.MessageBox(
                f"ScanBox could not confidently determine the orientation of page "
                f"{page_number}. It will be saved as scanned.",
                "Page orientation not determined",
                wx.OK | wx.ICON_INFORMATION,
                self,
            )
            return 0
        if rotation == 0:
            wx.MessageBox(
                f"ScanBox reports that page {page_number} is upright. It will be "
                "saved without rotation.",
                "Page is upright",
                wx.OK | wx.ICON_INFORMATION,
                self,
            )
            return 0

        dialog = wx.MessageDialog(
            self,
            f"ScanBox reports that page {page_number} needs to be rotated "
            f"{rotation} degrees clockwise to make its text upright. It will "
            "be rotated before export.",
            "Page will be rotated",
            wx.YES_NO | wx.CANCEL | wx.YES_DEFAULT | wx.ICON_QUESTION,
        )
        dialog.SetYesNoCancelLabels(
            "Rotate and Continue", "Keep as Scanned", "Cancel"
        )
        try:
            result = dialog.ShowModal()
        finally:
            dialog.Destroy()
        if result == wx.ID_CANCEL:
            logger.info("User cancelled export at page %s rotation prompt", page_number)
            return None
        if result == wx.ID_YES:
            logger.info(
                "User accepted %s-degree rotation for export page %s",
                rotation,
                page_number,
            )
            return rotation
        logger.info("User kept export page %s as scanned", page_number)
        return 0

    def _prompt_for_another_export_page(self):
        page_count = len(self.image_export_paths)
        dialog = wx.MessageDialog(
            self,
            f"Page {page_count} is ready. Scan another page, or finish and "
            "save all pages as one PDF?",
            "Multiple page document",
            wx.YES_NO | wx.CANCEL | wx.YES_DEFAULT | wx.ICON_QUESTION,
        )
        dialog.SetYesNoCancelLabels(
            "Scan Another Page", "Finish and Save", "Cancel"
        )
        try:
            result = dialog.ShowModal()
        finally:
            dialog.Destroy()
        if result == wx.ID_YES:
            self._scan_image_export_page()
        elif result == wx.ID_NO:
            self._save_multiple_page_pdf()
        else:
            self._finish_image_export()

    def _save_single_page_image(self):
        dialog = wx.FileDialog(
            self,
            "Save scanned image",
            defaultDir=IMAGES_DIR,
            defaultFile="Captured document.jpg",
            wildcard="JPEG image (*.jpg)|*.jpg",
            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        destination = ""
        try:
            if dialog.ShowModal() == wx.ID_OK:
                source = self.image_export_paths[0]
                destination = self.write_scanned_jpeg(
                    source,
                    dialog.GetPath(),
                    self.image_export_rotations.get(source, 0),
                )
        except Exception as exc:
            logger.exception("Could not save scanned image")
            wx.MessageBox(
                f"ScanBox could not save the scanned image.\n\n{exc}",
                "Save failed",
                wx.OK | wx.ICON_ERROR,
                self,
            )
        finally:
            dialog.Destroy()
        self._finish_image_export()
        if destination:
            self.SetStatusText(f"Scanned image saved: {destination}")
            announce("Scanned image saved.")

    def _save_multiple_page_pdf(self):
        dialog = wx.FileDialog(
            self,
            "Save scanned document",
            defaultDir=IMAGES_DIR,
            defaultFile="Captured document.pdf",
            wildcard="PDF document (*.pdf)|*.pdf",
            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        destination = ""
        try:
            if dialog.ShowModal() == wx.ID_OK:
                destination = self.write_scanned_image_pdf(
                    self.image_export_paths,
                    dialog.GetPath(),
                    self.image_export_rotations,
                )
        except Exception as exc:
            logger.exception("Could not save scanned document")
            wx.MessageBox(
                f"ScanBox could not save the scanned document.\n\n{exc}",
                "Save failed",
                wx.OK | wx.ICON_ERROR,
                self,
            )
        finally:
            dialog.Destroy()
        page_count = len(self.image_export_paths)
        self._finish_image_export()
        if destination:
            self.SetStatusText(f"Scanned document saved: {destination}")
            announce(f"Scanned document saved with {page_count} pages.")

    def _finish_image_export(self):
        _remove_quietly(*self.image_export_paths)
        self.image_export_active = False
        self.image_export_multiple = False
        self.image_export_paths = []
        self.image_export_rotations = {}
        self.update_controls()
        wx.CallAfter(self.scan_save_images_btn.SetFocus)

    def on_camera_capture(self, event, is_photo_mode):
        if self.photo_mode_blocked(is_photo_mode):
            return
        delay = max(
            0,
            min(120, int(self.app_settings.get("camera_delay_seconds", 5))),
        )
        interval = max(
            1,
            min(
                3600,
                int(self.app_settings.get("camera_interval_seconds", 5)),
            ),
        )
        capture_count = max(
            0,
            min(100, int(self.app_settings.get("camera_capture_count", 1))),
        )
        camera_index = self.choose_camera()
        if camera_index is None:
            return
        self.last_active_mode = "photo" if is_photo_mode else "document"
        self.prepare_output_buffer()
        self.camera_capture_active = True
        self.camera_capture_stopping = False
        self.camera_capture_completed = 0
        self.camera_capture_target = capture_count
        self.camera_capture_interval = interval
        self.camera_capture_index = camera_index
        self.camera_capture_photo_mode = is_photo_mode
        if sys.platform != "darwin":
            # Open before the countdown and retain the DirectShow connection
            # for every page in this capture workflow.
            self.open_camera_capture_device(camera_index)
        self.update_controls()
        self.schedule_camera_capture(delay, first=True)


    def schedule_camera_capture(self, seconds, first=False):
        if not self.camera_capture_active:
            return
        self.camera_capture_remaining = max(0, int(seconds))
        if self.camera_capture_remaining == 0:
            wx.CallAfter(self.perform_scheduled_camera_capture)
            return
        self.camera_capture_timer.Start(1000)

    def on_camera_countdown(self, event=None):
        if not self.camera_capture_active:
            self.camera_capture_timer.Stop()
            return
        self.camera_capture_remaining -= 1
        if self.camera_capture_remaining <= 0:
            self.camera_capture_timer.Stop()
            self.perform_scheduled_camera_capture()

    def perform_scheduled_camera_capture(self):
        if not self.camera_capture_active:
            return
        path = self.capture_camera_page(self.camera_capture_index)
        if not path:
            self.complete_camera_workflow()
            return
        self.process_image(
            path,
            self.camera_capture_photo_mode,
            "camera",
            prepare_buffer=False,
            completion=self.camera_capture_finished,
        )

    def camera_capture_finished(self, text):
        # This callback is invoked only after the OCR/description worker has
        # returned and _finish_image_processing has placed the complete result
        # in the Results area. The repeat interval therefore never overlaps
        # generation of the preceding result.
        if text == PHOTO_DESCRIPTION_CANCELLED:
            self.complete_camera_workflow()
            return
        self.camera_capture_completed += 1
        announce(text)
        if not self.camera_capture_active or self.camera_capture_stopping:
            self.complete_camera_workflow()
            return
        if (
            self.camera_capture_target > 0
            and self.camera_capture_completed >= self.camera_capture_target
        ):
            self.complete_camera_workflow()
            return
        wx.CallLater(
            250,
            self.schedule_camera_capture,
            self.camera_capture_interval,
        )

    def complete_camera_workflow(self):
        self.camera_capture_timer.Stop()
        self.close_camera_capture_device()
        self.camera_capture_active = False
        self.camera_capture_stopping = False
        self.camera_capture_index = None
        self.update_controls()

    def stop_camera_workflow(self, event=None):
        if self.camera_alignment_active:
            self.stop_camera_alignment()
            return
        if not self.camera_capture_active:
            return
        self.camera_capture_stopping = True
        self.camera_capture_timer.Stop()
        if not self.busy:
            self.complete_camera_workflow()

    def start_camera_alignment(self, event=None):
        """Guide a user until one face is centred in a front-facing camera."""
        launch_focus = wx.Window.FindFocus()
        camera_index = self.choose_camera(front_facing_only=True)
        if camera_index is None:
            return
        cascade_path = ""
        if sys.platform != "darwin":
            cascade_root = getattr(getattr(cv2, "data", None), "haarcascades", "")
            cascade_path = os.path.join(
                cascade_root, "haarcascade_frontalface_default.xml"
            )
            if not os.path.isfile(cascade_path):
                wx.MessageBox(
                    "The local face-position detector is missing from this "
                    "ScanBox installation.",
                    "FaceAlign unavailable",
                )
                return
        self.camera_alignment_active = True
        self.camera_alignment_stop_event = threading.Event()
        self.facealign_restore_focus = launch_focus
        self.show_facealign_dialog("Starting camera. Please wait.")
        self.update_controls()
        announce(
            "FaceAlign started. Keep facing the camera and follow "
            "the position announcements."
        )
        threading.Thread(
            target=self._camera_alignment_worker,
            args=(
                camera_index,
                cascade_path,
                self.camera_alignment_stop_event,
            ),
            name="ScanBox FaceAlign",
            daemon=True,
        ).start()

    def show_facealign_dialog(self, message):
        dialog = wx.Dialog(
            self,
            title="FaceAlign",
            style=wx.DEFAULT_DIALOG_STYLE,
        )
        status = wx.StaticText(dialog, label=message)
        status.SetName("FaceAlign guidance")
        cancel = wx.Button(dialog, wx.ID_CANCEL, label="Cancel")
        cancel.SetName("Cancel FaceAlign")
        cancel.Bind(wx.EVT_BUTTON, self.cancel_facealign)
        dialog.Bind(wx.EVT_CLOSE, self.cancel_facealign)
        sizer = wx.BoxSizer(wx.VERTICAL)
        sizer.Add(status, 0, wx.ALL | wx.EXPAND, 16)
        sizer.Add(cancel, 0, wx.ALL | wx.ALIGN_RIGHT, 12)
        dialog.SetSizerAndFit(sizer)
        self.facealign_dialog = dialog
        self.facealign_status = status
        dialog.Show()
        dialog.Raise()
        cancel.SetFocus()

    def close_facealign_dialog(self):
        dialog = self.facealign_dialog
        self.facealign_dialog = None
        self.facealign_status = None
        restore_focus = self.facealign_restore_focus
        self.facealign_restore_focus = None
        if dialog is None and restore_focus is None:
            return
        if dialog is not None:
            dialog.Destroy()
        if restore_focus is None:
            restore_focus = self.camera_alignment_btn
        wx.CallAfter(self._restore_facealign_keyboard_focus, restore_focus)

    def _restore_facealign_keyboard_focus(self, control):
        try:
            if control is not None and control.IsEnabled():
                control.SetFocus()
                return
        except RuntimeError:
            pass
        try:
            self.camera_alignment_btn.SetFocus()
        except RuntimeError:
            pass

    def cancel_facealign(self, event=None):
        if isinstance(event, wx.CloseEvent) and event.CanVeto():
            event.Veto()
        self.stop_camera_alignment()

    def _camera_alignment_worker(self, camera_index, cascade_path, stop_event):
        if sys.platform == "darwin":
            self._macos_camera_alignment_worker(camera_index, stop_event)
            return
        face_detector = cv2.CascadeClassifier(cascade_path)
        camera = cv2.VideoCapture(camera_index, self.camera_backend())
        error = ""
        try:
            if face_detector.empty():
                raise RuntimeError("The local face-position detector could not load.")
            if not camera.isOpened():
                raise RuntimeError("The selected camera could not be opened.")
            while not stop_event.is_set():
                check_started = time.monotonic()
                frame = None
                for _ in range(5):
                    ok, candidate = camera.read()
                    if ok:
                        frame = candidate
                if frame is None:
                    guidance = "No camera image was received."
                else:
                    frame_height, frame_width = frame.shape[:2]
                    faces = face_detector.detectMultiScale(
                        frame,
                        scaleFactor=1.1,
                        minNeighbors=6,
                        minSize=(60, 60),
                    )
                    if len(faces) == 0:
                        guidance = "No face."
                    else:
                        x, y, width, height = max(
                            faces,
                            key=lambda item: int(item[2]) + int(item[3]),
                        )
                        x_percent = 100 * (x + width / 2) / frame_width
                        y_percent = 100 * (y + height / 2) / frame_height
                        position = (
                            f"X {round(x_percent)} percent, "
                            f"Y {round(y_percent)} percent."
                        )
                        # Raw webcam images are not mirrored. Low X therefore
                        # means the person is to their own right, matching Can
                        # You See Me's coordinate descriptions.
                        if x_percent < 47.5:
                            guidance = f"{position} Move left."
                        elif x_percent > 52.5:
                            guidance = f"{position} Move right."
                        elif y_percent < 47.5:
                            guidance = f"{position} Move down."
                        elif y_percent > 52.5:
                            guidance = f"{position} Move up."
                        else:
                            guidance = f"{position} Centred."
                wx.CallAfter(self._camera_alignment_announcement, guidance)
                remaining = 2.0 - (time.monotonic() - check_started)
                if stop_event.wait(max(0.0, remaining)):
                    break
        except Exception as exc:
            logger.exception("FaceAlign failed")
            error = str(exc)
        finally:
            camera.release()
            wx.CallAfter(
                self._finish_camera_alignment,
                error,
                stop_event,
            )

    def _macos_camera_alignment_worker(self, camera_index, stop_event):
        error = ""
        process = None
        try:
            process = subprocess.Popen(
                [MACOS_CAPTURE_HELPER, "face-align", str(camera_index)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            while not stop_event.is_set():
                line = process.stdout.readline()
                if not line:
                    if process.poll() is not None:
                        detail = process.stderr.read().strip()
                        if process.returncode:
                            raise RuntimeError(detail or "FaceAlign stopped unexpectedly.")
                        break
                    continue
                fields = line.strip().split("\t")
                if fields[0] == "no-face":
                    guidance = "No face."
                elif len(fields) == 3 and fields[0] == "face":
                    x_percent = float(fields[1])
                    y_percent = float(fields[2])
                    position = (
                        f"X {round(x_percent)} percent, "
                        f"Y {round(y_percent)} percent."
                    )
                    if x_percent < 47.5:
                        guidance = f"{position} Move left."
                    elif x_percent > 52.5:
                        guidance = f"{position} Move right."
                    elif y_percent < 47.5:
                        guidance = f"{position} Move down."
                    elif y_percent > 52.5:
                        guidance = f"{position} Move up."
                    else:
                        guidance = f"{position} Centred."
                else:
                    continue
                wx.CallAfter(self._camera_alignment_announcement, guidance)
        except Exception as exc:
            logger.exception("Native macOS FaceAlign failed")
            error = str(exc)
        finally:
            if process is not None and process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=3)
                except subprocess.TimeoutExpired:
                    process.kill()
            wx.CallAfter(self._finish_camera_alignment, error, stop_event)

    def _camera_alignment_announcement(self, message):
        if self.camera_alignment_active:
            self.SetStatusText(message)
            if self.facealign_status is not None:
                self.facealign_status.SetLabel(message)
                self.facealign_dialog.Layout()
            announce(message)

    def _finish_camera_alignment(self, error, stop_event):
        if stop_event is not self.camera_alignment_stop_event:
            return
        was_active = self.camera_alignment_active
        self.camera_alignment_active = False
        self.camera_alignment_stop_event = None
        self.close_facealign_dialog()
        self.update_controls()
        if error:
            wx.MessageBox(
                f"FaceAlign could not continue.\n\n{error}",
                "FaceAlign failed",
            )
        elif was_active:
            announce("FaceAlign stopped.")

    def stop_camera_alignment(self):
        if self.camera_alignment_stop_event is not None:
            self.camera_alignment_stop_event.set()
        self.camera_alignment_active = False
        self.close_facealign_dialog()
        self.update_controls()
        announce("FaceAlign stopped.")

    def on_import_image(self, event, is_photo_mode):
        # Let the activating mouse/key event finish before Windows creates its
        # native picker. Otherwise the picker can inherit that input and put
        # Explorer's Search box into its active search state.
        if event is not None:
            wx.CallAfter(self.on_import_image, None, is_photo_mode)
            return
        if self.photo_mode_blocked(is_photo_mode):
            return
        self.last_active_mode = "photo" if is_photo_mode else "document"
        wildcard = (
            "Images and PDF (*.jpg;*.jpeg;*.png;*.tif;*.tiff;*.pdf)|"
            "*.jpg;*.jpeg;*.png;*.tif;*.tiff;*.pdf"
        )
        last_dir = self.app_settings.get("last_import_dir", "")
        start_dir = last_dir if last_dir and os.path.isdir(last_dir) else (
            os.path.expanduser("~") if sys.platform == "darwin" else IMAGES_DIR
        )

        def complete(source):
            self.app_settings["last_import_dir"] = os.path.dirname(source)
            write_app_settings(self.app_settings)
            try:
                if os.path.splitext(source)[1].lower() == ".pdf":
                    if is_photo_mode:
                        raise ValueError(
                            "PDF import is available via Import Document."
                        )
                    self.process_pdf(source)
                else:
                    path = self.copy_imported_image(source)
                    self.process_image(
                        path, is_photo_mode, "import", library_path=source
                    )
            except Exception as exc:
                wx.MessageBox(
                    f"ScanBox could not import this file: {exc}", "Import failed"
                )

        # macOS: use AppleScript's `choose file` via osascript, run as a
        # genuinely separate OS process with its own event loop, for real
        # Finder-style keyboard/VoiceOver navigation. wx.FileDialog and an
        # in-process NSOpenPanel via PyObjC were both tried first, but
        # neither could actually escape wxWidgets' own ownership of the
        # single NSApplication/event loop for this process - see git
        # history.
        if sys.platform == "darwin":
            choose_macos_files_async(
                self,
                "Import image or PDF",
                start_dir,
                lambda paths: complete(paths[0]) if paths else None,
                extensions=("jpg", "jpeg", "png", "tif", "tiff", "pdf"),
            )
            return

        dlg = wx.FileDialog(
            self,
            "Import image or PDF",
            defaultDir=start_dir,
            wildcard=wildcard,
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST,
        )
        try:
            if dlg.ShowModal() == wx.ID_OK:
                complete(dlg.GetPath())
        finally:
            dlg.Destroy()

    def on_document_import(self, event=None):
        # Let the activating mouse/key event finish before Windows creates its
        # native picker, same as on_import_image above.
        if event is not None:
            wx.CallAfter(self.on_document_import, None)
            return
        if self.photo_mode_blocked(False):
            return
        self.last_active_mode = "document"
        self.prepare_output_buffer()
        wildcard = (
            "Images and PDF (*.jpg;*.jpeg;*.png;*.tif;*.tiff;*.pdf)|"
            "*.jpg;*.jpeg;*.png;*.tif;*.tiff;*.pdf"
        )
        last_dir = self.app_settings.get("last_import_dir", "")
        start_dir = last_dir if last_dir and os.path.isdir(last_dir) else (
            os.path.expanduser("~") if sys.platform == "darwin" else IMAGES_DIR
        )

        def complete(source):
            self.app_settings["last_import_dir"] = os.path.dirname(source)
            write_app_settings(self.app_settings)
            try:
                if os.path.splitext(source)[1].lower() == ".pdf":
                    self.process_pdf(source)
                else:
                    path = self.copy_imported_image(source)
                    self.process_image(path, False, "import", library_path=source)
            except Exception as exc:
                wx.MessageBox(
                    f"ScanBox could not import this file: {exc}", "Import failed"
                )

        if sys.platform == "darwin":
            choose_macos_files_async(
                self,
                "Choose a document to import",
                start_dir,
                lambda paths: complete(paths[0]) if paths else None,
                extensions=("jpg", "jpeg", "png", "tif", "tiff", "pdf"),
            )
            return

        dlg = wx.FileDialog(
            self,
            "Choose a document to import",
            defaultDir=start_dir,
            wildcard=wildcard,
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST,
        )
        try:
            if dlg.ShowModal() == wx.ID_OK:
                complete(dlg.GetPath())
        finally:
            dlg.Destroy()

    def batch_import_photos(self, event=None):
        if self.photo_mode_blocked(True):
            return
        self.last_active_mode = "photo"
        last_dir = self.app_settings.get("last_import_dir")
        batch_dir = last_dir if last_dir and os.path.isdir(last_dir) else (
            os.path.expanduser("~") if sys.platform == "darwin" else IMAGES_DIR
        )
        if sys.platform == "darwin":
            choose_macos_files_async(
                self,
                "Choose photos for batch import",
                batch_dir,
                self._show_batch_import_dialog,
                multiple=True,
                extensions=("jpg", "jpeg", "png", "tif", "tiff"),
            )
            return

        picker = wx.FileDialog(
            self,
            "Choose photos for batch import",
            defaultDir=batch_dir,
            wildcard="Images (*.jpg;*.jpeg;*.png;*.tif;*.tiff)|*.jpg;*.jpeg;*.png;*.tif;*.tiff",
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST | wx.FD_MULTIPLE,
        )
        try:
            if picker.ShowModal() != wx.ID_OK:
                return
            paths = picker.GetPaths()
        finally:
            picker.Destroy()
        self._show_batch_import_dialog(paths)

    def _show_batch_import_dialog(self, paths):
        if not paths:
            return
        self.app_settings["last_import_dir"] = os.path.dirname(paths[0])
        write_app_settings(self.app_settings)

        dialog = wx.Dialog(self, title="Select photos to import")
        checklist = wx.ListCtrl(
            dialog, style=wx.LC_REPORT | wx.LC_SINGLE_SEL
        )
        checklist.EnableCheckBoxes(True)
        checklist.InsertColumn(0, "Photo", width=500)
        for path in paths:
            index = checklist.InsertItem(checklist.GetItemCount(), os.path.basename(path))
            checklist.CheckItem(index, True)
        sizer = wx.BoxSizer(wx.VERTICAL)
        sizer.Add(checklist, 1, wx.ALL | wx.EXPAND, 10)

        def on_batch_check(event):
            index = event.GetIndex()
            if index >= 0:
                state = "Checked" if checklist.IsItemChecked(index) else "Unchecked"
                announce(f"{state} {checklist.GetItemText(index)}")
            event.Skip()

        checklist.Bind(wx.EVT_LIST_ITEM_CHECKED, on_batch_check)
        add_more = wx.Button(dialog, label="Add more photos...")

        def append_more_paths(more_paths):
            if more_paths:
                self.app_settings["last_import_dir"] = os.path.dirname(more_paths[0])
                write_app_settings(self.app_settings)
            for path in more_paths:
                if path not in paths:
                    paths.append(path)
                    index = checklist.InsertItem(
                        checklist.GetItemCount(), os.path.basename(path)
                    )
                    checklist.CheckItem(index, True)

        def add_more_photos(event):
            more_dir = (
                self.app_settings.get("last_import_dir")
                if self.app_settings.get("last_import_dir")
                and os.path.isdir(self.app_settings.get("last_import_dir"))
                else (
                    os.path.expanduser("~") if sys.platform == "darwin" else IMAGES_DIR
                )
            )
            if sys.platform == "darwin":
                choose_macos_files_async(
                    dialog,
                    "Add photos",
                    more_dir,
                    append_more_paths,
                    multiple=True,
                    extensions=("jpg", "jpeg", "png", "tif", "tiff"),
                )
                return

            more_picker = wx.FileDialog(
                dialog,
                "Add photos",
                defaultDir=more_dir,
                wildcard="Images (*.jpg;*.jpeg;*.png;*.tif;*.tiff)|*.jpg;*.jpeg;*.png;*.tif;*.tiff",
                style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST | wx.FD_MULTIPLE,
            )
            try:
                if more_picker.ShowModal() != wx.ID_OK:
                    return
                append_more_paths(more_picker.GetPaths())
            finally:
                more_picker.Destroy()

        add_more.Bind(wx.EVT_BUTTON, add_more_photos)
        sizer.Add(add_more, 0, wx.ALL | wx.ALIGN_LEFT, 10)
        buttons = dialog.CreateButtonSizer(wx.OK | wx.CANCEL)
        sizer.Add(buttons, 0, wx.ALL | wx.ALIGN_RIGHT, 10)
        dialog.SetSizerAndFit(sizer)
        dialog.SetSize((600, 450))
        try:
            if dialog.ShowModal() != wx.ID_OK:
                return
            selected = [
                path
                for index, path in enumerate(paths)
                if checklist.IsItemChecked(index)
            ]
        finally:
            dialog.Destroy()
        if not selected:
            return
        self.prepare_output_buffer()
        self.busy = True
        self.busy_status_visible = False
        self.batch_cancel_event = threading.Event()
        wx.CallLater(1000, announce, "Working.")
        self.update_controls()
        threading.Thread(
            target=self._batch_import_worker,
            args=(selected,),
            name="ScanBox batch import",
            daemon=True,
        ).start()

    def _batch_import_worker(self, paths):
        completed = 0
        for source in paths:
            if self.batch_cancel_event.is_set():
                break
            try:
                path = self.copy_imported_image(source)
                text = self.process_photo(
                    path, cancel_event=self.batch_cancel_event
                )
                if text == PHOTO_DESCRIPTION_CANCELLED:
                    break
                if not text.startswith("Processing failed:"):
                    add_photo_description(source, text)
                wx.CallAfter(self._finish_batch_item, source, text, completed + 1, len(paths))
            except Exception as exc:
                logger.exception("Batch photo import failed for %s", source)
                wx.CallAfter(self._finish_batch_item, source, f"Processing failed: {exc}", completed + 1, len(paths))
            completed += 1
        wx.CallAfter(self._finish_batch_done)

    def cancel_batch_import(self, event=None):
        if getattr(self, "batch_cancel_event", None) is not None:
            self.batch_cancel_event.set()
            self.SetStatusText("Cancelling batch photo import...")
            announce("Cancelling batch photo import.")

    def _finish_batch_item(self, source, text, number, total):
        self.output_box.Show(True)
        self.output_box.Enable()
        self.append_output(f"{os.path.basename(source)}\n{text}\n\n")
        self.output_box.SetFocusFromKbd()
        self.Layout()
        self.SetStatusText(f"Batch photo {number} of {total}")

    def _finish_batch_done(self):
        was_cancelled = (
            self.batch_cancel_event is not None
            and self.batch_cancel_event.is_set()
        )
        self.busy = False
        self.batch_cancel_event = None
        self.photo_batch_import_btn.Enable(True)
        self.refresh_photo_library()
        self.update_controls()
        if (
            self.session_pages
            and self.mode_tabs.GetSelection() != self.TAB_LIBRARY
        ):
            self.output_box.SetFocusFromKbd()
        announce(
            "Batch photo import cancelled."
            if was_cancelled
            else "Batch photo import complete."
        )
        if not was_cancelled:
            wx.MessageBox(
                "All selected photos have been processed and stored in the Photo Library.",
                "Batch photo import complete",
            )

    def copy_imported_image(self, source):
        ext = os.path.splitext(source)[1].lower() or ".jpg"
        dest = os.path.join(TEMP_DIR, f"import_{uuid.uuid4().hex}{ext}")
        try:
            with Image.open(source) as img:
                img = ImageOps.exif_transpose(img)
                img.save(dest)
            logger.info("Imported image source=%s working_copy=%s", source, dest)
            return dest
        except UnidentifiedImageError as exc:
            raise ValueError("that file is not a supported image or PDF.") from exc

    def process_pdf(self, source):
        open_word = bool(
            sys.platform == "win32"
            and self.app_settings.get("open_word_after_pdf_conversion", False)
            and microsoft_word_available()
        )
        show_text = self.app_settings.get("render_converted_in_window", True)
        if not open_word and not show_text:
            raise ValueError(
                "Choose a PDF reading destination in Settings before importing "
                "a PDF."
            )
        if fitz is None and sys.platform != "darwin":
            raise ValueError("PDF page reading support is not installed in ScanBox.")

        self.prepare_output_buffer()
        # Only called from Import. On Windows the file dialog is an
        # in-process wx.FileDialog that never hands focus away, so
        # restore_focus=False avoids a redundant, announcement-clipping
        # re-focus. On macOS the picker is a separate native helper process
        # (see macos_choose_files) that genuinely takes foreground app
        # status - without restoring focus here, it can be left stuck on
        # whatever app was active before the picker (e.g. the Terminal
        # window a test build was launched from), with no obvious way back
        # to ScanBox for a screen reader user.
        self.begin_busy("Please wait.", restore_focus=sys.platform == "darwin")
        threading.Thread(
            target=self._process_pdf_worker,
            args=(source, open_word),
            daemon=True,
        ).start()

    def _process_pdf_worker(self, source, open_word):
        try:
            ocr_mode = self.app_settings.get("pdf_ocr_mode", "none")
            if not open_word:
                pages = (
                    self.read_pdf_pages_with_native_ocr(
                        source, all_pages=ocr_mode == "all_pages"
                    )
                    if ocr_mode != "none"
                    else self.read_pdf_pages(source)
                )
                wx.CallAfter(self._finish_pdf_text, pages)
                return

            if ocr_mode != "none":
                pages = self.read_pdf_pages_with_native_ocr(
                    source, all_pages=ocr_mode == "all_pages"
                )
                output_path = self.write_pdf_pages_to_docx(source, pages)
            else:
                self.pdf_suitable_for_pdf2word(source)
                output_path = self.convert_pdf_with_pdf2word(source)
            wx.CallAfter(self._finish_pdf_conversion, output_path)
        except PdfNeedsOcrError as exc:
            wx.CallAfter(self._finish_pdf_needs_ocr, str(exc))
        except Exception as exc:
            wx.CallAfter(self._finish_pdf_error, f"Processing failed: {exc}")
        finally:
            if os.path.basename(source).startswith("browser_"):
                _remove_quietly(source)

    def render_pdf_page(self, page, page_number):
        scale = 2.0
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        path = os.path.join(
            TEMP_DIR, f"pdf_{page_number + 1}_{uuid.uuid4().hex}.png"
        )
        pixmap.save(path)
        return path

    def pdf_suitable_for_pdf2word(self, source):
        """Confirm that a non-OCR PDF has selectable text for conversion."""
        with fitz.open(source) as pdf:
            if pdf.page_count == 0:
                raise ValueError("the PDF contains no pages.")
            indices = sorted({0, pdf.page_count // 2, pdf.page_count - 1})
            has_selectable_text = any(
                pdf.load_page(index).get_text("text").strip()
                for index in indices
            )
            if not has_selectable_text:
                raise PdfNeedsOcrError(
                    "This PDF doesn't appear to contain any readable "
                    "text without OCR. Choose an OCR option on the Import tab "
                    "and try again."
                )
        return True

    def read_pdf_pages(self, source):
        """Return only existing selectable PDF text, without running OCR."""
        if sys.platform == "darwin":
            return self.read_pdf_pages_macos(source, ocr_mode="none")
        pages = []
        with fitz.open(source) as pdf:
            if pdf.page_count == 0:
                raise ValueError("the PDF contains no pages.")
            for index in range(pdf.page_count):
                page = pdf.load_page(index)
                selectable = page.get_text("text").strip()
                pages.append(selectable)
        if not any(pages):
            raise PdfNeedsOcrError(
                "This PDF doesn't appear to contain any readable text without "
                "OCR. Choose an OCR option on the Import tab and try again."
            )
        return pages

    def read_pdf_pages_with_native_ocr(self, source, all_pages=False):
        """Read scanned pages with Windows OCR or Apple Vision."""
        if sys.platform == "darwin":
            return self.read_pdf_pages_macos(
                source, ocr_mode="all_pages" if all_pages else "missing_text"
            )
        pages = []
        with fitz.open(source) as pdf:
            if pdf.page_count == 0:
                raise ValueError("the PDF contains no pages.")
            for index in range(pdf.page_count):
                page = pdf.load_page(index)
                selectable = page.get_text("text").strip()
                if selectable and not all_pages:
                    pages.append(selectable)
                    continue
                image_path = self.render_pdf_page(page, index)
                try:
                    if sys.platform == "darwin":
                        text = macos_ocr(image_path)
                    else:
                        text = windows_ocr(image_path)
                finally:
                    _remove_quietly(image_path)
                text = (text or "").strip()
                if not text:
                    text = selectable
                if not text:
                    engine_name = (
                        "Apple Vision" if sys.platform == "darwin" else "Windows OCR"
                    )
                    text = f"This page could not be read by {engine_name}."
                pages.append(text)
        return pages

    def read_pdf_pages_macos(self, source, ocr_mode="none"):
        """Read PDF pages with PDFKit and, when requested, Apple Vision."""
        result = subprocess.run(
            [
                MACOS_CAPTURE_HELPER,
                "read-pdf",
                source,
                {
                    "none": "text",
                    "missing_text": "ocr-missing",
                    "all_pages": "ocr-all",
                }[ocr_mode],
            ],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=600,
        )
        pages = []
        for line in result.stdout.splitlines():
            kind, separator, encoded = line.partition("\t")
            if kind != "page" or not separator:
                continue
            text = base64.b64decode(encoded).decode("utf-8", errors="replace").strip()
            if not text and ocr_mode != "none":
                text = "This page could not be read by Apple Vision."
            pages.append(text)
        if not pages:
            raise ValueError("the PDF contains no pages.")
        if ocr_mode == "none" and not any(pages):
            raise PdfNeedsOcrError(
                "This PDF doesn't appear to contain any readable text without "
                "OCR. Choose an OCR option on the Import tab and try again."
            )
        return pages

    def write_pdf_pages_to_docx(self, source, pages):
        output_path = self.pdf2word_output_path(source)
        document = docx.Document()
        for index, text in enumerate(pages):
            if index:
                document.add_page_break()
            for block in xml_safe_text(text).split("\n\n"):
                document.add_paragraph(block)
        document.save(output_path)
        return output_path

    def convert_pdf_with_pdf2word(self, source):
        output_path = self.pdf2word_output_path(source)
        messages = []
        pdf_to_docx(
            source,
            output_path,
            log_fn=messages.append,
            use_ocr=False,
            use_tables=True,
        )
        logger.debug("PDF-to-Word output for %s: %s", source, " | ".join(messages))
        return output_path

    def pdf2word_output_path(self, source):
        base = os.path.splitext(os.path.basename(source))[0].strip() or "document"
        base = "".join(ch if ch not in '<>:"/\\|?*' else "_" for ch in base)
        output_path = os.path.join(OUTPUT_DIR, f"{base}.docx")
        if not os.path.exists(output_path):
            return output_path
        for index in range(2, 1000):
            candidate = os.path.join(OUTPUT_DIR, f"{base} {index}.docx")
            if not os.path.exists(candidate):
                return candidate
        return os.path.join(OUTPUT_DIR, f"{base} {uuid.uuid4().hex}.docx")

    def _finish_pdf_conversion(self, output_path):
        self.end_busy()
        self.page_counter += 1
        self.created_output_files.append(output_path)

        message = f"PDF converted to Word:\n{output_path}"
        self.session_pages.append(message)
        self.session_page_modes.append("document")
        self.append_output(message + "\n\n")
        announce("PDF converted to Word.")
        self.open_document_file(output_path)
        self.update_controls()

    def _finish_pdf_text(self, pages):
        self.end_busy()
        rendered_pages = []
        show_headings = self.app_settings.get("show_page_headings", False)
        for text in pages:
            self.page_counter += 1
            self.session_pages.append(text)
            self.session_page_modes.append("document")
            if show_headings:
                rendered_pages.append(
                    f"--- Document Page {self.page_counter} ---\n{text}"
                )
            else:
                rendered_pages.append(text)
        rendered = "\n\n".join(rendered_pages).strip()
        displayed = bool(
            rendered
            and self.app_settings.get("render_converted_in_window", True)
        )
        if displayed:
            self.append_output(rendered + "\n\n")
        elif rendered:
            self.SetStatusText("PDF reading completed.")
        self.update_controls()
        if displayed:
            # begin_busy parks focus on the frame while the import control is
            # disabled. Return it to the completed reading, matching the image
            # and batch completion paths, so screen readers expose the result.
            self.output_box.SetFocusFromKbd()
            announce("PDF reading completed.")

    def show_local_ai_acceleration(self, event=None):
        """Show the processor used by the selected local AI model."""
        report = local_ai_acceleration_report(
            self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
        )
        logger.info("Local AI acceleration status: %s", report.replace("\n", " | "))
        wx.MessageBox(
            report,
            "Local AI Acceleration",
            wx.OK | wx.ICON_INFORMATION,
            self,
        )

    def open_document_file(self, path):
        try:
            open_with_default_application(path)
        except Exception:
            pass

    def _finish_pdf_error(self, message):
        self.end_busy()
        wx.MessageBox(message, "Import failed")

    def _finish_pdf_needs_ocr(self, message):
        # Deliberately does not touch session_pages/output_box/created_
        # output_files - nothing was produced, so there's nothing to show
        # or open, just the explanation.
        self.end_busy()
        wx.MessageBox(message, "This PDF needs OCR")

    def show_photo_cancel_dialog(self):
        self.close_photo_cancel_dialog()
        dialog = wx.Dialog(
            self,
            title="Processing",
            style=wx.DEFAULT_DIALOG_STYLE,
        )
        cancel = wx.Button(dialog, label="Cancel")
        cancel.SetName("Cancel")
        cancel.SetToolTip("Stop processing this photograph")
        cancel.Bind(wx.EVT_BUTTON, self.cancel_photo_description)
        dialog.Bind(wx.EVT_CLOSE, self.cancel_photo_description)
        sizer = wx.BoxSizer(wx.VERTICAL)
        sizer.Add(cancel, 1, wx.ALL | wx.EXPAND, 24)
        dialog.SetSizer(sizer)
        dialog.SetSize(self.GetSize())
        dialog.SetPosition(self.GetPosition())
        self.photo_cancel_dialog = dialog
        self.photo_cancel_button = cancel
        dialog.Show()
        dialog.Raise()
        wx.CallAfter(cancel.SetFocus)

    def cancel_photo_description(self, event=None):
        if self.photo_cancel_event is None:
            return
        self.photo_cancel_event.set()
        if self.camera_capture_active:
            self.camera_capture_stopping = True
        button = getattr(self, "photo_cancel_button", None)
        if button is not None:
            button.SetLabel("Cancelling")
            button.Enable(False)
        self.SetStatusText("Cancelling...")
        if isinstance(event, wx.CloseEvent) and event.CanVeto():
            event.Veto()

    def close_photo_cancel_dialog(self):
        dialog = self.photo_cancel_dialog
        self.photo_cancel_dialog = None
        self.photo_cancel_button = None
        if dialog is not None:
            dialog.Destroy()

    def process_image(
        self,
        path,
        is_photo_mode,
        source_kind,
        library_path=None,
        prepare_buffer=True,
        completion=None,
    ):
        if prepare_buffer:
            self.prepare_output_buffer()
        self.page_counter += 1
        self.session_files.append(path)
        self.session_file_sources.append(source_kind)
        self.session_file_modes.append("photo" if is_photo_mode else "document")

        # source_kind == "import" skips the re-focus on Windows, where the
        # file dialog is in-process and never hands focus away. macOS's
        # Import uses a separate native picker process that does take
        # foreground app status (see macos_choose_files / process_pdf
        # above), so it still needs restore_focus there.
        self.begin_busy(
            "" if source_kind == "camera" else "Please wait.",
            restore_focus=source_kind != "import" or sys.platform == "darwin",
        )
        if is_photo_mode:
            self.photo_cancel_event = threading.Event()
            self.show_photo_cancel_dialog()
        threading.Thread(
            target=self._process_image_worker,
            args=(
                path,
                self.page_counter,
                is_photo_mode,
                library_path or path,
                source_kind == "import",
                completion,
                self.photo_cancel_event if is_photo_mode else None,
            ),
            daemon=True,
        ).start()

    def _process_image_worker(
        self,
        path,
        page_number,
        is_photo_mode,
        library_path,
        store_in_library,
        completion,
        cancel_event,
    ):
        try:
            if is_photo_mode:
                text = self.process_photo(path, cancel_event=cancel_event)
            else:
                text = self.process_document(
                    path,
                    allow_install_prompt=False,
                    page_number=page_number,
                )
        except Exception as exc:
            logger.exception("Image processing failed for %s", path)
            text = f"Processing failed: {exc}"
        logger.info(
            "Processed image page=%s photo_mode=%s path=%s\nresult:\n%s",
            page_number,
            is_photo_mode,
            path,
            text,
        )
        wx.CallAfter(
            self._finish_image_processing,
            page_number,
            is_photo_mode,
            text,
            library_path,
            store_in_library,
            completion,
        )

    def _finish_image_processing(
        self,
        page_number,
        is_photo_mode,
        text,
        source_path,
        store_in_library=True,
        completion=None,
    ):
        self.close_photo_cancel_dialog()
        self.photo_cancel_event = None
        if text == PHOTO_DESCRIPTION_CANCELLED:
            self.end_busy()
            self.SetStatusText(PHOTO_DESCRIPTION_CANCELLED)
            if completion is not None:
                wx.CallAfter(completion, text)
            return
        if is_photo_mode:
            rendered = text + "\n\n"
        elif self.app_settings.get("show_page_headings", False):
            rendered = f"--- Document Page {page_number} ---\n{text}\n\n"
        else:
            rendered = text + "\n\n"
        self.end_busy()
        self.session_pages.append(text)
        self.session_page_modes.append("photo" if is_photo_mode else "document")
        if is_photo_mode and not text.startswith("Processing failed:"):
            self.session_photo_descriptions[source_path] = text
        if (
            is_photo_mode
            and store_in_library
            and not text.startswith("Processing failed:")
        ):
            try:
                add_photo_description(source_path, text)
                self.refresh_photo_library()
            except OSError:
                logger.exception("Could not store photo description for %s", source_path)
        self.append_output(rendered)
        self.output_box.SetFocusFromKbd()
        # Camera sequences announce through their completion callback. A
        # single scanner/import operation has no callback and must announce
        # its result here regardless of whether it is a document or photo.
        if completion is None:
            wx.CallLater(1000, announce, text)
        self.update_controls()
        if completion is not None:
            wx.CallAfter(completion, text)

    def process_document(
        self, path, allow_install_prompt=True, detect_page=True, page_number=None,
    ):
        # Camera-based captures (a document camera, or a handheld/desktop
        # scanner like a Pearl/IRIScan-style device) photograph the page
        # against a desk/background rather than scanning it edge-to-edge.
        # Find and straighten the page before OCR. Safe no-op on flatbed/WIA
        # scans, which are already edge-to-edge.
        if detect_page:
            detect_and_crop_page(path)
        # Choosing another local AI is an explicit request to use that model
        # for document OCR, not merely as a fallback after native OCR.
        external_attempted = external_ai_config(self.app_settings) is not None
        if external_attempted:
            vision_text = run_vision_task("transcribe", path)
            if (
                not vision_text.startswith("Local AI service")
                and not looks_like_ocr_repetition_garbage(vision_text)
            ):
                return vision_text
            logger.warning(
                "External local AI OCR failed; falling back to native OCR: %s",
                vision_text,
            )
        if sys.platform == "win32":
            native_text = windows_ocr(path).strip()
            native_ocr_name = "Windows.Media.Ocr"
        elif sys.platform == "darwin":
            native_text = macos_ocr(path).strip()
            native_ocr_name = "Apple Vision"
        else:
            native_text = ""
            native_ocr_name = ""
        if native_text:
            logger.info("Document OCR used %s", native_ocr_name)
            return native_text

        if not external_attempted and vision_ready("transcribe"):
            vision_text = run_vision_task("transcribe", path)
            if not looks_like_ocr_repetition_garbage(vision_text):
                return vision_text
            logger.info(
                "Vision transcribe result for %s also looked like repetition "
                "garbage; falling back to the unreadable-page message.",
                path,
            )

        if not allow_install_prompt:
            return (
                "This document could not be read reliably. Install the local AI "
                "pack for improved document reading."
            )
        if self.offer_install_vision(
            "This document could not be read reliably. The local AI pack can "
            "provide a better reading.",
            DEFAULT_VISION_MODEL_ID,
        ):
            return (
                "The local AI pack is being installed. Scan or import this page "
                "again after installation."
            )
        return "This document could not be read reliably."

    def process_photo(self, path, include_ocr=True, cancel_event=None):
        # Runs on a worker thread, so it must not show any wx dialog. The
        # main-thread photo_mode_blocked() check already prompts to install
        # before we get here; this is just a defensive fallback.
        if not vision_ready():
            return "Photo description needs the local AI model. It is not installed."

        jpg_path = self.ensure_jpg(path)

        # Auto-crop is deliberately NOT run here. General photos (people,
        # scenes) very often contain strong rectangular shapes in the
        # background - door frames, windows, buildings - and cropping to
        # the wrong rectangle silently discards the actual subject with no
        # way for a user who can't see the photo to notice. It's only safe
        # to assume "the frame is one page" for a document scan/import.

        # No auto-rotation here. The local model's own "is this upright,
        # yes/no" self-judgment turned out to be unreliable - it rotated at
        # least one already-correct photo (a book cover) because it
        # confidently misjudged its own orientation. There's no dependable
        # way to detect this with a model this size, so it isn't attempted.

        recognised_text = ""
        # Native OCR (Windows.Media.Ocr / Apple Vision) is a separate,
        # model-independent pipeline - it has no dependency on which local
        # AI describes the photo, so it's worth running alongside either
        # model rather than only alongside Qwen3-VL.
        if include_ocr and sys.platform in {"win32", "darwin"}:
            native_ocr = windows_ocr if sys.platform == "win32" else macos_ocr
            recognised_text = native_ocr(jpg_path).strip()
            recognised_words = re.findall(r"[A-Za-z0-9]+", recognised_text)
            recognised_characters = sum(len(word) for word in recognised_words)
            if len(recognised_words) < 3 or recognised_characters < 10:
                recognised_text = ""
        if cancel_event is not None and cancel_event.is_set():
            return PHOTO_DESCRIPTION_CANCELLED
        description = run_vision_task("describe", jpg_path, cancel_event)
        if description == PHOTO_DESCRIPTION_CANCELLED:
            return description
        if recognised_text:
            description += "\n\nRecognised text:\n" + recognised_text
        return description

    def ensure_jpg(self, path):
        jpg_path = os.path.splitext(path)[0] + ".jpg"
        if os.path.abspath(path).lower() == os.path.abspath(jpg_path).lower():
            return path
        with Image.open(path) as img:
            ImageOps.exif_transpose(img).convert("RGB").save(jpg_path, "JPEG", quality=95)
        return jpg_path

    def append_output(self, text):
        self.output_box.Enable()
        start_pos = self.output_box.GetLastPosition()
        self.output_box.AppendText(text)
        self.output_box.SetInsertionPoint(start_pos)
        self.output_box.ShowPosition(start_pos)

    def prepare_output_buffer(self):
        """Start a fresh result unless the user enabled multi-page appending."""
        if self.app_settings.get("append_text_to_buffer", False):
            return
        self.session_files = []
        self.session_file_sources = []
        self.session_file_modes = []
        self.session_pages = []
        self.session_page_modes = []
        self.session_photo_descriptions = {}
        self.page_counter = 0
        self.output_box.Clear()
        self.output_box.Disable()

    def open_settings(self, event=None):
        dlg = wx.Dialog(self, title="Settings")
        root = wx.BoxSizer(wx.VERTICAL)

        notebook = wx.Notebook(dlg)

        general_panel = wx.Panel(notebook)
        _set_named_page_accessible(general_panel, "General")
        general_sizer = wx.BoxSizer(wx.VERTICAL)
        delete_output = wx.CheckBox(general_panel, label="Delete output files upon exit")
        pdf_destination_choices = ["Show PDF reading in ScanBox"]
        if sys.platform == "win32":
            pdf_destination_choices.append(
                "Create a DOCX and open it in Microsoft Word"
            )
        pdf_destination = wx.RadioBox(
            general_panel,
            label="PDF reading destination",
            choices=pdf_destination_choices,
            majorDimension=1,
            style=wx.RA_SPECIFY_ROWS,
        )
        append_buffer = wx.CheckBox(
            general_panel,
            label="Append text to buffer for each Scan or Import",
        )
        show_page_headings = wx.CheckBox(
            general_panel,
            label="Show document page headings in results",
        )
        diagnostic_logging = wx.CheckBox(
            general_panel, label="Enable diagnostic logging"
        )
        check_updates_on_startup = wx.CheckBox(
            general_panel, label="Check for updates on startup"
        )
        diagnostic_logging.SetToolTip(
            f"Write scan, import, and local AI details to {LOG_FILE}"
        )
        delete_output.SetValue(self.app_settings.get("delete_output_files_on_exit", False))
        word_available = (
            sys.platform == "win32" and microsoft_word_available()
        )
        pdf_destination.SetSelection(
            1
            if word_available
            and self.app_settings.get("open_word_after_pdf_conversion", False)
            else 0
        )
        if sys.platform == "win32":
            pdf_destination.EnableItem(1, word_available)
        if sys.platform == "win32" and not word_available:
            pdf_destination.SetToolTip(
                "Microsoft Word was not found; PDF readings will be shown in ScanBox."
            )
        append_buffer.SetValue(
            self.app_settings.get("append_text_to_buffer", False)
        )
        show_page_headings.SetValue(
            self.app_settings.get("show_page_headings", False)
        )
        diagnostic_logging.SetValue(
            self.app_settings.get("diagnostic_logging", False)
        )
        check_updates_on_startup.SetValue(
            self.app_settings.get("check_for_updates_on_startup", True)
        )
        general_sizer.Add(delete_output, 0, wx.ALL, 10)
        general_sizer.Add(
            check_updates_on_startup, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10
        )
        general_sizer.Add(pdf_destination, 0, wx.ALL | wx.EXPAND, 10)
        general_sizer.Add(append_buffer, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        general_sizer.Add(
            show_page_headings, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10
        )
        general_sizer.Add(
            diagnostic_logging, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10
        )
        general_panel.SetSizer(general_sizer)
        notebook.AddPage(general_panel, "General")

        scanner_panel = wx.Panel(notebook)
        _set_named_page_accessible(scanner_panel, "Scanner")
        scanner_sizer = wx.BoxSizer(wx.VERTICAL)
        saved_scanner_id = self.app_settings.get("scanner_id", "")
        saved_scanner_name = self.app_settings.get(
            "scanner_name", "Selected scanner"
        )
        scanner_options = (
            [{"id": saved_scanner_id, "name": saved_scanner_name}]
            if saved_scanner_id
            else []
        )
        scanner_choices = ["Ask me each time"] + [
            scanner["name"] for scanner in scanner_options
        ]
        scanner_choice = wx.Choice(scanner_panel, choices=scanner_choices)
        scanner_choice.SetName("Choose default scanner")
        scanner_choice.SetSelection(1 if scanner_options else 0)
        scanner_status = wx.StaticText(scanner_panel, label="Looking for scanners…")
        scanner_status.SetName("Scanner discovery status")
        refresh_scanners_btn = wx.Button(scanner_panel, label="Find Scanners")
        scanner_sizer.Add(
            wx.StaticText(scanner_panel, label="Choose default scanner"),
            0, wx.LEFT | wx.RIGHT | wx.TOP, 10,
        )
        scanner_sizer.Add(scanner_choice, 0, wx.ALL | wx.EXPAND, 10)
        scanner_sizer.Add(scanner_status, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        scanner_sizer.Add(
            refresh_scanners_btn, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10
        )
        scanner_panel.SetSizer(scanner_sizer)
        notebook.AddPage(scanner_panel, "Scanner")

        def finish_scanner_discovery(scanners):
            nonlocal scanner_options
            try:
                previous_id = (
                    scanner_options[scanner_choice.GetSelection() - 1]["id"]
                    if scanner_choice.GetSelection() > 0
                    else ""
                )
                saved_missing = bool(
                    saved_scanner_id
                    and not any(
                        scanner["id"] == saved_scanner_id for scanner in scanners
                    )
                )
                scanner_options = (
                    [{"id": saved_scanner_id, "name": saved_scanner_name}]
                    if saved_missing
                    else []
                ) + scanners
                scanner_choice.Clear()
                scanner_choice.Append("Ask me each time")
                names = [scanner["name"] for scanner in scanner_options]
                for index, scanner in enumerate(scanner_options):
                    label = scanner["name"]
                    if saved_missing and index == 0:
                        label += " (not currently detected)"
                    if names.count(label) > 1:
                        label += f" ({index + 1})"
                    scanner_choice.Append(label)
                selected = next(
                    (
                        index + 1 for index, scanner in enumerate(scanner_options)
                        if scanner["id"] == previous_id
                    ),
                    0,
                )
                scanner_choice.SetSelection(selected)
                scanner_status.SetLabel(
                    f"Found {len(scanners)} scanner"
                    + ("." if len(scanners) == 1 else "s.")
                    if scanners
                    else "No scanners were found."
                )
                refresh_scanners_btn.Enable(True)
                scanner_panel.Layout()
            except RuntimeError:
                pass

        def find_scanners(event=None):
            refresh_scanners_btn.Enable(False)
            scanner_status.SetLabel("Looking for scanners…")

            def worker():
                wx.CallAfter(finish_scanner_discovery, self.available_scanners())

            threading.Thread(
                target=worker,
                name="ScanBox scanner discovery",
                daemon=True,
            ).start()

        refresh_scanners_btn.Bind(wx.EVT_BUTTON, find_scanners)
        find_scanners()

        if sys.platform == "darwin":
            permissions_panel = wx.Panel(notebook)
            _set_named_page_accessible(permissions_panel, "Permissions")
            permissions_sizer = wx.BoxSizer(wx.VERTICAL)
            permissions_help = wx.StaticText(
                permissions_panel,
                label=(
                    "Open ScanBox's Mac permissions guide. Use this if "
                    "screen description, global shortcuts or camera capture "
                    "are not working."
                ),
            )
            permissions_help.Wrap(480)
            open_permissions_btn = wx.Button(
                permissions_panel, label="Open Mac Permissions"
            )
            open_permissions_btn.Bind(wx.EVT_BUTTON, self.show_macos_permissions)
            permissions_sizer.Add(
                permissions_help, 0, wx.ALL | wx.EXPAND, 10
            )
            permissions_sizer.Add(
                open_permissions_btn, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10
            )
            permissions_panel.SetSizer(permissions_sizer)
            notebook.AddPage(permissions_panel, "Permissions")

        camera_panel = wx.Panel(notebook)
        _set_named_page_accessible(camera_panel, "Camera")
        camera_sizer = wx.BoxSizer(wx.VERTICAL)
        camera_help = wx.StaticText(
            camera_panel,
            label=(
                "These settings apply whenever OCR using Camera or a photo "
                "camera command is started. Set number of captures to zero "
                "for continuous capture."
            ),
        )
        camera_help.Wrap(480)
        camera_sizer.Add(camera_help, 0, wx.ALL | wx.EXPAND, 10)

        camera_delay_label = wx.StaticText(
            camera_panel, label="Delay before first capture, in seconds"
        )
        camera_delay = wx.SpinCtrl(camera_panel, min=0, max=120)
        camera_delay.SetName("Delay before first capture, in seconds")
        camera_delay.SetValue(
            max(
                0,
                min(
                    120,
                    int(self.app_settings.get("camera_delay_seconds", 5)),
                ),
            )
        )
        camera_sizer.Add(camera_delay_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        camera_sizer.Add(camera_delay, 0, wx.ALL | wx.EXPAND, 10)

        camera_interval_label = wx.StaticText(
            camera_panel,
            label="Delay after a completed result before the next capture, in seconds",
        )
        camera_interval = wx.SpinCtrl(camera_panel, min=1, max=3600)
        camera_interval.SetName(
            "Delay after a completed result before the next capture, in seconds"
        )
        camera_interval.SetValue(
            max(
                1,
                min(
                    3600,
                    int(self.app_settings.get("camera_interval_seconds", 5)),
                ),
            )
        )
        camera_sizer.Add(
            camera_interval_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10
        )
        camera_sizer.Add(camera_interval, 0, wx.ALL | wx.EXPAND, 10)

        camera_count_label = wx.StaticText(
            camera_panel,
            label="Number of captures; enter zero for continuous capture",
        )
        camera_count = wx.SpinCtrl(camera_panel, min=0, max=100)
        camera_count.SetName(
            "Number of captures; enter zero for continuous capture"
        )
        camera_count.SetValue(
            max(
                0,
                min(
                    100,
                    int(self.app_settings.get("camera_capture_count", 1)),
                ),
            )
        )
        camera_sizer.Add(camera_count_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        camera_sizer.Add(camera_count, 0, wx.ALL | wx.EXPAND, 10)
        camera_panel.SetSizer(camera_sizer)
        notebook.AddPage(camera_panel, "Camera")

        ai_panel = wx.Panel(notebook)
        _set_named_page_accessible(ai_panel, "AI")
        ai_sizer = wx.BoxSizer(wx.VERTICAL)

        ai_source = wx.RadioBox(
            ai_panel,
            label="AI source",
            choices=[
                "ScanBox local AI (recommended)",
                "Another local AI on this computer",
            ],
            majorDimension=1,
            style=wx.RA_SPECIFY_ROWS,
        )
        ai_source.SetSelection(
            1 if self.app_settings.get("ai_provider") == "external" else 0
        )
        ai_sizer.Add(ai_source, 0, wx.ALL | wx.EXPAND, 10)

        model_ids = list(VISION_MODELS)
        model_labels = []
        for model_id in model_ids:
            model = VISION_MODELS[model_id]
            installed = (
                _find_florence_files() is not None
                if model["runner"] == "florence"
                else _find_mtmd_model_files(model_id) is not None
            )
            model_labels.append(
                model.get("choice_label", model["name"])
                + ("; installed" if installed else "; not installed")
            )
        model_choice = wx.Choice(ai_panel, choices=model_labels)
        model_choice.SetName("Image description model")
        current_model = self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
        model_choice.SetSelection(model_ids.index(current_model) if current_model in model_ids else 0)
        ai_sizer.Add(wx.StaticText(ai_panel, label="Image description model"), 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        ai_sizer.Add(model_choice, 0, wx.ALL | wx.EXPAND, 10)

        find_local_ai_btn = wx.Button(ai_panel, label="Find Local AI")
        found_ai_choice = wx.Choice(ai_panel, choices=[])
        found_ai_choice.SetName("Detected local AI models")
        external_results = []
        saved_external = external_ai_config(self.app_settings)
        if saved_external:
            external_results.append({
                **saved_external,
                "label": (
                    self.app_settings.get("external_ai_label", "Local AI")
                    + ": " + saved_external["model"]
                ),
            })
            found_ai_choice.Append(external_results[0]["label"])
            found_ai_choice.SetSelection(0)
        external_status = wx.StaticText(
            ai_panel,
            label=("Saved local AI connection." if saved_external else ""),
        )
        external_status.SetName("Local AI discovery status")
        external_url_label = wx.StaticText(ai_panel, label="Server address (advanced)")
        external_url = wx.TextCtrl(
            ai_panel,
            value=self.app_settings.get("external_ai_url", ""),
        )
        external_url.SetName("Local AI server address")
        external_model_label = wx.StaticText(ai_panel, label="Model name (advanced)")
        external_model = wx.TextCtrl(
            ai_panel,
            value=self.app_settings.get("external_ai_model", ""),
        )
        external_model.SetName("Local AI model name")
        external_url_label.Hide()
        external_url.Hide()
        external_model_label.Hide()
        external_model.Hide()
        ai_sizer.Add(find_local_ai_btn, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        ai_sizer.Add(
            wx.StaticText(ai_panel, label="Detected local AI models"),
            0, wx.LEFT | wx.RIGHT | wx.TOP, 10,
        )
        ai_sizer.Add(found_ai_choice, 0, wx.ALL | wx.EXPAND, 10)
        ai_sizer.Add(external_status, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        ai_sizer.Add(external_url_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        ai_sizer.Add(external_url, 0, wx.ALL | wx.EXPAND, 10)
        ai_sizer.Add(external_model_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        ai_sizer.Add(external_model, 0, wx.ALL | wx.EXPAND, 10)

        install_ai_btn = wx.Button(ai_panel, label="Install or update local AI model")
        cancel_install_btn = wx.Button(ai_panel, label="Cancel AI Download")
        delete_model_btn = wx.Button(ai_panel, label="Delete local AI model")
        install_status = wx.StaticText(ai_panel, label="")
        install_status.SetName("AI download status")
        install_progress = wx.Gauge(ai_panel, range=100)
        install_status.Hide()
        install_progress.Hide()
        self._ai_install_status = install_status
        self._ai_install_progress = install_progress

        def model_is_installed(model_id):
            model = VISION_MODELS[model_id]
            return (
                _find_florence_files() is not None
                if model["runner"] == "florence"
                else (
                    _find_mtmd_model_files(model_id) is not None
                    and _find_mtmd_runner() is not None
                )
            )

        def refresh_model_controls():
            builtin_selected = ai_source.GetSelection() == 0
            for index, model_id in enumerate(model_ids):
                if self.installing and self._installing_model_id == model_id:
                    state = "installing"
                else:
                    state = "installed" if model_is_installed(model_id) else "not installed"
                model = VISION_MODELS[model_id]
                model_choice.SetString(
                    index, f"{model.get('choice_label', model['name'])}; {state}"
                )
            selected_id = model_ids[model_choice.GetSelection()]
            delete_model_btn.Enable(
                builtin_selected and model_is_installed(selected_id)
            )
            install_ai_btn.Enable(builtin_selected and not self.installing)
            cancel_install_btn.Show(self.installing)
            cancel_install_btn.Enable(self.installing)
            model_choice.Enable(builtin_selected and not self.installing)
            find_local_ai_btn.Enable(not builtin_selected)
            found_ai_choice.Enable(not builtin_selected)
            external_url.Enable(not builtin_selected)
            external_model.Enable(not builtin_selected)
            ai_panel.Layout()

        self._refresh_ai_model_controls = refresh_model_controls
        refresh_model_controls()

        def select_external_result(event=None):
            selected = found_ai_choice.GetSelection()
            if 0 <= selected < len(external_results):
                item = external_results[selected]
                external_url.SetValue(item["url"])
                external_model.SetValue(item["model"])
                external_status.SetLabel(f"Selected {item['label']}.")
                ai_panel.Layout()

        def find_local_ai(event=None):
            nonlocal external_results
            external_status.SetLabel("Looking for local AI services…")
            ai_panel.Layout()
            wx.SafeYield(ai_panel, onlyIfNeeded=True)
            results = discover_local_ai_models()
            external_results = results
            found_ai_choice.Clear()
            for item in results:
                found_ai_choice.Append(item["label"])
            if results:
                wanted = (
                    self.app_settings.get("external_ai_kind", ""),
                    self.app_settings.get("external_ai_url", ""),
                    self.app_settings.get("external_ai_model", ""),
                )
                selected = next(
                    (
                        index for index, item in enumerate(results)
                        if (item["kind"], item["url"], item["model"]) == wanted
                    ),
                    0,
                )
                found_ai_choice.SetSelection(selected)
                select_external_result()
                external_status.SetLabel(
                    f"Found {len(results)} local AI model"
                    + ("." if len(results) == 1 else "s.")
                )
                found_ai_choice.SetFocusFromKbd()
            else:
                external_status.SetLabel(
                    "No local AI service was found. Start it and try again, "
                    "or enter its server address and model name."
                )
                external_url_label.Show()
                external_url.Show()
                external_model_label.Show()
                external_model.Show()
            ai_panel.Layout()

        ai_source.Bind(wx.EVT_RADIOBOX, lambda event: refresh_model_controls())
        find_local_ai_btn.Bind(wx.EVT_BUTTON, find_local_ai)
        found_ai_choice.Bind(wx.EVT_CHOICE, select_external_result)

        def on_install_model(event):
            selected_index = model_choice.GetSelection()
            selected_id = model_ids[selected_index]
            self.app_settings["vision_model"] = selected_id
            write_app_settings(self.app_settings)
            self.install_local_ai(event)
            if self.installing:
                model_choice.SetString(
                    selected_index,
                    f"{VISION_MODELS[selected_id].get('choice_label', VISION_MODELS[selected_id]['name'])}; installing",
                )
                model_choice.Enable(False)
                install_ai_btn.Enable(False)
                cancel_install_btn.Show(True)
                cancel_install_btn.Enable(True)
                delete_model_btn.Enable(False)
                install_status.SetLabel("Downloading model…")
                install_progress.SetValue(0)
                install_status.Show()
                install_progress.Show()
                ai_panel.Layout()

        def on_delete_model(event):
            selected_id = model_ids[model_choice.GetSelection()]
            selected = VISION_MODELS[selected_id]
            installed = (_find_florence_files() is not None if selected["runner"] == "florence"
                         else _find_mtmd_model_files(selected_id) is not None)
            if not installed:
                wx.MessageBox("That model is not installed.", "Delete Model")
                return
            if wx.MessageBox(
                f"Delete {selected['name']} from this computer? "
                "You can reinstall it later from 'Install or update local AI "
                "model'. Settings will close after this.",
                "Delete Model",
                wx.YES_NO | wx.ICON_WARNING,
            ) != wx.YES:
                return
            target_dir = os.path.join(VISION_DIR, selected["subdir"])
            try:
                if selected_id == "qwen3_vl_2b":
                    stop_mtmd_server()
                shutil.rmtree(target_dir)
            except OSError as exc:
                wx.MessageBox(
                    f"Could not delete the local AI model: {exc}", "Delete Model"
                )
                return
            if selected["runner"] == "florence":
                _florence_cache.pop("base", None)
            wx.MessageBox(
                "The local AI model has been deleted. Settings will now "
                "close; reopen it to continue.",
                "Delete Model",
            )
            dlg.EndModal(wx.ID_CANCEL)

        install_ai_btn.Bind(wx.EVT_BUTTON, on_install_model)
        cancel_install_btn.Bind(wx.EVT_BUTTON, self.cancel_ai_download)
        delete_model_btn.Bind(wx.EVT_BUTTON, on_delete_model)
        model_choice.Bind(wx.EVT_CHOICE, lambda event: refresh_model_controls())
        ai_sizer.Add(install_ai_btn, 0, wx.ALL, 10)
        ai_sizer.Add(cancel_install_btn, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        ai_sizer.Add(install_status, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        ai_sizer.Add(install_progress, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM | wx.EXPAND, 10)
        ai_sizer.Add(delete_model_btn, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        ai_panel.SetSizer(ai_sizer)
        notebook.AddPage(ai_panel, "AI")

        buttons = dlg.CreateButtonSizer(wx.OK | wx.CANCEL)
        root.Add(notebook, 1, wx.ALL | wx.EXPAND, 10)
        root.Add(buttons, 0, wx.ALL | wx.ALIGN_RIGHT, 10)
        dlg.SetSizer(root)
        dlg.SetSize((540, 440))
        dlg.Fit()

        try:
            if dlg.ShowModal() == wx.ID_OK:
                self.app_settings["delete_output_files_on_exit"] = delete_output.GetValue()
                self.app_settings["check_for_updates_on_startup"] = (
                    check_updates_on_startup.GetValue()
                )
                self.app_settings["minimize_to_notification_area"] = False
                open_in_word = (
                    word_available and pdf_destination.GetSelection() == 1
                )
                self.app_settings["open_word_after_pdf_conversion"] = open_in_word
                self.app_settings["render_converted_in_window"] = not open_in_word
                self.app_settings["append_text_to_buffer"] = append_buffer.GetValue()
                self.app_settings["show_page_headings"] = (
                    show_page_headings.GetValue()
                )
                self.app_settings["camera_delay_seconds"] = (
                    camera_delay.GetValue()
                )
                self.app_settings["camera_interval_seconds"] = (
                    camera_interval.GetValue()
                )
                self.app_settings["camera_capture_count"] = (
                    camera_count.GetValue()
                )
                scanner_selection = scanner_choice.GetSelection()
                selected_scanner = (
                    scanner_options[scanner_selection - 1]
                    if 0 < scanner_selection <= len(scanner_options)
                    else None
                )
                self.app_settings["scanner_id"] = (
                    selected_scanner["id"] if selected_scanner else ""
                )
                self.app_settings["scanner_name"] = (
                    selected_scanner["name"]
                    if selected_scanner else "Ask me each time"
                )
                logging_changed = (
                    self.app_settings.get("diagnostic_logging", False)
                    != diagnostic_logging.GetValue()
                )
                self.app_settings["diagnostic_logging"] = diagnostic_logging.GetValue()
                self.app_settings["vision_model"] = model_ids[model_choice.GetSelection()]
                self.app_settings["ai_provider"] = (
                    "external" if ai_source.GetSelection() == 1 else "builtin"
                )
                if self.app_settings["ai_provider"] == "external":
                    selected_external = found_ai_choice.GetSelection()
                    selected_item = (
                        external_results[selected_external]
                        if 0 <= selected_external < len(external_results)
                        else None
                    )
                    self.app_settings["external_ai_kind"] = (
                        selected_item["kind"]
                        if selected_item
                        else _local_ai_kind_for_url(external_url.GetValue())
                    )
                    self.app_settings["external_ai_label"] = (
                        selected_item["name"] if selected_item else "Local AI"
                    )
                    self.app_settings["external_ai_url"] = _normalise_local_ai_url(
                        external_url.GetValue()
                    )
                    self.app_settings["external_ai_model"] = (
                        external_model.GetValue().strip()
                    )
                    if external_ai_config(self.app_settings) is None:
                        self.app_settings["ai_provider"] = "builtin"
                        wx.MessageBox(
                            "The local AI connection was incomplete, so ScanBox "
                            "will continue using its built-in AI.",
                            "Local AI not selected",
                            wx.OK | wx.ICON_INFORMATION,
                            self,
                        )
                write_app_settings(self.app_settings)
                if self.app_settings["ai_provider"] == "external":
                    stop_mtmd_server()
                elif self.app_settings["vision_model"] == "qwen3_vl_2b":
                    threading.Thread(
                        target=start_mtmd_server,
                        args=("qwen3_vl_2b",),
                        name="ScanBox AI preload",
                        daemon=True,
                    ).start()
                else:
                    stop_mtmd_server()
                if logging_changed:
                    configure_logging(diagnostic_logging.GetValue())
                    self.SetStatusText(
                        "Diagnostic logging "
                        + ("on" if diagnostic_logging.GetValue() else "off")
                    )
                self.update_controls()
        finally:
            if getattr(self, "_refresh_ai_model_controls", None) is refresh_model_controls:
                self._refresh_ai_model_controls = None
            dlg.Destroy()

    def on_close(self, event):
        global _mac_announce_process
        if getattr(self, "_close_after_announcement", False):
            self._close_after_announcement = False
        else:
            processing = self.installing or self.busy or self.image_export_active
            if processing:
                # Closing the application is also an emergency stop. Signal all
                # cancellable workers and leave their temporary inputs untouched;
                # a clean sweep at the next startup is safer than trapping the user
                # in an application they may be closing because it stopped responding.
                self._skip_exit_file_cleanup = True
                self.shutdown_event.set()
                if self.install_cancel_event is not None:
                    self.install_cancel_event.set()
                if self.photo_cancel_event is not None:
                    self.photo_cancel_event.set()
                if getattr(self, "batch_cancel_event", None) is not None:
                    self.batch_cancel_event.set()
                logger.info("Close requested during processing; cancelling work and exiting")
            self.browser_url_executor.shutdown(wait=False, cancel_futures=True)
            # A delayed capture or FaceAlign check is not represented by self.busy,
            # so stop those workflows before the delayed exit announcement.
            self.camera_capture_timer.Stop()
            self.camera_capture_active = False
            self.close_camera_capture_device()
            if self.camera_alignment_stop_event is not None:
                self.camera_alignment_stop_event.set()
            self.camera_alignment_active = False
            self.close_facealign_dialog()
            announce("Exiting.")
            if sys.platform != "darwin" and event.CanVeto():
                self._close_after_announcement = True
                event.Veto()
                wx.CallLater(500, self.Close)
                return
        self.camera_capture_timer.Stop()
        self.camera_capture_active = False
        self.close_camera_capture_device()
        if self.camera_alignment_stop_event is not None:
            self.camera_alignment_stop_event.set()
            self.camera_alignment_stop_event = None
        self.camera_alignment_active = False
        self.close_facealign_dialog()
        stop_mtmd_server()
        if not self._skip_exit_file_cleanup:
            clear_temp_directory()
        if (
            not self._skip_exit_file_cleanup
            and self.app_settings.get("delete_output_files_on_exit")
        ):
            self.delete_created_output_files()
        if sys.platform == "win32":
            try:
                self.UnregisterHotKey(self.screen_hotkey_id)
            except Exception:
                pass
            try:
                self.UnregisterHotKey(self.screen_ocr_hotkey_id)
            except Exception:
                pass
            try:
                self.UnregisterHotKey(self.screen_question_hotkey_id)
            except Exception:
                pass
            try:
                self.UnregisterHotKey(self.browser_pdf_hotkey_id)
            except Exception:
                pass
            try:
                self.UnregisterHotKey(self.toggle_window_hotkey_id)
            except Exception:
                pass
        if self._mac_helper_process is not None:
            _mac_announce_process = None
            try:
                self._mac_helper_process.terminate()
                self._mac_helper_process.wait(timeout=2)
            except subprocess.TimeoutExpired:
                try:
                    self._mac_helper_process.kill()
                except Exception:
                    pass
            except Exception:
                pass
            self._mac_helper_process = None
        if self.tray_icon is not None:
            self.tray_icon.RemoveIcon()
            self.tray_icon.Destroy()
            self.tray_icon = None
        if sys.platform == "darwin":
            logger.info("macOS close cleanup finished; exiting process")
            logging.shutdown()
            try:
                self.Destroy()
            except Exception:
                pass
            os._exit(0)
            return
        event.Skip()

    def delete_created_output_files(self):
        # A real sweep of the output folder, rather than only the paths in
        # self.created_output_files (an in-memory list scoped to this one
        # run). That old approach meant anything left behind by a session
        # that crashed, was force-quit, or simply ended before this one
        # started was never touched by any later session, no matter how
        # many times ScanBox was restarted - the "delete on exit" setting
        # only ever cleaned up after itself. This clears everything
        # actually sitting in Output except the active log file.
        output_root = os.path.abspath(OUTPUT_DIR)
        log_path = os.path.abspath(LOG_FILE)
        try:
            entries = os.listdir(output_root)
        except OSError:
            return
        for name in entries:
            target = os.path.join(output_root, name)
            if os.path.abspath(target) == log_path:
                continue
            try:
                if os.path.isfile(target):
                    os.remove(target)
            except OSError:
                pass
        self.created_output_files = []

    def save_text(self, event):
        if event is not None:
            wx.CallAfter(self.save_text, None)
            return
        merged = "\n\n".join(
            text
            for text, mode in zip(self.session_pages, self.session_page_modes)
            if mode == "document"
        )
        dlg = wx.FileDialog(
            self,
            "Save text",
            defaultDir=OUTPUT_DIR,
            # Leave the extension off so the selected file-type filter controls
            # whether this becomes TXT, DOCX or PDF unless the user explicitly
            # types an extension of their own.
            defaultFile=suggested_document_filename(merged, ""),
            wildcard="Text (*.txt)|*.txt|Word (*.docx)|*.docx|PDF (*.pdf)|*.pdf",
            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        formats = ["txt", "docx", "pdf"]
        try:
            if dlg.ShowModal() == wx.ID_OK:
                path = dlg.GetPath()
                fmt = os.path.splitext(path)[1].lower().lstrip(".")
                if fmt not in formats:
                    # No (or unknown) extension typed: use the chosen file type.
                    fmt = formats[dlg.GetFilterIndex()]
                    path = path + "." + fmt
                save_text_output(merged, path, fmt)
        finally:
            dlg.Destroy()

    def save_images(self, event):
        if event is not None:
            wx.CallAfter(self.save_images, None)
            return
        active_mode = self.last_active_mode
        scanned_paths = [
            path
            for path, source, mode in zip(
                self.session_files,
                self.session_file_sources,
                self.session_file_modes,
            )
            if source in {"scan", "camera"} and mode == active_mode
        ]
        if not scanned_paths:
            return
        if len(scanned_paths) == 1:
            default_name = (
                "Captured photo.jpg"
                if active_mode == "photo"
                else "Captured document.jpg"
            )
            dlg = wx.FileDialog(
                self,
                "Save scanned image",
                defaultDir=IMAGES_DIR,
                defaultFile=default_name,
                wildcard="JPEG image (*.jpg)|*.jpg",
                style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
            )
            try:
                if dlg.ShowModal() == wx.ID_OK:
                    destination = self.write_scanned_jpeg(
                        scanned_paths[0],
                        dlg.GetPath(),
                    )
                    if active_mode == "photo":
                        description = self.session_photo_descriptions.get(
                            scanned_paths[0]
                        )
                        if description:
                            add_photo_description(destination, description)
                        else:
                            update_photo_library_path(scanned_paths[0], destination)
                        self.refresh_photo_library()
            finally:
                dlg.Destroy()
            return

        name_dlg = wx.TextEntryDialog(self, "Base filename for scanned images:")
        try:
            if name_dlg.ShowModal() != wx.ID_OK:
                return
            base = self.safe_filename_stem(name_dlg.GetValue())
        finally:
            name_dlg.Destroy()
        folder_dlg = wx.DirDialog(
            self, "Choose a folder for scanned images", defaultPath=IMAGES_DIR
        )
        try:
            if folder_dlg.ShowModal() != wx.ID_OK:
                return
            folder = folder_dlg.GetPath()
        finally:
            folder_dlg.Destroy()
        for index, image_path in enumerate(scanned_paths, 1):
            target = self.available_image_path(
                folder, f"{base} {index}.jpg"
            )
            destination = self.write_scanned_jpeg(image_path, target)
            if active_mode == "photo":
                description = self.session_photo_descriptions.get(image_path)
                if description:
                    add_photo_description(destination, description)
                else:
                    update_photo_library_path(image_path, destination)
        if active_mode == "photo":
            self.refresh_photo_library()

    def safe_filename_stem(self, value):
        cleaned = "".join(
            character if character not in '<>:"/\\|?*' else "_"
            for character in value.strip()
        ).rstrip(" .")
        return cleaned or "Scanned image"

    def available_image_path(self, folder, filename):
        target = os.path.join(folder, filename)
        if not os.path.exists(target):
            return target
        stem, extension = os.path.splitext(filename)
        for number in range(2, 1000):
            candidate = os.path.join(folder, f"{stem} ({number}){extension}")
            if not os.path.exists(candidate):
                return candidate
        return os.path.join(folder, f"{stem} {uuid.uuid4().hex}{extension}")

    def write_scanned_jpeg(self, source, destination, rotation=None):
        stem, extension = os.path.splitext(destination)
        if extension.lower() not in {".jpg", ".jpeg"}:
            destination = (stem if extension else destination) + ".jpg"
        with Image.open(source) as img:
            img = ImageOps.exif_transpose(img)
            transpose = {
                90: Image.Transpose.ROTATE_270,
                180: Image.Transpose.ROTATE_180,
                270: Image.Transpose.ROTATE_90,
            }.get(rotation)
            if transpose is not None:
                img = img.transpose(transpose)
            img.convert("RGB").save(
                destination, "JPEG", quality=95
            )
        return destination

    def write_scanned_image_pdf(self, sources, destination, rotations=None):
        stem, extension = os.path.splitext(destination)
        if extension.lower() != ".pdf":
            destination = (stem if extension else destination) + ".pdf"
        rotations = rotations or {}
        pages = []
        try:
            for source in sources:
                with Image.open(source) as img:
                    img = ImageOps.exif_transpose(img)
                    transpose = {
                        90: Image.Transpose.ROTATE_270,
                        180: Image.Transpose.ROTATE_180,
                        270: Image.Transpose.ROTATE_90,
                    }.get(rotations.get(source, 0))
                    if transpose is not None:
                        img = img.transpose(transpose)
                    pages.append(img.convert("RGB"))
            if not pages:
                raise ValueError("No scanned pages were available to save.")
            pages[0].save(
                destination,
                "PDF",
                save_all=True,
                append_images=pages[1:],
                resolution=300.0,
            )
        finally:
            for page in pages:
                page.close()
        return destination

    def install_local_ai(self, event):
        if self.installing:
            wx.MessageBox("The local AI model is already downloading.", "ScanBox")
            return
        model_id = self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
        model_name = VISION_MODELS.get(model_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])["name"]
        if wx.MessageBox(
            f"Install {model_name} for local image descriptions? Images remain "
            "on this computer.",
            "Install Local AI",
            wx.YES_NO | wx.ICON_QUESTION,
        ) != wx.YES:
            return
        self.start_install(model_id)

    def offer_install_vision(self, reason, model_id=None):
        """Ask whether to download the local AI model and start it if agreed.
        Returns True if a download was started (or is already running)."""
        if self.installing:
            return True
        selected_id = model_id or self.app_settings.get("vision_model", DEFAULT_VISION_MODEL_ID)
        selected_model = VISION_MODELS.get(selected_id, VISION_MODELS[DEFAULT_VISION_MODEL_ID])
        if wx.MessageBox(
            reason + f" ScanBox is about to download {selected_model['name']}. "
            "Install it now? Processing remains on this computer.",
            "Install Local AI",
            wx.YES_NO | wx.ICON_QUESTION,
        ) != wx.YES:
            return False
        self.start_install(selected_id)
        return True

    def start_install(self, model_id=None):
        if self.installing:
            return
        self.installing = True
        self.install_cancel_event = threading.Event()
        self._installing_model_id = model_id or self.app_settings.get(
            "vision_model", DEFAULT_VISION_MODEL_ID
        )
        self.install_gauge.SetValue(0)
        self.install_gauge.Show()
        self.cancel_ai_download_btn.Enable(True)
        self.cancel_ai_download_btn.Show()
        self.Layout()
        self.SetStatusText("Installing local AI pack. Downloading...")
        announce(
            "Installing local AI. Downloading. This can take several minutes."
        )
        model_id = self._installing_model_id
        threading.Thread(
            target=self._run_install,
            args=(model_id, self.install_cancel_event),
            daemon=True,
        ).start()

    def _run_install(self, model_id, cancel_event):
        def status(message):
            wx.CallAfter(self._install_status, message)

        result = install_local_ai_pack(status, model_id, cancel_event)
        wx.CallAfter(self._install_done, result)

    def cancel_ai_download(self, event=None):
        if not self.installing or self.install_cancel_event is None:
            return
        self.install_cancel_event.set()
        self.cancel_ai_download_btn.Enable(False)
        if event is not None and hasattr(event, "GetEventObject"):
            control = event.GetEventObject()
            if control is not None:
                control.Enable(False)
        self.SetStatusText("Cancelling AI download...")

    def _install_status(self, message):
        if self.install_cancel_event is not None and self.install_cancel_event.is_set():
            return
        progress = re.search(r":\s*(\d+)%\s+of\s+", message)
        if progress:
            percentage = int(progress.group(1))
            self.install_gauge.SetValue(percentage)
            # The native progress event is continuous; keep speech and the text
            # log concise enough to remain useful.
            if percentage not in (0, 100) and percentage % 25:
                return
        else:
            self.install_gauge.Pulse()
        self.SetStatusText(message)
        status = getattr(self, "_ai_install_status", None)
        progress_control = getattr(self, "_ai_install_progress", None)
        try:
            if status is not None:
                status.SetLabel(message)
            if progress_control is not None:
                if progress:
                    progress_control.SetValue(percentage)
                else:
                    progress_control.Pulse()
        except RuntimeError:
            self._ai_install_status = None
            self._ai_install_progress = None
        announce(message)

    def _install_done(self, result):
        installed_model_id = self._installing_model_id
        self.installing = False
        self._installing_model_id = None
        self.install_cancel_event = None
        if result.startswith("Local AI pack installed"):
            self.install_gauge.SetValue(100)
        else:
            self.install_gauge.SetValue(0)
        self.install_gauge.Hide()
        self.cancel_ai_download_btn.Hide()
        self.Layout()
        refresh_controls = getattr(self, "_refresh_ai_model_controls", None)
        if refresh_controls is not None:
            try:
                refresh_controls()
            except RuntimeError:
                self._refresh_ai_model_controls = None
        self.SetStatusText(result)
        status = getattr(self, "_ai_install_status", None)
        progress_control = getattr(self, "_ai_install_progress", None)
        try:
            if status is not None:
                status.SetLabel(result)
            if progress_control is not None:
                progress_control.SetValue(100 if result.startswith("Local AI pack installed") else 0)
        except RuntimeError:
            self._ai_install_status = None
            self._ai_install_progress = None
        announce(result)
        if (
            result.startswith("Local AI pack installed")
            and installed_model_id == "qwen3_vl_2b"
            and self.app_settings.get("vision_model") == "qwen3_vl_2b"
        ):
            threading.Thread(
                target=start_mtmd_server,
                args=("qwen3_vl_2b",),
                name="ScanBox AI preload",
                daemon=True,
            ).start()
        # Completion is reported through the status bar and screen reader.
        # A modal dialog can become stranded behind another application and
        # interfere with returning to ScanBox through Alt+Tab.
        if self._close_when_install_stops:
            self._close_when_install_stops = False
            self.Close()

    def reset_session(self, event=None):
        for file_path in self.session_files:
            try:
                temp_root = os.path.abspath(TEMP_DIR)
                target = os.path.abspath(file_path)
                if os.path.commonpath([temp_root, target]) == temp_root and os.path.exists(target):
                    os.remove(target)
            except OSError:
                pass

        self.session_files = []
        self.session_file_sources = []
        self.session_file_modes = []
        self.session_pages = []
        self.session_page_modes = []
        self.session_photo_descriptions = {}
        self.page_counter = 0
        self.SetStatusText("No active session")
        self.output_box.Clear()
        self.output_box.Disable()

        self.update_controls()
        wx.CallAfter(self._restore_current_tab_focus)

    def close_or_reset_session(self, event=None):
        is_import_photo_session = bool(
            self.last_active_mode == "photo"
            and self.session_file_sources
            and all(
                source == "import" and mode == "photo"
                for source, mode in zip(
                    self.session_file_sources, self.session_file_modes
                )
            )
        )
        if is_import_photo_session:
            # Move accessibility focus before Close is renamed, disabled and
            # hidden; otherwise screen readers announce the stale control as
            # "Discard Session unavailable".
            self.photo_import_btn.SetFocus()
            wx.CallAfter(self.reset_session)
            return
        self.reset_session()

    def _restore_current_tab_focus(self):
        selected_tab = self.mode_tabs.GetSelection()
        if selected_tab == self.TAB_IMPORT:
            control = (
                self.photo_import_btn
                if self.last_active_mode == "photo"
                else self.document_import_btn
            )
        elif selected_tab == self.TAB_LIBRARY:
            control = self.library_list
        else:
            control = (
                self.photo_scan_btn
                if self.last_active_mode == "photo"
                else self.document_scan_btn
            )
        try:
            control.SetFocus()
        except RuntimeError:
            self.mode_tabs.SetFocus()

    def choose_another_photo(self, event=None):
        """Close the current imported-photo result and reopen its picker."""
        self.reset_session()
        wx.CallAfter(self.on_import_image, None, True)


def main():
    # Guard against two copies running at once - a second copy would
    # silently confuse someone using a screen reader (which window is
    # which?), and two processes could also collide over the same
    # temp/output files or a local AI download in progress.
    instance_checker = wx.SingleInstanceChecker(f"ScanBox-{wx.GetUserId()}")
    if instance_checker.IsAnotherRunning():
        app = wx.App(False)
        existing_hwnd = (
            ctypes.windll.user32.FindWindowW(None, "ScanBox")
            if sys.platform == "win32"
            else None
        )
        if existing_hwnd:
            SW_RESTORE = 9
            ctypes.windll.user32.ShowWindow(existing_hwnd, SW_RESTORE)
            ctypes.windll.user32.SetForegroundWindow(existing_hwnd)
            message = "ScanBox is already running. Switching to the existing window."
        else:
            message = (
                "ScanBox is already running. Check the Dock or notification "
                "area for the open window."
            )
        wx.MessageBox(message, "ScanBox", wx.OK | wx.ICON_INFORMATION)
        return

    # A forced close deliberately leaves in-use temporary inputs alone. Once
    # the single-instance check proves no older ScanBox is still using them,
    # they are safe to remove at the beginning of this new session.
    clear_temp_directory()
    app = wx.App(False)
    app.SetExitOnFrameDelete(True)
    frame = ScanBox()
    frame.Show()
    if sys.platform == "darwin":
        _macos_activate_self()
        wx.CallAfter(frame.focus_scan_tab_on_startup)
    try:
        app.MainLoop()
    finally:
        stop_mtmd_server()


def _macos_activate_self():
    """Force ScanBox to become the frontmost/active application.

    Double-clicking an app in Finder, or launching it with `open`, makes
    macOS activate it automatically. A bare executable started directly
    from a shell - as test_macos.sh does with `nohup` - does not get that
    same automatic activation, so keyboard focus (and with it, things like
    Cmd+Q) can silently stay with whichever app launched it instead, even
    though ScanBox's window is visibly on screen. This asks macOS directly
    to make ScanBox the active app, regardless of how it was started.
    """
    if NSRunningApplication is None:
        return
    try:
        NSRunningApplication.currentApplication().activateWithOptions_(
            NSApplicationActivateIgnoringOtherApps
        )
    except Exception:
        logger.exception("Could not activate ScanBox")


if __name__ == "__main__":
    main()
