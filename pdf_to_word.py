"""Convert selectable-text PDF files to accessible Word documents.

This is the restored source for the converter previously distributed only as
``engines/pdf2word/system/python/pdf_to_word_cli.exe``.  ScanBox calls the
module directly; it can also be run from a command prompt.

Copyright (C) 2026 Sam Taylor
SPDX-License-Identifier: AGPL-3.0-or-later
"""

from __future__ import annotations

import argparse
import os
import re
import time
from collections.abc import Callable

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


LogFunction = Callable[[str], None]


def sanitise(text: str) -> str:
    """Remove control characters that are invalid in Word XML."""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.find(qn("w:tblPr"))
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def add_fitz_table(document: Document, fitz_table) -> None:
    try:
        data = fitz_table.extract()
    except Exception:
        return
    rows = [row for row in data if any(cell is not None and str(cell).strip() for cell in row)]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    word_table = document.add_table(rows=len(rows), cols=column_count)
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            value = row[column_index] if column_index < len(row) else ""
            text = str(value).strip() if value is not None else ""
            cell = word_table.cell(row_index, column_index)
            cell.text = sanitise(text)
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    set_table_borders(word_table)
    document.add_paragraph()


def ocr_page_windows(pdf_path: str, page_index: int) -> str:
    """Use the Windows OCR service as an optional scanned-page fallback."""
    import asyncio

    from winsdk.windows.data.pdf import PdfDocument, PdfPageRenderOptions
    from winsdk.windows.foundation import Uri
    from winsdk.windows.graphics.imaging import BitmapDecoder
    from winsdk.windows.media.ocr import OcrEngine
    from winsdk.windows.storage.streams import InMemoryRandomAccessStream

    async def run() -> str:
        path_uri = Uri("file:///" + pdf_path.replace("\\", "/"))
        pdf_document = await PdfDocument.load_from_uri_async(path_uri)
        if page_index >= pdf_document.page_count:
            return ""
        page = pdf_document.get_page(page_index)
        image_stream = InMemoryRandomAccessStream()
        options = PdfPageRenderOptions()
        options.destination_width = 1700
        await page.render_to_stream_async(image_stream, options)
        image_stream.seek(0)
        decoder = await BitmapDecoder.create_async(image_stream)
        bitmap = await decoder.get_software_bitmap_async()
        engine = OcrEngine.try_create_from_user_profile_languages()
        if engine is None:
            return ""
        result = await engine.recognize_async(bitmap)
        return result.text

    return asyncio.run(run())


def _rectangles_overlap(first, second) -> bool:
    ax0, ay0, ax1, ay1 = first
    bx0, by0, bx1, by1 = second
    return ax0 < bx1 and ax1 > bx0 and ay0 < by1 and ay1 > by0


def convert_page(
    page,
    document: Document,
    pdf_path: str,
    page_index: int,
    use_ocr: bool = False,
    use_tables: bool = True,
    log_fn: LogFunction = print,
) -> None:
    text = (page.get_text("text") or "").strip()
    if not text:
        if use_ocr:
            log_fn(f"  Page {page_index + 1}: scanned, running OCR...")
            try:
                text = ocr_page_windows(pdf_path, page_index).strip()
                if text:
                    paragraph = document.add_paragraph()
                    paragraph.add_run(f"[Page {page_index + 1} - OCR]").bold = True
                    for line in text.splitlines():
                        if line.strip():
                            document.add_paragraph(sanitise(line.strip()))
            except Exception as exc:
                log_fn(f"  OCR failed: {exc}")
        return

    tables = []
    table_boxes = []
    added_tables: set[int] = set()
    if use_tables:
        started = time.time()
        try:
            drawings = page.get_drawings()
            # Table detection is most useful on pages with a moderate number
            # of ruled lines. Avoid expensive detection on graphic-heavy pages.
            if 3 < len(drawings) < 200:
                found = page.find_tables()
                if found and found.tables:
                    tables = found.tables
                    table_boxes = [table.bbox for table in tables]
        except Exception:
            pass
        elapsed = time.time() - started
        if elapsed > 3:
            log_fn(f"  Page {page_index + 1}: slow table detection ({round(elapsed, 1)}s)")

    page_dict = page.get_text("dict", flags=11)
    blocks = sorted(
        page_dict["blocks"],
        key=lambda block: (round(block["bbox"][1] / 5) * 5, block["bbox"][0]),
    )

    for block in blocks:
        block_box = block["bbox"]
        if table_boxes:
            for index, table_box in enumerate(table_boxes):
                if _rectangles_overlap(block_box, table_box):
                    if index not in added_tables:
                        add_fitz_table(document, tables[index])
                        added_tables.add(index)
                    break
            if any(
                _rectangles_overlap(block_box, table_boxes[index])
                for index in added_tables
            ):
                continue

        if block["type"] != 0:
            continue

        line_texts = [
            "".join(span["text"] for span in line["spans"])
            for line in block["lines"]
        ]

        # Rejoin words split by a PDF line wrap, except where the following
        # line looks like a heading.
        joined = []
        index = 0
        while index < len(line_texts):
            line_text = line_texts[index]
            while (
                line_text.rstrip().endswith("-")
                and index + 1 < len(line_texts)
                and line_texts[index + 1].strip()
                and not line_texts[index + 1].strip()[0].isupper()
            ):
                line_text = line_text.rstrip()[:-1] + line_texts[index + 1].strip()
                index += 1
            joined.append(line_text)
            index += 1

        first_span = None
        if block["lines"] and block["lines"][0]["spans"]:
            first_span = block["lines"][0]["spans"][0]
        for line_text in joined:
            line_text = line_text.strip()
            if not line_text:
                continue
            paragraph = document.add_paragraph()
            run = paragraph.add_run(sanitise(line_text))
            if first_span:
                run.font.size = Pt(round(first_span["size"]))
                run.bold = bool(first_span["flags"] & 16)


def pdf_to_docx(
    pdf_path: str,
    output_path: str,
    log_fn: LogFunction = print,
    use_ocr: bool = False,
    use_tables: bool = True,
) -> None:
    pdf = fitz.open(pdf_path)
    try:
        total = len(pdf)
        log_fn(f"TOTAL_PAGES:{total}")
        document = Document()
        for paragraph in document.paragraphs:
            paragraph._element.getparent().remove(paragraph._element)
        for index, page in enumerate(pdf):
            log_fn(f"({index + 1}/{total}) Page {index + 1}")
            if index > 0:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            convert_page(
                page,
                document,
                pdf_path,
                index,
                use_ocr,
                use_tables,
                log_fn,
            )
        document.save(output_path)
        log_fn(f"Done. Saved: {output_path}")
    finally:
        pdf.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="PDF to Word converter")
    parser.add_argument("input", help="Input PDF path")
    parser.add_argument("output", help="Output DOCX path")
    parser.add_argument(
        "--ocr",
        action="store_true",
        help="Enable OCR fallback for scanned pages (Windows 10/11)",
    )
    parser.add_argument(
        "--no-tables",
        action="store_true",
        help="Skip table detection for faster conversion",
    )
    args = parser.parse_args(argv)
    if not os.path.isfile(args.input):
        print(f"ERROR: File not found: {args.input}", flush=True)
        return 1
    try:
        pdf_to_docx(
            args.input,
            args.output,
            lambda message: print(message, flush=True),
            use_ocr=args.ocr,
            use_tables=not args.no_tables,
        )
    except Exception as exc:
        print(f"ERROR: {exc}", flush=True)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
